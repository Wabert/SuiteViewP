"""Strict, read-only CKCVDVPD dividend extract parsing.

The current WL loader uses a 45-character natural HEADER_ID, not an identity.
Valid source dates become ``date`` objects; the explicit CyberLife sentinels
``00/00/1900`` and ``00/00/0000`` remain strings because SQL stores them verbatim.
MAINT_DT belongs to the loading service and is deliberately not generated here.

PUA references must be present and unambiguous in the supplied extract. Parse
the complete extract before filtering keys or effective dates for a load.
Each selected file must be self-contained, including in multi-file imports.
PUA resolution and lookup-only filtering are scoped by USER_CODE and RECORD_TYPE;
there is no implicit cross-user or user-00 fallback.

Import safety: header/rate changes are one atomic dependency group. Preserving
absent database rows is not a replacement: reject duration-range contractions
that would retain stale rates, and do not apply only half of a changed schedule.
Corrections to shared PUA schedules need every affected parent re-denormalized.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
import re


class DividendValidationError(ValueError):
    """The source cannot produce a complete, unambiguous dividend package."""


HEADER_COLUMNS = (
    "HEADER_ID", "RECORD_TYPE", "RECORD_TYPE_DESC",
    "USER_CODE", "DIV_KEY", "CLASS", "BASE_SERIES", "SUBSERIES",
    "CONTENT_CODE", "USER_DEFINED", "ISSUE_AGE",
    "ISSUE_DATE", "EFFECTIVE_DATE", "END_DATE",
    "FIRST_DURATION", "LAST_DURATION", "PRO_RATA", "EI_DEPOSIT",
    "INTEREST_RATE", "RESTRICT_CODE", "EI_ADDITIONS", "MORT_TABLE",
    "INT_FUNCTION", "PUA_INTEREST", "PUA_CLASS", "PUA_MATURE",
    "PUA_PARTICIPATING", "PUA_KEY", "PUA_KEY_USER_DEFINED", "GROSS_INTEREST",
)
RATE_COLUMNS = (
    "HEADER_ID", "RECORD_TYPE", "DURATION", "CASH_RATE", "PUA_RATE", "OYT_RATE",
    "PUA_DIV_CASH", "PUA_DIV_PUA", "PUA_DIV_OYT",
)
PLANKEY_COLUMNS = (
    "PLANCODE", "SEX", "RATECLASS", "SUBSERIES", "USER_KEY", "CV_DIV_KEY",
)
TABLE_KEYS = {
    "WL_DIV_HEADER": ("HEADER_ID",),
    "WL_RATE_DIV": ("HEADER_ID", "DURATION"),
    "WL_DIV_PLANKEY_MAP": ("PLANCODE", "SEX", "RATECLASS"),
}
DATE_COLUMNS = frozenset({"ISSUE_DATE", "EFFECTIVE_DATE", "END_DATE"})
GENERATED_COLUMNS = frozenset({"MAINT_DT"})
RECORD_TYPES = {
    "PREMIUM PAYING DIVIDEND": "D",
    "REDUCED PAID-UP DIVIDEND": "R",
    "DIRECT RECOGNITION DIVIDEND": "L",
    "DIRECT RECOGNITION RPU DIVIDEND": "P",
    "TERMINATION DIVIDEND": "T",
}

_DATE_SENTINELS = frozenset({"00/00/1900", "00/00/0000"})
_DATE = re.compile(r"[0-9]{2}/[0-9]{2}/[0-9]{4}")
_DECIMAL = re.compile(r"-?(?:[0-9]+(?:\.[0-9]+)?|\.[0-9]+)")
_RATE_NUMBER = re.compile(r"-?[0-9]+\.[0-9]{2}")
_RATE_LINE = re.compile(r"(CASH|PUA|1YT)?\s*([0-9]{3})-([0-9]{3})\s*(.*)")
_NULL_PADDING = re.compile(r"\x00+")
_RATE_LABEL_PREFIX = re.compile(r"\s*(?:CASH|PUA|1YT)")
_RATE_RANGE_PREFIX = re.compile(r"\s*(?:CASH|PUA|1YT)?\s*[0-9]{3}-[0-9]{3}")
_BASE_RATE_COLUMNS = ("CASH_RATE", "PUA_RATE", "OYT_RATE")
_PUA_RATE_COLUMNS = ("PUA_DIV_CASH", "PUA_DIV_PUA", "PUA_DIV_OYT")


def _fail(source: str, message: str) -> None:
    raise DividendValidationError(f"{source}: {message}")


def _print_line(raw: str, source: str) -> str:
    """Accept print padding, never repair NUL-corrupted field contents."""
    if "\x00" not in raw:
        return raw.strip()
    for padding in _NULL_PADDING.finditer(raw):
        start, end = padding.span()
        if (
            start == 0 or end == len(raw)
            or raw[start - 1].isspace() or raw[end].isspace()
        ):
            continue
        prefix = raw[:start].replace("\x00", " ")
        suffix = raw[end:]
        # CKCVDVPD prints NUL padding at the label/range and range/value
        # boundaries, including when overflowing values leave no spaces.
        if (
            _RATE_LABEL_PREFIX.fullmatch(prefix)
            and re.match(r"[0-9]{3}-[0-9]{3}", suffix)
        ) or (
            _RATE_RANGE_PREFIX.fullmatch(prefix)
            and _RATE_NUMBER.match(suffix)
        ):
            continue
        _fail(source, "NUL inside a source token or outside verified print padding")
    return raw.replace("\x00", " ").strip()


def _text(value: str, field: str, maximum: int, source: str) -> str:
    text = " ".join(value.split()).upper()
    if not text or len(text) > maximum or not text.isascii():
        _fail(source, f"{field} must contain 1–{maximum} ASCII characters")
    if not re.fullmatch(r"[A-Z0-9 #]+", text):
        _fail(source, f"{field} contains invalid identifier characters")
    return text


def _integer(value: str, field: str, source: str, maximum: int = 32767) -> int:
    if not re.fullmatch(r"[0-9]+", value) or not 0 <= int(value) <= maximum:
        _fail(source, f"{field} must be an integer between 0 and {maximum}")
    return int(value)


def _decimal(value: str, field: str, source: str) -> Decimal:
    if not _DECIMAL.fullmatch(value):
        _fail(source, f"{field} is not a finite decimal")
    return Decimal(value)


def _date(value: str, field: str, source: str) -> date | str:
    if value in _DATE_SENTINELS:
        return value
    if _DATE.fullmatch(value):
        month, day, year = (int(part) for part in value.split("/"))
        try:
            return date(year, month, day)
        except ValueError:
            pass
    _fail(source, f"{field} is not a valid MM/DD/YYYY date or known zero-date sentinel")


def _key_date(value: date | str) -> str:
    if isinstance(value, date):
        return f"{value.year:04d}{value.month:02d}{value.day:02d}"
    month, day, year = value.split("/")
    return year + month + day


def _header_key(header: dict[str, object]) -> str:
    # These spaces are part of the authoritative natural key; never strip it.
    fields = (
        str(header["USER_CODE"]).ljust(2),
        str(header["RECORD_TYPE"]),
        str(header["CLASS"]),
        str(header["BASE_SERIES"]).ljust(3),
        str(header["SUBSERIES"]).ljust(2),
        str(header["CONTENT_CODE"]),
        str(header["USER_DEFINED"] or "").ljust(7),
        f'{header["ISSUE_AGE"]:03d}',
        _key_date(header["ISSUE_DATE"]),
        _key_date(header["EFFECTIVE_DATE"]),
    )
    return "_".join(fields)


def _control_date_indexes(tokens: list[str], source: str) -> list[int]:
    dates = [index for index, token in enumerate(tokens) if "/" in token]
    if len(dates) != 3 or dates[0] < 4 or dates != list(range(dates[0], dates[0] + 3)):
        _fail(source, "malformed control row: expected three consecutive dates")
    return dates


def _control_plan_fields(tokens: list[str], issue_index: int, source: str) -> dict[str, object]:
    plan = tokens[1].split("-")
    if len(plan) != 3:
        _fail(source, "malformed CLASS-BASE-SUB key")
    user_code = _text(tokens[0], "USER_CODE", 2, source)
    if not re.fullmatch(r"[0-9]{2}", user_code):
        _fail(source, "USER_CODE must contain two digits")
    cls = _text(plan[0], "CLASS", 1, source)
    base = _text(plan[1], "BASE_SERIES", 3, source)
    sub = _text(plan[2], "SUBSERIES", 2, source)
    user_defined = " ".join(tokens[3:issue_index - 1])
    return {
        "USER_CODE": user_code,
        "DIV_KEY": cls + base + sub,
        "CLASS": cls,
        "BASE_SERIES": base,
        "SUBSERIES": sub,
        "CONTENT_CODE": _text(tokens[2], "CONTENT_CODE", 1, source),
        "USER_DEFINED": _text(user_defined, "USER_DEFINED", 6, source) if user_defined else None,
        "ISSUE_AGE": _integer(tokens[issue_index - 1], "ISSUE_AGE", source, 999),
    }


def _base_control_header(
    tokens: list[str],
    description: str,
    dates: list[int],
    rest: list[str],
    source: str,
) -> dict[str, object]:
    if len(rest) < 12:
        _fail(source, "incomplete rate parameters in control row")
    issue_index = dates[0]
    header = {
        "RECORD_TYPE": RECORD_TYPES[description],
        "RECORD_TYPE_DESC": description,
        **_control_plan_fields(tokens, issue_index, source),
        "ISSUE_DATE": _date(tokens[issue_index], "ISSUE_DATE", source),
        "EFFECTIVE_DATE": _date(tokens[issue_index + 1], "EFFECTIVE_DATE", source),
        "END_DATE": _date(tokens[issue_index + 2], "END_DATE", source),
        "FIRST_DURATION": _integer(rest[0], "FIRST_DURATION", source),
        "LAST_DURATION": _integer(rest[1], "LAST_DURATION", source),
        "PRO_RATA": _text(rest[2], "PRO_RATA", 2, source),
        "EI_DEPOSIT": _text(rest[3], "EI_DEPOSIT", 2, source),
        "INTEREST_RATE": _decimal(rest[4], "INTEREST_RATE", source),
        "RESTRICT_CODE": _text(rest[5], "RESTRICT_CODE", 2, source),
        "EI_ADDITIONS": _text(rest[6], "EI_ADDITIONS", 2, source),
        "MORT_TABLE": _text(rest[7], "MORT_TABLE", 4, source),
        "INT_FUNCTION": None,
        "PUA_KEY": None,
        "PUA_KEY_USER_DEFINED": None,
        "GROSS_INTEREST": None,
    }
    if header["FIRST_DURATION"] > header["LAST_DURATION"]:
        _fail(source, "FIRST_DURATION exceeds LAST_DURATION")
    return header


def _control_pua_fields(header: dict[str, object], tail: list[str], source: str) -> None:
    if "." not in tail[0]:
        header["INT_FUNCTION"] = _text(tail.pop(0), "INT_FUNCTION", 2, source)
    if len(tail) < 4:
        _fail(source, "incomplete PUA parameters")
    header.update({
        "PUA_INTEREST": _decimal(tail[0], "PUA_INTEREST", source),
        "PUA_CLASS": _text(tail[1], "PUA_CLASS", 2, source),
        "PUA_MATURE": _integer(tail[2], "PUA_MATURE", source),
        "PUA_PARTICIPATING": tail[3],
    })
    if tail[3] not in {"0", "1", "2"}:
        _fail(source, "PUA_PARTICIPATING must be 0, 1 or 2")
    optional = tail[4:]
    if len(optional) > 2:
        _fail(source, "unexpected trailing control fields")
    if len(optional) == 2 or (optional and "." not in optional[0]):
        header["PUA_KEY"] = _text(optional.pop(0), "PUA_KEY", 8, source)
    if optional:
        header["GROSS_INTEREST"] = _decimal(optional[0], "GROSS_INTEREST", source)
    if header["PUA_PARTICIPATING"] == "2" and not header["PUA_KEY"]:
        _fail(source, "participating PUA schedule requires PUA_KEY")


def _control_line(text: str, description: str, source: str) -> dict[str, object]:
    tokens = text.split()
    dates = _control_date_indexes(tokens, source)
    rest = tokens[dates[-1] + 1:]
    header = _base_control_header(tokens, description, dates, rest, source)
    _control_pua_fields(header, rest[8:], source)
    header["HEADER_ID"] = _header_key(header)
    return {column: header[column] for column in HEADER_COLUMNS}


def _rate_values(text: str, source: str) -> list[Decimal]:
    values = []
    for token in text.split():
        # Printed rates have two decimals. Full-token coverage prevents the old
        # regex from silently discarding corruption or a third decimal digit.
        matches = list(_RATE_NUMBER.finditer(token))
        if not matches or "".join(match.group() for match in matches) != token:
            _fail(source, "invalid rate value (expected two decimal places)")
        values.extend(Decimal(match.group()) for match in matches)
    return values


@dataclass
class _Record:
    header: dict[str, object]
    rates: list[dict[str, object]]
    source: str


def _denormalize(records: list[_Record]) -> list[_Record]:
    index, lookup_keys = _dividend_reference_index(records)
    for record in records:
        h = record.header
        participation = h["PUA_PARTICIPATING"]
        if participation == "0":
            continue
        if participation == "1":
            _copy_self_pua(record)
            continue
        _copy_referenced_pua(record, _resolve_pua_reference(record, index))
    result = _base_dividend_records(records, lookup_keys)
    if not result:
        raise DividendValidationError("Extract contains no base dividend schedules after PUA lookup filtering")
    return result


def _dividend_reference_index(records: list[_Record]):
    index = defaultdict(list)
    lookup_keys = set()
    for record in records:
        h = record.header
        index[h["USER_CODE"], h["DIV_KEY"], h["RECORD_TYPE"], h["ISSUE_AGE"]].append(record)
        if h["PUA_KEY"]:
            lookup_keys.add((h["USER_CODE"], h["RECORD_TYPE"], h["PUA_KEY"]))
    return index, lookup_keys


def _copy_self_pua(record: _Record) -> None:
    record.header["PUA_KEY_USER_DEFINED"] = record.header["USER_DEFINED"]
    for row in record.rates:
        row.update(zip(_PUA_RATE_COLUMNS, (row[key] for key in _BASE_RATE_COLUMNS)))


def _resolve_pua_reference(record: _Record, index) -> _Record:
    h = record.header
    reference_scope = (h["USER_CODE"], h["PUA_KEY"], h["RECORD_TYPE"])
    candidates = index.get((*reference_scope, h["ISSUE_AGE"]), [])
    if not candidates:
        candidates = index.get((*reference_scope, 0), [])
    candidates = _prefer_matching_reference(record, candidates)
    if len(candidates) != 1:
        _fail(record.source, "missing or ambiguous PUA reference; supply a complete, unambiguous extract")
    reference = candidates[0]
    if reference is record:
        _fail(record.source, "PUA schedule references itself")
    return reference


def _prefer_matching_reference(record: _Record, candidates: list[_Record]) -> list[_Record]:
    h = record.header
    if len(candidates) > 1:
        same_date = [ref for ref in candidates if ref.header["EFFECTIVE_DATE"] == h["EFFECTIVE_DATE"]]
        if same_date:
            candidates = same_date
    if len(candidates) > 1:
        same_base = [
            ref for ref in candidates
            if str(ref.header["USER_DEFINED"] or "")[:3].strip() == h["BASE_SERIES"]
        ]
        if same_base:
            candidates = same_base
    return candidates


def _copy_referenced_pua(record: _Record, reference: _Record) -> None:
    record.header["PUA_KEY_USER_DEFINED"] = reference.header["USER_DEFINED"]
    reference_rates = {row["DURATION"]: row for row in reference.rates}
    age_shift = record.header["ISSUE_AGE"] if reference.header["ISSUE_AGE"] == 0 else 0
    for row in record.rates:
        reference_row = reference_rates.get(row["DURATION"] + age_shift)
        if reference_row is None:
            _fail(record.source, f'PUA reference is missing duration {row["DURATION"] + age_shift}')
        row.update(zip(_PUA_RATE_COLUMNS, (reference_row[key] for key in _BASE_RATE_COLUMNS)))


def _base_dividend_records(records: list[_Record], lookup_keys: set[tuple]) -> list[_Record]:
    result = [
        record for record in records
        if (
            record.header["USER_CODE"],
            record.header["RECORD_TYPE"],
            record.header["DIV_KEY"],
        ) not in lookup_keys
    ]
    return result


class _DividendParser:
    """State machine for CKCVDVPD dividend print extracts."""

    def __init__(self, path: Path):
        self.path = path
        self.records: dict[str, _Record] = {}
        self.current_header = None
        self.current_type = None
        self.description = None
        self.current_source = str(path)
        self.buffers = {column: {} for column in _BASE_RATE_COLUMNS}
        self.saw_report = False
        self.pending_description = False

    def parse(self) -> dict[str, list[dict[str, object]]]:
        try:
            with self.path.open("r", encoding="utf-8-sig", errors="strict") as handle:
                for number, raw in enumerate(handle, 1):
                    source = f"{self.path.name}:{number}"
                    self._parse_text(_print_line(raw, source), source)
        except UnicodeError as exc:
            raise DividendValidationError(f"{self.path.name}: invalid UTF-8 source text") from exc
        self._flush()
        if not self.records or self.pending_description:
            _fail(self.path.name, "empty or incomplete dividend extract")
        selected = _denormalize(list(self.records.values()))
        return {
            "WL_DIV_HEADER": [record.header for record in selected],
            "WL_RATE_DIV": [row for record in selected for row in record.rates],
        }

    def _parse_text(self, text: str, source: str) -> None:
        if not text or text == "0":
            return
        if text.startswith("1CKCVDVPD"):
            self._handle_report_heading(text, source)
            return
        if not self.saw_report:
            _fail(source, "expected CKCVDVPD report heading")
        if text.startswith("RECORD TYPE"):
            self._handle_record_type(text, source)
            return
        if _is_dividend_formatting_line(text):
            return
        rate = _RATE_LINE.fullmatch(text)
        if rate:
            self._handle_rate_line(rate, source)
            return
        if re.match(r"[0-9]{2}\s+", text):
            self._start_control_row(text, source)
            return
        _fail(source, "unrecognized or malformed dividend source row")

    def _handle_report_heading(self, text: str, source: str) -> None:
        run = re.search(r"RUN DATE\s*=\s*([0-9]{2}/[0-9]{2}/[0-9]{2})(?![0-9])", text)
        if not run or "DIVIDEND RATES" not in text:
            _fail(source, "invalid CKCVDVPD report heading")
        month, day, year = (int(part) for part in run[1].split("/"))
        try:
            date(2000 + year, month, day)
        except ValueError:
            _fail(source, "invalid report run date")
        self.saw_report = True

    def _handle_record_type(self, text: str, source: str) -> None:
        match = re.fullmatch(r"RECORD TYPE\s*=\s*(.+)", text)
        new_description = " ".join(match[1].split()) if match else ""
        if new_description not in RECORD_TYPES:
            _fail(source, "unknown dividend RECORD TYPE")
        if self.current_header and new_description != self.description:
            self._flush()
        self.description = new_description
        self.pending_description = self.current_header is None

    def _handle_rate_line(self, rate, source: str) -> None:
        if self.current_header is None:
            _fail(source, "rate row appears before a control row")
        if rate[1]:
            self.current_type = {"CASH": "CASH_RATE", "PUA": "PUA_RATE", "1YT": "OYT_RATE"}[rate[1]]
        if self.current_type is None:
            _fail(source, "continuation rate row has no CASH/PUA/1YT type")
        start, end = int(rate[2]), int(rate[3])
        first, last = self.current_header["FIRST_DURATION"], self.current_header["LAST_DURATION"]
        if not first <= start <= last or not start <= end <= start + 9:
            _fail(source, "rate duration range is outside the header range or print grid")
        values = _rate_values(rate[4], source)
        if len(values) != min(end, last) - start + 1:
            _fail(source, "rate value count does not match the printed duration range")
        self._store_rate_values(start, values, source)

    def _store_rate_values(self, start: int, values: list[Decimal], source: str) -> None:
        target = self.buffers[self.current_type]
        for duration, value in enumerate(values, start):
            if duration in target and target[duration] != value:
                _fail(source, f"conflicting duplicate {self.current_type} duration {duration}")
            target[duration] = value

    def _start_control_row(self, text: str, source: str) -> None:
        self._flush()
        if self.description is None:
            _fail(source, "control row has no RECORD TYPE")
        self.current_header = _control_line(text, self.description, source)
        self.current_source = source
        self.pending_description = False

    def _flush(self) -> None:
        if self.current_header is None:
            return
        expected = set(range(self.current_header["FIRST_DURATION"], self.current_header["LAST_DURATION"] + 1))
        self._validate_complete_rates(expected)
        rows = self._rate_rows(expected)
        key = self.current_header["HEADER_ID"]
        previous = self.records.get(key)
        if previous and (previous.header != self.current_header or previous.rates != rows):
            _fail(self.current_source, f"conflicting duplicate HEADER_ID (first seen at {previous.source})")
        if previous is None:
            self.records[key] = _Record(self.current_header, rows, self.current_source)
        self.current_header = None
        self.current_type = None
        self.buffers = {column: {} for column in _BASE_RATE_COLUMNS}

    def _validate_complete_rates(self, expected: set[int]) -> None:
        for column, rates in self.buffers.items():
            if set(rates) != expected:
                missing = sorted(expected - set(rates))
                _fail(self.current_source, f"incomplete {column}; missing durations {missing[:10]}")

    def _rate_rows(self, expected: set[int]) -> list[dict[str, object]]:
        return [
            {
                "HEADER_ID": self.current_header["HEADER_ID"],
                "RECORD_TYPE": self.current_header["RECORD_TYPE"],
                "DURATION": duration,
                **{column: self.buffers[column][duration] for column in _BASE_RATE_COLUMNS},
                **dict.fromkeys(_PUA_RATE_COLUMNS),
            }
            for duration in sorted(expected)
        ]


def _is_dividend_formatting_line(text: str) -> bool:
    return (
        re.fullmatch(r"[- ]+", text)
        or text.startswith(("USR CLASS", "ID  BASE-SUB"))
        or re.match(r"0?TYPE\s+DURATION\b", text)
    )


def parse_dividend(path: str | Path) -> dict[str, list[dict[str, object]]]:
    """Read a complete CKAS/CKMO CKCVDVPD print extract without DB access.

    Returns ``WL_DIV_HEADER`` and ``WL_RATE_DIV`` rows using physical column
    names, Decimal rates and date objects (except documented zero-date strings).
    Identical duplicate schedules are collapsed; conflicting duplicates, missing
    rates/references and unrecognized non-formatting lines raise
    :class:`DividendValidationError` with a source line number.
    """
    return _DividendParser(Path(path)).parse()


def _plan_key_heading(path: Path, rows) -> None:
    heading = list(next(rows, ()))
    while heading and heading[-1].value is None:
        heading.pop()
    if tuple(str(cell.value or "").strip().upper() for cell in heading) != PLANKEY_COLUMNS:
        _fail(path.name, f"expected mapping columns {', '.join(PLANKEY_COLUMNS)}")


def _plan_key_row(cells, source: str) -> dict[str, object]:
    if any(cell.value is not None for cell in cells[len(PLANKEY_COLUMNS):]):
        _fail(source, "unexpected data beyond the six mapping columns")
    row = {}
    for column, limit, cell in zip(PLANKEY_COLUMNS, (8, 1, 1, 2, 2, 6), cells):
        if cell.data_type in {"f", "e"} or isinstance(cell.value, bool):
            _fail(source, f"{column} must be a literal identifier, not a formula/error/boolean")
        value = "" if cell.value is None else str(cell.value).strip()
        if column == "USER_KEY" and (not value or value.upper() == "(BLANK)"):
            row[column] = None
        else:
            row[column] = _text(value, column, limit, source)
    return row


def parse_plan_key_map(path: str | Path) -> list[dict[str, object]]:
    """Read the authoritative six-column WL_DIV_PLANKEY_MAP workbook.

    Only USER_KEY may be blank; its source workbook's explicit ``(blank)``
    marker is normalized to None. Formula cells are rejected, not cached.
    """
    from openpyxl import load_workbook

    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows()
        _plan_key_heading(path, rows)
        unique = {}
        for number, cells in enumerate(rows, 2):
            if all(cell.value is None for cell in cells):
                continue
            source = f"{path.name}:{sheet.title}:{number}"
            row = _plan_key_row(cells, source)
            key = tuple(row[column] for column in PLANKEY_COLUMNS[:3])
            if key in unique and unique[key] != row:
                _fail(source, "conflicting duplicate (PLANCODE, SEX, RATECLASS)")
            unique[key] = row
        if not unique:
            _fail(path.name, "mapping workbook has no data rows")
        return list(unique.values())
    finally:
        workbook.close()
