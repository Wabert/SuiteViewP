"""Workbook specification and writers for ABR quote detail exports.

The calculation viewer and output panel share the same workbook model.  The
spec places every value at an explicit row/column with a named style so both
writers render the identical formatted report:

* ``write_openpyxl`` saves an ``.xlsx`` for Print Detail (policy folder).
* ``write_excel_com`` opens an unsaved Excel workbook from the Calc Viewer.

Layout contract (matches the historical Print Detail workbook): a dark-red
section banner per group, bold labels in column A with values in column B,
the Derived Substandard Values and APV figures in columns C-D, and grey-header
monthly tables with per-column number formats and red summary rows.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

from ..models.abr_constants import MODAL_LABELS, PLAN_CODE_INFO

STYLE_TITLE = "title"
STYLE_SECTION = "section"
STYLE_LABEL = "label"
STYLE_VALUE = "value"
STYLE_SUBHEADER = "subheader"
STYLE_MESSAGE = "message"
STYLE_SUMMARY = "summary"


@dataclass(frozen=True)
class CellStyle:
    """Font/fill description shared by the openpyxl and COM writers."""

    bold: bool = False
    size: int = 11
    color: str | None = None
    fill: str | None = None
    underline: bool = False


CELL_STYLES: dict[str, CellStyle] = {
    STYLE_TITLE: CellStyle(bold=True, size=12),
    STYLE_SECTION: CellStyle(bold=True, size=11, color="FFFFFF", fill="8B0000"),
    STYLE_LABEL: CellStyle(bold=True, size=11),
    STYLE_VALUE: CellStyle(size=11),
    STYLE_SUBHEADER: CellStyle(bold=True, size=11, underline=True),
    STYLE_MESSAGE: CellStyle(bold=True, size=10, color="C62828"),
    STYLE_SUMMARY: CellStyle(bold=True, size=11, color="8B0000"),
}
TABLE_HEADER_STYLE = CellStyle(bold=True, size=10, fill="D3D3D3")
TABLE_DATA_STYLE = CellStyle(size=10)


@dataclass(frozen=True)
class CellSpec:
    """One positioned, styled cell (1-based row/column)."""

    row: int
    column: int
    value: Any
    style: str = STYLE_VALUE
    number_format: str | None = None


@dataclass(frozen=True)
class MergeSpec:
    """A single-row merged range."""

    row: int
    first_column: int
    last_column: int


@dataclass(frozen=True)
class TableSpec:
    """A header row plus data rows starting at A1.

    ``number_formats`` is per column; a format applies only to float cells so
    integer counters and blank placeholders keep Excel's General format.
    """

    headers: tuple[str, ...]
    rows: list[tuple[Any, ...]] = field(default_factory=list)
    number_formats: tuple[str | None, ...] = ()


@dataclass(frozen=True)
class SheetSpec:
    """A workbook sheet: an optional table plus free-positioned cells."""

    name: str
    cells: list[CellSpec] = field(default_factory=list)
    merges: list[MergeSpec] = field(default_factory=list)
    column_widths: dict[str, float] = field(default_factory=dict)
    table: TableSpec | None = None


@dataclass(frozen=True)
class WorkbookSpec:
    """Complete ABR detail workbook specification."""

    sheets: list[SheetSpec]


class _FormSheet:
    """Row cursor that lays out titled label/value forms."""

    def __init__(self, name: str):
        self.name = name
        self.row = 1
        self.cells: list[CellSpec] = []
        self.merges: list[MergeSpec] = []

    def put(self, row: int, column: int, value: Any, style: str = STYLE_VALUE) -> None:
        self.cells.append(CellSpec(row, column, value, style))

    def title(self, text: str) -> None:
        self.put(self.row, 1, text, STYLE_TITLE)
        self.row += 2

    def section(self, text: str) -> None:
        self.put(self.row, 1, text, STYLE_SECTION)
        self.merges.append(MergeSpec(self.row, 1, 2))
        self.row += 1

    def field(self, label: str, value: Any) -> None:
        self.put(self.row, 1, label, STYLE_LABEL)
        self.put(self.row, 2, value, STYLE_VALUE)
        self.row += 1

    def subheader(self, text: str) -> None:
        self.put(self.row, 1, text, STYLE_SUBHEADER)
        self.row += 1

    def message(self, text: str) -> None:
        self.put(self.row, 1, f"\u2022 {text}", STYLE_MESSAGE)
        self.merges.append(MergeSpec(self.row, 1, 2))
        self.row += 1

    def side_fields(self, start_row: int, pairs: list[tuple[str, Any]]) -> None:
        """Write label/value pairs into columns C-D beside an existing block."""
        for offset, (label, value) in enumerate(pairs):
            self.put(start_row + offset, 3, label, STYLE_LABEL)
            self.put(start_row + offset, 4, value, STYLE_VALUE)

    def build(self, column_widths: dict[str, float]) -> SheetSpec:
        return SheetSpec(
            name=self.name,
            cells=list(self.cells),
            merges=list(self.merges),
            column_widths=dict(column_widths),
        )


def _money(value: float | int | None, dash: str = "—") -> str:
    return f"${value:,.2f}" if value else dash


def _date(value) -> str:
    return value.strftime("%m/%d/%Y") if value else "—"


def _plan_description(policy) -> str:
    if not policy or not policy.plan_code:
        return "—"
    info = PLAN_CODE_INFO.get(policy.plan_code.upper())
    return f"{info[1]} ({info[0]}-Year Level)" if info else "—"


def _is_ul(policy) -> bool:
    return bool(policy) and policy.product_type in ("UL", "IUL", "ISWL")


def _policy_sheet(policy, result) -> SheetSpec:
    sheet = _FormSheet("Policy Info")
    sheet.title("ABR Quote — Policy Information")

    sheet.section("Policy Details")
    sheet.field("Policy Number:", policy.policy_number)
    sheet.field("Insured:", policy.insured_name or "—")
    sheet.field("Plancode:", policy.plan_code or "—")
    sheet.field("Plan Description:", _plan_description(policy))
    sheet.field("Sex:", {"M": "Male", "F": "Female", "U": "Unisex"}.get(policy.sex, policy.sex or "—"))
    sheet.field("Rate Sex:", policy.rate_sex or "—")
    sheet.field("Issue Age:", policy.issue_age)
    sheet.field("Attained Age:", policy.attained_age)
    sheet.field("Rate Class:", policy.rate_class or "—")
    sheet.field("Face Amount:", _money(policy.face_amount))
    sheet.field("Min Face:", f"${policy.min_face_amount:,.0f}")
    sheet.field("Issue State:", policy.issue_state or "—")
    sheet.field("Issue Date:", _date(policy.issue_date))
    sheet.field("Policy Year:", policy.policy_year)
    sheet.field("Month of Year:", policy.policy_month)
    sheet.field("Base Plancode:", policy.base_plancode or "—")
    sheet.field("Billing Mode:", MODAL_LABELS.get(policy.billing_mode, str(policy.billing_mode)))
    sheet.field("Modal Premium:", _money(policy.modal_premium))
    sheet.field("Table Rating:", policy.table_rating)
    sheet.field("Annual Flat Extra:", f"${policy.flat_extra:.2f}" if policy.flat_extra > 0 else "None")
    sheet.field("Flat Cease Date:", _date(policy.flat_cease_date))
    sheet.field("Reinsurers:", policy.reinsurers or "(none)")

    sheet.row += 1
    sheet.section("Quote Parameters")
    sheet.field("Quote Date:", _date(result.quote_date))
    sheet.field("ABR Interest Rate:", f"{result.abr_interest_rate * 100:.2f}%")
    sheet.field("Per Diem (Daily):", f"${result.per_diem_daily:,.2f}")
    sheet.field("Per Diem (Annual):", f"${result.per_diem_annual:,.2f}")

    sheet.row += 1
    sheet.section("Riders / Coverages")
    if policy.riders:
        for rider in policy.riders:
            desc = f"{rider.plancode} ({rider.rider_type})"
            if rider.benefit_type:
                desc += f" — BNF {rider.benefit_type}{rider.benefit_subtype or ''}"
            sheet.field(desc, f"${rider.fallback_premium:,.2f}/yr")
    else:
        sheet.field("No riders.", "")

    return sheet.build({"A": 22, "B": 35})


def _assessment_inputs(sheet: _FormSheet, assessment) -> None:
    a = assessment
    if not a:
        sheet.field("No assessment data.", "")
        return
    if a.use_five_year:
        sheet.field("5-Year Survival Rate:", f"{a.five_year_survival}")
        sheet.field("  Return to Normal:", "Yes" if a.use_return_5yr else "No")
    if a.use_ten_year:
        sheet.field("10-Year Survival Rate:", f"{a.ten_year_survival}")
        sheet.field("  Return to Normal:", "Yes" if a.use_return_10yr else "No")
    if a.use_le:
        sheet.field("Life Expectancy:", f"{a.life_expectancy_years} years")
    if a.use_increased_decrement:
        sheet.field("Increased Decrement:", f"{a.direct_increased_decrement:.0f}%")
        sheet.field("  Start/Stop Year:", f"{a.incr_decrement_start_year} — {a.incr_decrement_stop_year}")
    if a.use_table:
        sheet.field("Table (rating):", f"{a.direct_table_rating}")
        sheet.field("  Start/Stop Year:", f"{a.table_start_year} — {a.table_stop_year}")
    if a.use_flat:
        sheet.field("Flat ($/1000):", f"${a.direct_flat_extra:.2f}")
        sheet.field("  Start/Stop Year:", f"{a.flat_start_year} — {a.flat_stop_year}")
    if a.use_table_2:
        sheet.field("Table 2 (rating):", f"{a.direct_table_rating_2}")
        sheet.field("  Start/Stop Year:", f"{a.table_2_start_year} — {a.table_2_stop_year}")
    if a.use_flat_2:
        sheet.field("Flat 2 ($/1000):", f"${a.direct_flat_extra_2:.2f}")
        sheet.field("  Start/Stop Year:", f"{a.flat_2_start_year} — {a.flat_2_stop_year}")
    sheet.field("In Lieu Of:", "Yes" if a.in_lieu_of else "No (In Addition To)")


_DERIVED_PAIRS = (
    ("5-Year Survival:", "std_survival_5yr", "5-Year Survival:", "mod_survival_5yr"),
    ("10-Year Survival:", "std_survival_10yr", "10-Year Survival:", "mod_survival_10yr"),
    ("Life Expectancy:", "std_le", "Life Expectancy:", "mod_le"),
    ("Table Rating:", "std_table_rating", "Table Ratings:", "table_rating"),
    ("Flat Extra:", "std_flat_extra", "Flat Extras:", "flat_extra"),
)


def _derived_values(sheet: _FormSheet, assessment, derived_values: dict[str, str]) -> None:
    if derived_values:
        sheet.put(sheet.row, 1, "Current (Unmodified)", STYLE_SUBHEADER)
        sheet.put(sheet.row, 3, "Modified (Substandard Applied)", STYLE_SUBHEADER)
        sheet.row += 1
        for std_label, std_key, mod_label, mod_key in _DERIVED_PAIRS:
            sheet.put(sheet.row, 1, std_label, STYLE_LABEL)
            sheet.put(sheet.row, 2, derived_values.get(std_key, "—"), STYLE_VALUE)
            sheet.put(sheet.row, 3, mod_label, STYLE_LABEL)
            sheet.put(sheet.row, 4, derived_values.get(mod_key, "—"), STYLE_VALUE)
            sheet.row += 1
    elif assessment:
        sheet.field("Derived Table Rating:", f"{assessment.derived_table_rating:.4f}")
        if assessment.use_five_year and assessment.use_ten_year:
            sheet.field("  5yr Table Rating:", f"{assessment.derived_table_rating_5yr:.4f}")
            sheet.field("  10yr Table Rating:", f"{assessment.derived_table_rating_10yr:.4f}")
        sheet.field("Life Expectancy (rounded):", f"{assessment.life_expectancy_rounded}")


def _acceleration_block(sheet: _FormSheet, prefix: str, result) -> int:
    """Write one Eligible DB → Accelerated Benefit block; return its first row."""
    start_row = sheet.row
    sheet.field("Eligible Death Benefit:", f"${getattr(result, f'{prefix}_eligible_db'):,.2f}")
    sheet.field("Actuarial Discount:", f"${getattr(result, f'{prefix}_actuarial_discount'):,.2f}")
    sheet.field("Administrative Fee:", f"${getattr(result, f'{prefix}_admin_fee'):,.2f}")
    loan_repayment = getattr(result, f"{prefix}_loan_repayment")
    if loan_repayment > 0:
        sheet.field("Loan Repayment:", f"${loan_repayment:,.2f}")
    sheet.field("Calculated Benefit:", f"${max(0.0, getattr(result, f'{prefix}_accel_benefit')):,.2f}")
    sheet.field("Benefit Ratio:", f"{getattr(result, f'{prefix}_benefit_ratio') * 100:.2f}%")
    surrender_value = getattr(result, f"{prefix}_surrender_value")
    if surrender_value > 0:
        sheet.field("Surrender Value:", f"${surrender_value:,.2f}")
        sheet.field("Accelerated Benefit:", f"${getattr(result, f'{prefix}_accelerated_benefit'):,.2f}")
    return start_row


def _apv_pairs(result, ratio: float = 1.0) -> list[tuple[str, str]]:
    return [
        ("APV_FB:", f"${result.apv_fb * ratio:,.2f}"),
        ("APV_FP:", f"${result.apv_fp * ratio:,.2f}"),
        ("APV_FD:", f"${result.apv_fd * ratio:,.2f}"),
    ]


def _assessment_sheet(
    policy,
    assessment,
    result,
    derived_values: dict[str, str],
    accel_amount: float,
    min_face_amount: float,
    after_partial_override: str,
    messages: list[str],
) -> SheetSpec:
    sheet = _FormSheet("Assessment")
    sheet.title("ABR Quote — Assessment")

    sheet.section("Rider Configuration")
    sheet.field("Rider Type:", assessment.rider_type if assessment else "—")

    sheet.row += 1
    sheet.section("Assessment Inputs")
    _assessment_inputs(sheet, assessment)

    sheet.row += 1
    sheet.section("Derived Substandard Values")
    _derived_values(sheet, assessment, derived_values)

    sheet.row += 1
    sheet.section("Results Summary")
    sheet.subheader("FULL ACCELERATION")
    sheet.field("Acceleration Amount Input:", f"${accel_amount:,.2f}" if accel_amount else "—")
    full_start_row = _acceleration_block(sheet, "full", result)
    sheet.side_fields(full_start_row, _apv_pairs(result))

    sheet.row += 1
    if result.partial_eligible_db > 0:
        sheet.subheader("MAX PARTIAL ACCELERATION")
        sheet.field("Min Face Amount Input:", f"${min_face_amount:,.0f}")
        partial_start_row = _acceleration_block(sheet, "partial", result)
        # Partial APV values are the full APVs scaled to the partial eligible DB.
        ratio = result.partial_eligible_db / result.full_eligible_db if result.full_eligible_db > 0 else 0.0
        sheet.side_fields(partial_start_row, _apv_pairs(result, ratio))
    else:
        sheet.field("Partial Acceleration:", "NOT ALLOWED — At Minimum Face")

    sheet.row += 1
    is_ul = _is_ul(policy)
    sheet.field("Last Monthly Deduction:" if is_ul else "Premium Before:", result.premium_before)
    sheet.field("After (Full Accel):", f"${result.premium_after_full:,.2f}")
    if result.partial_eligible_db > 0:
        after_partial = (after_partial_override if is_ul else "") or result.premium_after_partial
        sheet.field("After (Partial):", after_partial)
    else:
        sheet.field("After (Partial):", "NOT ALLOWED")

    if messages:
        sheet.row += 1
        sheet.section("Messages")
        for message in messages:
            sheet.message(message)

    widths = {"A": 28, "B": 35, "C": 30, "D": 35} if derived_values else {"A": 28, "B": 35, "C": 20, "D": 25}
    return sheet.build(widths)


def _column_letter(column: int) -> str:
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _column_widths(column_count: int, width: float) -> dict[str, float]:
    return {_column_letter(index): width for index in range(1, column_count + 1)}


def _mortality_sheet(rows: list[dict]) -> SheetSpec:
    headers = (
        "Quote Month", "Policy Year", "Mo in Yr", "Att Age",
        "qx VBT (annual)", "qx × Mult (annual)", "qx Improved (annual)",
        "Table Rating", "qx + Table (annual)",
        "Flat Extra ($/1000)", "qx + Flat (annual)", "qx Capped (annual)",
        "qx Monthly", "px Monthly", "Cum Survival",
    )
    data = []
    for row in rows:
        table_rating = row.get("table_rating_applied", 0.0)
        flat_extra = row.get("flat_extra_applied", 0.0)
        data.append((
            row["quote_month"], row["duration_year"],
            row["month_in_year"], row["attained_age"],
            row["qx_vbt"], row["qx_multiplied"],
            row["qx_improved"],
            table_rating if table_rating > 0 else "",
            row["qx_table_rated"],
            flat_extra if flat_extra > 0 else "",
            row["qx_flat_extra"], row["qx_capped"],
            row["qx_monthly"], row["px_monthly"],
            row["cum_survival"],
        ))
    rate = "0.00000000"
    return SheetSpec(
        "Mortality Derivation",
        column_widths=_column_widths(len(headers), 14),
        table=TableSpec(headers, data, (None, None, None, None) + (rate,) * 11),
    )


def _life_expectancy_sheet(rows: list[dict]) -> SheetSpec:
    headers = (
        "Quote Month", "Policy Year", "Att Age",
        "qx Monthly", "px Monthly", "tPx (cum surv)",
        "Sum tPx (months)", "Curtate LE (years)",
    )
    tp_x = 1.0
    sum_tpx = 0.0
    data = []
    for row in rows:
        qx_monthly = row["qx_monthly"]
        px_monthly = 1.0 - qx_monthly
        tp_x *= px_monthly
        sum_tpx += tp_x
        data.append((
            row["quote_month"], row["duration_year"], row["attained_age"],
            qx_monthly, px_monthly, tp_x,
            sum_tpx, sum_tpx / 12.0,
        ))

    curtate_le = sum_tpx / 12.0 if rows else 0.0
    summary_row = len(data) + 3
    cells = []
    for offset, (label, value) in enumerate((
        ("Sum tPx (months):", sum_tpx),
        ("Curtate LE (years):", curtate_le),
        ("Complete LE (+ 0.5):", curtate_le + 0.5),
    )):
        cells.append(CellSpec(summary_row + offset, 6, label, STYLE_SUMMARY))
        cells.append(CellSpec(summary_row + offset, 8, value, STYLE_SUMMARY, "0.0000"))

    rate = "0.000000"
    return SheetSpec(
        "Life Expectancy",
        cells=cells,
        column_widths=_column_widths(len(headers), 16),
        table=TableSpec(headers, data, (None, None, None) + (rate,) * 5),
    )


def _apv_sheet(rows: list[dict], apv_summary: dict | None, apv_sheet_name: str) -> SheetSpec:
    headers = (
        "Month", "t", "qx Monthly", "px Monthly", "tpx (cum surv)",
        "v^(t+1) (benefit)", "v^t (premium)", "Death Benefit",
        "PVDB(t) (this mo)",
        "PVDB Cum", "Prem Rate (per $1K)", "PVFP(t) (this mo)",
        "PVFP Cum", "tpx End",
    )
    data = [
        (
            row["month"], row["t"],
            row["qx_monthly"], row["px_monthly"], row["tp_x"],
            row["v_benefit"], row["v_premium"],
            row.get("death_benefit", 0.0),
            row["pvdb_t"], row["pvdb_cum"],
            row["prem_rate"] if row["prem_rate"] > 0 else "",
            row["pvfp_t"], row["pvfp_cum"],
            row["tp_x_end"],
        )
        for row in rows
    ]

    cells = []
    if apv_summary:
        summary_row = len(data) + 3
        for offset, (label, key, number_format) in enumerate((
            ("PVFB (raw sum):", "pvfb_raw", "0.000000"),
            ("Cont Mort Adj:", "cont_mort_adj", "0.0000000000"),
            ("PVFB (adj × 1000):", "pvfb_adjusted", "#,##0.00"),
            ("PVFP:", "pvfp", "#,##0.00"),
            ("Actuarial Discount:", "actuarial_discount", "#,##0.00"),
        )):
            cells.append(CellSpec(summary_row + offset, 9, label, STYLE_SUMMARY))
            cells.append(CellSpec(
                summary_row + offset, 10, apv_summary.get(key, 0), STYLE_SUMMARY, number_format,
            ))

    rate = "0.000000"
    return SheetSpec(
        apv_sheet_name,
        cells=cells,
        column_widths=_column_widths(len(headers), 16),
        table=TableSpec(headers, data, (None, None) + (rate,) * 12),
    )


def build_detail_workbook_spec(
    policy,
    result,
    assessment,
    mortality_rows: list[dict],
    apv_rows: list[dict],
    apv_summary: dict | None = None,
    derived_values: dict[str, str] | None = None,
    accel_amount: float = 0.0,
    min_face_amount: float = 0.0,
    after_partial_override: str = "",
    apv_sheet_name: str = "APV - Present Value",
    messages: list[str] | None = None,
) -> WorkbookSpec:
    """Build the five-sheet ABR detail workbook model.

    ``after_partial_override`` is the user-entered UL monthly deduction and is
    only used for UL/IUL/ISWL policies.  ``messages`` defaults to
    ``result.messages``.
    """
    return WorkbookSpec([
        _policy_sheet(policy, result),
        _assessment_sheet(
            policy,
            assessment,
            result,
            derived_values or {},
            accel_amount,
            min_face_amount,
            after_partial_override,
            list(messages) if messages is not None else list(result.messages or []),
        ),
        _mortality_sheet(mortality_rows),
        _life_expectancy_sheet(mortality_rows),
        _apv_sheet(apv_rows, apv_summary, apv_sheet_name),
    ])


# ── openpyxl writer ──────────────────────────────────────────────────────


class _OpenpyxlStyles:
    """Cache of openpyxl Font/Fill objects per CellStyle."""

    def __init__(self):
        from openpyxl.styles import Alignment, Font, PatternFill

        self._font_cls = Font
        self._fill_cls = PatternFill
        self.header_alignment = Alignment(horizontal="center", wrap_text=True)
        self._fonts: dict[CellStyle, Any] = {}
        self._fills: dict[CellStyle, Any] = {}

    def font(self, style: CellStyle):
        if style not in self._fonts:
            self._fonts[style] = self._font_cls(
                bold=style.bold,
                size=style.size,
                color=style.color,
                underline="single" if style.underline else None,
            )
        return self._fonts[style]

    def fill(self, style: CellStyle):
        if style.fill is None:
            return None
        if style not in self._fills:
            self._fills[style] = self._fill_cls("solid", fgColor=style.fill)
        return self._fills[style]

    def apply(self, cell, style: CellStyle) -> None:
        cell.font = self.font(style)
        fill = self.fill(style)
        if fill is not None:
            cell.fill = fill


def _write_openpyxl_table(ws, table: TableSpec, styles: _OpenpyxlStyles) -> None:
    for column, header in enumerate(table.headers, 1):
        cell = ws.cell(row=1, column=column, value=header)
        styles.apply(cell, TABLE_HEADER_STYLE)
        cell.alignment = styles.header_alignment
    data_font = styles.font(TABLE_DATA_STYLE)
    for row_index, values in enumerate(table.rows, 2):
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row_index, column=column, value=value)
            cell.font = data_font
            number_format = table.number_formats[column - 1] if column <= len(table.number_formats) else None
            if number_format and isinstance(value, float):
                cell.number_format = number_format


def _write_openpyxl_sheet(ws, sheet: SheetSpec, styles: _OpenpyxlStyles) -> None:
    for column, width in sheet.column_widths.items():
        ws.column_dimensions[column].width = width
    if sheet.table:
        _write_openpyxl_table(ws, sheet.table, styles)
    for spec in sheet.cells:
        cell = ws.cell(row=spec.row, column=spec.column, value=spec.value)
        styles.apply(cell, CELL_STYLES[spec.style])
        if spec.number_format:
            cell.number_format = spec.number_format
    for merge in sheet.merges:
        ws.merge_cells(
            start_row=merge.row,
            start_column=merge.first_column,
            end_row=merge.row,
            end_column=merge.last_column,
        )


def write_openpyxl(spec: WorkbookSpec, path: str) -> str:
    """Write an ABR detail workbook to ``path`` with openpyxl."""
    import openpyxl

    workbook = openpyxl.Workbook()
    styles = _OpenpyxlStyles()
    for index, sheet in enumerate(spec.sheets):
        if index == 0:
            ws = workbook.active
            ws.title = sheet.name
        else:
            ws = workbook.create_sheet(sheet.name)
        _write_openpyxl_sheet(ws, sheet, styles)
    workbook.save(path)
    return path


# ── Excel COM writer ─────────────────────────────────────────────────────

_XL_CENTER = -4108


def _bgr(hex_rgb: str) -> int:
    """Convert ``RRGGBB`` to the BGR integer Excel COM expects."""
    red, green, blue = (int(hex_rgb[i:i + 2], 16) for i in (0, 2, 4))
    return red | (green << 8) | (blue << 16)


def _com_apply_style(target, style: CellStyle) -> None:
    target.Font.Bold = style.bold
    target.Font.Size = style.size
    if style.color:
        target.Font.Color = _bgr(style.color)
    if style.underline:
        target.Font.Underline = True
    if style.fill:
        target.Interior.Color = _bgr(style.fill)


def _com_write_table(ws, table: TableSpec) -> None:
    data = [table.headers, *table.rows]
    column_count = len(table.headers)
    ws.Range(ws.Cells(1, 1), ws.Cells(len(data), column_count)).Value = data
    header = ws.Range(ws.Cells(1, 1), ws.Cells(1, column_count))
    _com_apply_style(header, TABLE_HEADER_STYLE)
    header.HorizontalAlignment = _XL_CENTER
    header.WrapText = True
    if not table.rows:
        return
    body = ws.Range(ws.Cells(2, 1), ws.Cells(len(data), column_count))
    _com_apply_style(body, TABLE_DATA_STYLE)
    for column, number_format in enumerate(table.number_formats, 1):
        if number_format:
            letter = _column_letter(column)
            ws.Range(f"{letter}2:{letter}{len(data)}").NumberFormat = number_format


def _com_write_sheet(ws, sheet: SheetSpec) -> None:
    if sheet.table:
        _com_write_table(ws, sheet.table)
    for spec in sheet.cells:
        cell = ws.Cells(spec.row, spec.column)
        cell.Value = spec.value
        _com_apply_style(cell, CELL_STYLES[spec.style])
        if spec.number_format:
            cell.NumberFormat = spec.number_format
    for merge in sheet.merges:
        merged = ws.Range(ws.Cells(merge.row, merge.first_column), ws.Cells(merge.row, merge.last_column))
        merged.Merge()
    for column, width in sheet.column_widths.items():
        ws.Columns(f"{column}:{column}").ColumnWidth = width


def write_excel_com(spec: WorkbookSpec):
    """Open an unsaved Excel workbook and populate it from ``spec``."""
    from win32com.client import dynamic

    excel = dynamic.Dispatch("Excel.Application")
    excel.Visible = True
    excel.ScreenUpdating = False
    try:
        workbook = excel.Workbooks.Add()
        previous = None
        for index, sheet in enumerate(spec.sheets):
            if index == 0:
                ws = workbook.Worksheets(1)
            else:
                ws = workbook.Worksheets.Add(After=previous)
            ws.Name = sheet.name
            _com_write_sheet(ws, sheet)
            previous = ws
        workbook.Worksheets(1).Activate()
        workbook.Worksheets(1).Range("A1").Select()
    finally:
        excel.ScreenUpdating = True
    return workbook
