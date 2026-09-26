"""ABR Quote — "Explanation of Benefit" content builder.

This module is the single source of truth for the customer-facing explanation of
how an Accelerated Death Benefit (ABR) offer is calculated. The same structured
document is rendered two ways:

    * ``explanation_to_html(doc)`` — for the on-screen Resources panel.
    * ``explanation_to_docx(doc, out_path)`` — for the "Export to Word" button.

The explanation is written to comprehensively answer the kind of detailed,
itemized disclosure request a policyholder (or their attorney) may send, while
staying inside the boundary of what the company will disclose:

    * DISCLOSED — the calculation method, the published *industry base* mortality
      table, the discount/interest basis, and every dollar amount and deduction
      in the offer (the ABR disclosure statement contemplated by the NAIC
      Accelerated Benefits Model Regulation #620 and good-faith claims practice).
    * NOT DISCLOSED — the individualized *modified* mortality table, the specific
      multipliers/debits, and the medical-director underwriting methodology,
      which are proprietary and confidential actuarial/underwriting work product.
      The actuarial method itself is on file with the Department of Insurance.

Where a precise value from the current quote is available it is injected;
otherwise a ``[Bracketed Placeholder]`` is emitted so the document can be
completed by hand for an ad-hoc request.

NOTE FOR MAINTAINERS: This is customer-facing legal/compliance language. Wording
changes should be cleared with Legal/Compliance and Actuarial before release.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from typing import List, Optional, Union

from ..models.abr_data import ABRPolicyData, ABRQuoteResult, MedicalAssessment

# ── Facts used in the explanation, sourced from the ABR forms and memoranda ──
# See docs/ABRQuote/. Form numbers and filing references are intentionally kept
# out of the customer-facing wording (this is a plain-language explanation, not
# a contract citation), but are noted here for maintainers.

# Chronic and Critical riders: the underlying standard-risk table is a 75%
# multiple of the 2008 VBT, then modified with an individualized table rating
# and/or flat extra derived from the health evaluation.
BASE_TABLE_CHRONIC_CRITICAL = (
    "a 75% multiple of the 2008 Valuation Basic Table (VBT), a standard industry "
    "mortality table that reflects sex and smoking status (Male/Female, "
    "Smoker/Nonsmoker)"
)
# Terminal rider: a mortality table for terminally ill insureds.
BASE_TABLE_TERMINAL = (
    "a mortality table for terminally ill insureds, under which the chance of "
    "death is 500 per 1,000 (50%) in each policy year"
)
MORTALITY_SOURCE = "the 2008 VBT is published by the Society of Actuaries"
MORTALITY_IMPROVEMENT_TEXT = "1% per year, up to attained age 105"
INTEREST_RATE_RULE = (
    "based on the Moody's Corporate Bond Yield Average \u2014 the same published "
    "index we use to set the maximum policy loan interest rate on the policy. We "
    "apply that same rate here to bring the future death benefit into today's "
    "dollars"
)
ADMIN_CHARGE_DEFAULT = 250.0
ADMIN_CHARGE_MAX = 500.0
PLACEHOLDER = "[to be completed]"

# Company letterhead / contact block (from the reference claims letter).
LETTERHEAD = [
    "AMERICAN NATIONAL INSURANCE COMPANY",
    "Life Insurance & Annuity Claims Department   P.O. Box 10466, Springfield, MO 65808-0466",
    "TEL: (800) 615-7372    FAX: (281) 538-6757",
]
CONTACT_PHONE = "(800) 615-7372"

# Rider-type → human-readable label (no form numbers in customer-facing text).
_RIDER_LABELS = {
    "terminal": "terminal illness",
    "chronic": "chronic illness",
    "critical": "critical illness",
}


# ── Structured document model ───────────────────────────────────────────────
@dataclass
class Para:
    """A body paragraph."""
    text: str
    bold: bool = False
    italic: bool = False


@dataclass
class Bullets:
    """A bulleted list."""
    items: List[str] = field(default_factory=list)


@dataclass
class WorksheetRow:
    label: str
    value: str
    emphasis: bool = False


@dataclass
class Worksheet:
    """A two-column label/value calculation worksheet."""
    rows: List[WorksheetRow] = field(default_factory=list)


@dataclass
class Note:
    """A set-apart callout (used for confidentiality / boundary statements)."""
    text: str


Block = Union[Para, Bullets, Worksheet, Note]


@dataclass
class Section:
    heading: str
    blocks: List[Block] = field(default_factory=list)


@dataclass
class ExplanationDoc:
    title: str
    subtitle: str
    sections: List[Section] = field(default_factory=list)
    # Letter framing (claims-letter format). When these are populated the
    # document renders as a formal letter rather than a titled report.
    letterhead: List[str] = field(default_factory=list)
    letter_date: str = ""
    recipient: List[str] = field(default_factory=list)
    re_lines: List[str] = field(default_factory=list)
    salutation: str = ""
    closing: str = ""
    signer: str = ""
    signer_title: str = ""
    cc_lines: List[str] = field(default_factory=list)


@dataclass(frozen=True)
class _RiderContext:
    """Normalized rider inputs that control customer-facing wording."""

    label: str
    is_terminal: bool


@dataclass(frozen=True)
class _ProductContext:
    """Product-family flags used by premium explanation sections."""

    is_ul: bool
    is_term: bool


@dataclass(frozen=True)
class _CalculationAmounts:
    """Money and rate values displayed in the explanation worksheet."""

    eligible_db: Optional[float]
    admin_fee: float
    loan_repay: Optional[float]
    payable: Optional[float]
    surrender: Optional[float]
    pv_benefits: Optional[float]
    pv_premiums: Optional[float]
    pv_dividends: Optional[float]
    interest_rate: Optional[float]
    accelerated_amount: Optional[float]
    net_pv: Optional[float]
    calc_benefit: Optional[float]
    discount: Optional[float]
    accel_pct: Optional[float]


@dataclass(frozen=True)
class _AmountComponents:
    eligible_db: Optional[float]
    admin_fee: float
    loan_repay: Optional[float]
    payable: Optional[float]
    surrender: Optional[float]
    pv_benefits: Optional[float]
    pv_premiums: Optional[float]
    pv_dividends: Optional[float]
    interest_rate: Optional[float]


# ── Formatting helpers ──────────────────────────────────────────────────────
def _money(value: Optional[float]) -> str:
    if value is None:
        return PLACEHOLDER
    try:
        return f"${value:,.2f}"
    except (TypeError, ValueError):
        return PLACEHOLDER


def _pct(value: Optional[float], decimals: int = 2) -> str:
    if value is None:
        return PLACEHOLDER
    try:
        return f"{value * 100:.{decimals}f}%"
    except (TypeError, ValueError):
        return PLACEHOLDER


def _rate_pct(value: Optional[float], decimals: int = 3) -> str:
    """A rate already expressed as a fraction (e.g. 0.0545) → '5.450%'."""
    if not value:
        return PLACEHOLDER
    try:
        return f"{value * 100:.{decimals}f}%"
    except (TypeError, ValueError):
        return PLACEHOLDER


def _fmt_date(d: Optional[date]) -> str:
    return d.strftime("%B %d, %Y") if d else PLACEHOLDER


def _rider_context(assessment: Optional[MedicalAssessment]) -> _RiderContext:
    rider_key = ""
    if assessment and assessment.rider_type:
        rider_key = assessment.rider_type.strip().lower()
    is_terminal = rider_key.startswith("terminal")
    is_chronic = rider_key.startswith("chronic")
    is_critical = rider_key.startswith("critical")
    label = _RIDER_LABELS.get(
        "terminal" if is_terminal else
        "chronic" if is_chronic else
        "critical" if is_critical else "", "",
    )
    return _RiderContext(label=label, is_terminal=is_terminal)


def _product_context(policy: Optional[ABRPolicyData]) -> _ProductContext:
    product_type = policy.product_type.upper() if policy and policy.product_type else ""
    return _ProductContext(
        is_ul=product_type in ("UL", "IUL", "VUL", "ISWL"),
        is_term=product_type == "TERM",
    )


def _amount_components(result: Optional[ABRQuoteResult]) -> _AmountComponents:
    eligible_db = result.full_eligible_db if result else None
    admin_fee = (
        result.full_admin_fee
        if result and result.full_admin_fee
        else ADMIN_CHARGE_DEFAULT
    )
    loan_repay = result.full_loan_repayment if result else None
    payable = result.full_accelerated_benefit if result else None
    surrender = result.full_surrender_value if result else None
    pv_benefits = result.apv_fb if result else None
    pv_premiums = result.apv_fp if result else None
    pv_dividends = result.apv_fd if result else None
    interest_rate = result.abr_interest_rate if result else None
    return _AmountComponents(
        eligible_db=eligible_db,
        admin_fee=admin_fee,
        loan_repay=loan_repay,
        payable=payable,
        surrender=surrender,
        pv_benefits=pv_benefits,
        pv_premiums=pv_premiums,
        pv_dividends=pv_dividends,
        interest_rate=interest_rate,
    )


def _has_eligible_override(
    result: Optional[ABRQuoteResult],
    eligible_db_override: Optional[float],
    eligible_db: Optional[float],
) -> bool:
    return bool(
        result and eligible_db_override and eligible_db
        and abs(eligible_db_override - eligible_db) >= 0.01
    )


def _with_eligible_override(
    amounts: _AmountComponents,
    eligible_db_override: float,
) -> _AmountComponents:
    ratio = eligible_db_override / amounts.eligible_db
    return _AmountComponents(
        eligible_db=eligible_db_override,
        admin_fee=amounts.admin_fee,
        loan_repay=(
            round(amounts.loan_repay * ratio, 2)
            if amounts.loan_repay else amounts.loan_repay
        ),
        payable=amounts.payable,
        surrender=(
            round(amounts.surrender * ratio, 2)
            if amounts.surrender else amounts.surrender
        ),
        pv_benefits=(amounts.pv_benefits or 0.0) * ratio,
        pv_premiums=(amounts.pv_premiums or 0.0) * ratio,
        pv_dividends=(amounts.pv_dividends or 0.0) * ratio,
        interest_rate=amounts.interest_rate,
    )


def _net_and_calculated_benefit(
    result: Optional[ABRQuoteResult],
    amounts: _AmountComponents,
) -> tuple[Optional[float], Optional[float]]:
    if not result:
        return None, None
    net_pv = (
        (amounts.pv_benefits or 0.0)
        - (amounts.pv_premiums or 0.0)
        + (amounts.pv_dividends or 0.0)
    )
    calc_benefit = round(
        net_pv - amounts.admin_fee - (amounts.loan_repay or 0.0), 2
    )
    return net_pv, calc_benefit


def _payable_amount(
    amounts: _AmountComponents,
    overridden: bool,
    calc_benefit: Optional[float],
) -> Optional[float]:
    if overridden and calc_benefit is not None:
        return round(max(max(0.0, calc_benefit), amounts.surrender or 0.0), 2)
    return amounts.payable


def _actuarial_discount(
    result: Optional[ABRQuoteResult],
    overridden: bool,
    accelerated_amount: Optional[float],
    net_pv: Optional[float],
) -> Optional[float]:
    discount = None if overridden else (result.full_actuarial_discount if result else None)
    if discount is None and accelerated_amount is not None and net_pv is not None:
        return round(accelerated_amount - net_pv, 2)
    return discount


def _calculation_amounts(
    result: Optional[ABRQuoteResult],
    eligible_db_override: Optional[float],
) -> _CalculationAmounts:
    amounts = _amount_components(result)
    overridden = _has_eligible_override(result, eligible_db_override, amounts.eligible_db)
    if overridden:
        amounts = _with_eligible_override(amounts, eligible_db_override)

    accelerated_amount = amounts.eligible_db
    net_pv, calc_benefit = _net_and_calculated_benefit(result, amounts)
    payable = _payable_amount(amounts, overridden, calc_benefit)
    discount = _actuarial_discount(result, overridden, accelerated_amount, net_pv)
    accel_pct = accelerated_amount / amounts.eligible_db if accelerated_amount else None

    return _CalculationAmounts(
        eligible_db=amounts.eligible_db,
        admin_fee=amounts.admin_fee,
        loan_repay=amounts.loan_repay,
        payable=payable,
        surrender=amounts.surrender,
        pv_benefits=amounts.pv_benefits,
        pv_premiums=amounts.pv_premiums,
        pv_dividends=amounts.pv_dividends,
        interest_rate=amounts.interest_rate,
        accelerated_amount=accelerated_amount,
        net_pv=net_pv,
        calc_benefit=calc_benefit,
        discount=discount,
        accel_pct=accel_pct,
    )


def _letter_shell(
    policy_number: str,
    quote_date: str,
    insured_name: str,
    rider_label: str,
    has_result: bool,
) -> ExplanationDoc:
    return ExplanationDoc(
        title="How Your Accelerated Benefit Was Determined",
        subtitle=(
            f"Policy {policy_number}"
            + (f"   \u2022   Accelerated benefit for {rider_label}" if rider_label else "")
            + (f"   \u2022   Prepared {quote_date}" if has_result else "")
        ),
        letterhead=list(LETTERHEAD),
        letter_date=(quote_date if has_result else PLACEHOLDER),
        recipient=[
            insured_name if insured_name else "[Policyholder Name]",
            "[Street Address]",
            "[City, State  ZIP]",
        ],
        re_lines=[f"RE:  Accelerated Benefit    Policy: {policy_number}"],
        salutation=(
            f"Dear {insured_name}:"
            if insured_name
            else "Dear [Policyholder Name]:"
        ),
        closing="Sincerely,",
        signer="[Claims Specialist Name]",
        signer_title="Claims Specialist",
        cc_lines=["cc:  Agent \u2013 [Agent Name]"],
    )


def _opening_section() -> Section:
    return Section("", [
        Para(
            "This letter responds to your questions regarding the Accelerated "
            "Benefit available under your policy and explains how the benefit "
            "amount was determined."
        ),
        Para(
            "The calculation involves several actuarial components. The "
            "sections below describe each of them and show the figures used in "
            "this determination."
        ),
    ])


def _why_amount_section(is_terminal: bool) -> Section:
    blocks: List[Block] = [
        Para(
            "The Accelerated Death Benefit Rider is not a disability or "
            "income-replacement benefit. It is an early payment of the policy's "
            "death benefit: when the rider's eligibility conditions are met, a "
            "portion of the death benefit that would otherwise be paid in the "
            "future can be paid now, reduced by an actuarial discount that "
            "reflects how early it is being paid."
        ),
    ]
    if not is_terminal:
        blocks.append(Para(
            "Eligibility and amount are therefore determined by different "
            "things. Eligibility depends on the condition described in the "
            "rider. The amount, however, depends on how significantly the "
            "insured's condition affects expected mortality — the likelihood "
            "and timing of death — because that is what determines how early "
            "the death benefit is being paid. Some conditions profoundly "
            "affect daily functioning (morbidity) while having only a modest "
            "effect on expected mortality. In those situations the "
            "rider's conditions for making a claim are met, but the actuarial "
            "value of paying the death benefit early is small."
        ))
    blocks.append(Para(
        "The sooner the death benefit is expected to be paid, the less it "
        "must be reduced for being paid early, and the larger the amount that "
        "can be advanced today. Conversely, when expected mortality changes "
        "only modestly, the benefit is expected further in the future, a "
        "larger reduction applies, and the accelerated amount is smaller."
    ))
    return Section(
        "Why qualifying for the benefit does not itself determine the amount",
        blocks,
    )


def _mortality_basis_section(is_terminal: bool) -> Section:
    if is_terminal:
        blocks: List[Block] = [
            Para(
                "As part of the accelerated benefit calculation, we use a mortality "
                f"table built for terminally ill insureds \u2014 {BASE_TABLE_TERMINAL}. "
                "This table reflects a condition expected to significantly affect "
                "mortality, and it is what drives the present-value calculation "
                "described below."
            ),
        ]
    else:
        blocks = [
            Para(
                "As part of the accelerated benefit calculation, underwriting "
                "factors are applied to a standard industry mortality table to "
                "develop a modified mortality table specific to the insured. Based "
                "on the medical documentation we received, our Medical Director's "
                "Office reviewed the insured's condition and developed the modified "
                "mortality table used in this calculation."
            ),
            Para(
                "We begin with a standard industry mortality table \u2014 "
                f"{BASE_TABLE_CHRONIC_CRITICAL} ({MORTALITY_SOURCE}) \u2014 and "
                f"apply a mortality-improvement adjustment of "
                f"{MORTALITY_IMPROVEMENT_TEXT}. Our Medical Director's Office then "
                "adjusts that table to reflect the insured's specific condition. "
                "This modified mortality table is the primary driver of the "
                "calculation: it sets the year-by-year likelihood of the claim, "
                "which in turn determines the present value of the future death "
                "benefit and the future premiums."
            ),
        ]
    return Section("The mortality basis used in the calculation", blocks)


def _rider_provisions_section(admin_fee: float) -> Section:
    return Section(
        "The rider provisions that govern the calculation", [
        Para(
            "Under the terms of the rider, the Accelerated Death Benefit equals "
            "the portion of the Eligible Death Benefit requested, less the "
            "following deductions:"
        ),
        Bullets([
            "The actuarial discount (see details below);",
            f"An administrative charge (currently ${admin_fee:,.0f}, "
            f"not to exceed ${ADMIN_CHARGE_MAX:,.0f}); and,",
            "Any outstanding policy debt, if the qualifying insured is "
            "also the base policy insured.",
        ]),
        Para(
            "The Accelerated Death Benefit for the base policy insured will never "
            "be less than the cash surrender value of the base policy, if any."
        ),
        Para("A copy of the rider is available to you at any time on request."),
    ])


def _actuarial_discount_section(
    product: _ProductContext,
    level_annual_premium: Optional[float],
) -> Section:
    blocks: List[Block] = [
        Para(
            "The actuarial discount is the Eligible Death Benefit less the "
            "value today — the present value — of the benefit being "
            "accelerated:"
        ),
        Para("Actuarial Discount  =  Eligible Death Benefit  \u2212  Value Today",
             bold=True),
        Para(
            "The value today is built from three present-value components:"
        ),
        Bullets([
            "Present Value of Future Benefits (PVFB) \u2014 today's value of the "
            "death benefit that would otherwise be paid in the future.",
            "Present Value of Future Premiums (PVFP) \u2014 today's value of the "
            "future premiums needed to keep this policy in force to maturity.",
            "Present Value of Future Dividends (PVFDivs) \u2014 today's value of "
            "any future dividends the policy would have paid (zero for policies "
            "that do not pay dividends).",
        ]),
        Para("Value Today  =  PVFB  \u2212  PVFP  +  PVFDivs", bold=True),
        Para(
            "In plain terms: we start with today's value of the future death "
            "benefit, subtract today's value of the future premiums that would "
            "have been required to keep it in force, and add back today's value of "
            "any future dividends. From the value today we then subtract the "
            "administrative fee and any outstanding policy debt to arrive at "
            "the amount payable — subject to the cash-surrender-value minimum "
            "described above."
        ),
    ]
    if product.is_ul:
        lvl = level_annual_premium if level_annual_premium else None
        lvl_text = _money(lvl) if lvl else "[Annual Level Premium]"
        blocks.append(Para(
            "For this universal life policy, the future premiums are "
            "determined by calculating a level annual premium that keeps the "
            "policy in force to maturity, using the same interest rate applied "
            "elsewhere in this calculation. For this quote, that annual "
            f"premium was determined to be {lvl_text}."
        ))
    elif product.is_term:
        blocks.append(Para(
            "For this term policy, no separate premium calculation is needed: "
            "the future premiums used in this calculation are the premiums "
            "already scheduled under the policy's premium schedule for the "
            "remainder of the term."
        ))
    return Section("How the actuarial discount is calculated", blocks)


def _interest_rate_section(interest_rate: Optional[float]) -> Section:
    return Section("The interest rate", [
        Para(
            "Because a future benefit is being paid early, an interest rate is "
            "used to express future amounts in today's dollars. For this "
            f"calculation the rate is {_rate_pct(interest_rate)}, "
            f"{INTEREST_RATE_RULE}. Under the rider, this rate may not exceed "
            "the greater of the yield on 90-day Treasury Bills on the election "
            "date or the maximum adjustable policy loan interest rate allowed "
            "by law."
        ),
    ])


def _summary_section(amounts: _CalculationAmounts) -> Section:
    return Section("Summary of this calculation", [
        Para(
            "The figures behind this determination are summarized below. They "
            "show the key present-value components and how the accelerated "
            "benefit follows from them:"
        ),
        Worksheet([
            WorksheetRow("Eligible Death Benefit", _money(amounts.eligible_db)),
            WorksheetRow("Portion requested for acceleration",
                         _pct(amounts.accel_pct, 1)
                         if amounts.accel_pct is not None else PLACEHOLDER),
            WorksheetRow("Death benefit being accelerated",
                         _money(amounts.accelerated_amount)),
            WorksheetRow("Present Value of Future Benefits (PVFB)",
                         _money(amounts.pv_benefits)),
            WorksheetRow("Less: Present Value of Future Premiums (PVFP)",
                         _money(amounts.pv_premiums)),
            WorksheetRow("Plus: Present Value of Future Dividends (PVFDivs)",
                         _money(amounts.pv_dividends)),
            WorksheetRow("Value today of the accelerated benefit", _money(amounts.net_pv)),
            WorksheetRow("Actuarial discount (amount accelerated \u2212 value today)",
                         _money(amounts.discount)),
            WorksheetRow("Less: Administrative Fee", _money(amounts.admin_fee)),
            WorksheetRow("Less: outstanding policy loan, if any",
                         _money(amounts.loan_repay)),
            WorksheetRow("Calculated accelerated benefit", _money(amounts.calc_benefit)),
            WorksheetRow("Policy's current cash surrender value",
                         _money(amounts.surrender)),
            WorksheetRow("Accelerated benefit payable (the greater of the two above)",
                         _money(amounts.payable), emphasis=True),
        ]),
    ])


def _confidentiality_section() -> Section:
    return Section("A note about the information in this letter", [
        Para(
            "A few items \u2014 the individualized modified mortality table, the "
            "specific rating factors, and the internal medical evaluations "
            "behind them \u2014 are kept confidential in the claim file. We handle "
            "them this way to protect the insured's health information, which "
            "these documents reflect throughout, and because they are part of "
            "our internal underwriting and actuarial work. Please know this is "
            "not meant to withhold anything you need to understand the benefit: "
            "the summary above contains every figure used to arrive at the "
            "amount offered, and the actuarial methodology for this rider is on "
            "file with the insurance regulator as part of the approved form "
            "filing."
        ),
    ])


def _additional_medical_section() -> Section:
    return Section("Additional medical information", [
        Para(
            "Our goal is for the determination to reflect a complete picture of "
            "the insured's health. If there is additional medical information "
            "you would like us to consider, you are welcome to send it to us, "
            "and we will review it."
        ),
    ])


def _required_notices_section() -> Section:
    return Section("", [
        Note(
            "Important notices: Depending on your circumstances, an accelerated "
            "death benefit may be fully or partially excludable from income "
            "under federal tax law (Section 101(g) of the Internal Revenue "
            "Code). Because individual situations vary and the exclusions are "
            "subject to limits, we recommend reviewing the tax treatment with a "
            "personal tax advisor before accepting the benefit. Receipt of an "
            "accelerated benefit may also affect eligibility for means-tested "
            "government programs such as Medicaid or Supplemental Security "
            "Income (SSI); the agency that administers the program can confirm "
            "how a payment would be treated."
        ),
    ])


def _contact_section() -> Section:
    return Section("", [
        Para(
            "If you have questions about this determination or would like to "
            f"discuss it further, please contact our office at {CONTACT_PHONE}."
        ),
    ])


def _explanation_sections(
    rider: _RiderContext,
    product: _ProductContext,
    amounts: _CalculationAmounts,
    level_annual_premium: Optional[float],
) -> List[Section]:
    return [
        _opening_section(),
        _why_amount_section(rider.is_terminal),
        _mortality_basis_section(rider.is_terminal),
        _rider_provisions_section(amounts.admin_fee),
        _actuarial_discount_section(product, level_annual_premium),
        _interest_rate_section(amounts.interest_rate),
        _summary_section(amounts),
        _confidentiality_section(),
        _additional_medical_section(),
        _required_notices_section(),
        _contact_section(),
    ]


# ── Builder ─────────────────────────────────────────────────────────────────
def build_explanation(
    policy: Optional[ABRPolicyData] = None,
    result: Optional[ABRQuoteResult] = None,
    assessment: Optional[MedicalAssessment] = None,
    level_annual_premium: Optional[float] = None,
    eligible_db_override: Optional[float] = None,
) -> ExplanationDoc:
    """Build the Accelerated Benefit explanation as a formal claims letter.

    All arguments are optional. When a quote object is missing, its values are
    rendered as ``[to be completed]`` placeholders so the document remains a
    usable template for a manual response.

    ``level_annual_premium`` is the user-entered UL annual level premium (used to
    fund the coverage to maturity); when provided for a UL-type product, the
    letter explains how the future premiums were determined.

    ``eligible_db_override`` is the user-entered Eligible Death Benefit (e.g.
    Option B policies where the death benefit is face plus account value). When
    it differs from the quoted eligible DB, all dollar figures are rescaled
    proportionally — the same way the on-screen Full Acceleration recalc does.
    The letter always presents a full acceleration of that amount.
    """
    p = policy
    r = result
    policy_number = (p.policy_number if p and p.policy_number else PLACEHOLDER)
    quote_date = _fmt_date(r.quote_date if r else None)
    insured_name = (p.insured_name.strip() if p and p.insured_name else "")

    rider = _rider_context(assessment)
    product = _product_context(policy)
    amounts = _calculation_amounts(result, eligible_db_override)
    doc = _letter_shell(
        policy_number,
        quote_date,
        insured_name,
        rider.label,
        r is not None,
    )
    doc.sections.extend(
        _explanation_sections(rider, product, amounts, level_annual_premium)
    )
    return doc


def _html_letter_frame(doc: ExplanationDoc, crimson_dark: str, slate: str, gray: str) -> List[str]:
    parts: List[str] = []
    if doc.letterhead:
        parts.append(
            f'<div style="text-align:center;color:{crimson_dark};'
            f'border-bottom:1px solid {slate};padding-bottom:6px;margin-bottom:10px;">'
        )
        for i, line in enumerate(doc.letterhead):
            weight = "font-weight:bold;font-size:13px;" if i == 0 else "font-size:10px;"
            parts.append(f'<div style="{weight}">{line}</div>')
        parts.append("</div>")
    if doc.letter_date:
        parts.append(f'<p style="margin:8px 0;color:{gray};">{doc.letter_date}</p>')
    if doc.recipient:
        parts.append('<div style="margin:8px 0;line-height:1.3;color:%s;">' % gray)
        for line in doc.recipient:
            parts.append(f'<div>{line}</div>')
        parts.append("</div>")
    for line in doc.re_lines:
        parts.append(
            f'<p style="margin:8px 0;font-weight:bold;color:{crimson_dark};">{line}</p>'
        )
    if doc.salutation:
        parts.append(f'<p style="margin:8px 0;color:{gray};">{doc.salutation}</p>')
    return parts


def _html_report_title(doc: ExplanationDoc, crimson_dark: str, slate: str) -> List[str]:
    if doc.letterhead:
        return []
    return [
        f'<h2 style="color:{crimson_dark};margin:0 0 2px 0;">{doc.title}</h2>',
        f'<div style="color:{slate};font-size:11px;margin:0 0 12px 0;">'
        f'{doc.subtitle}</div>',
    ]


def _html_para(block: Para) -> str:
    style = "margin:4px 0;line-height:1.4;"
    if block.bold:
        style += "font-weight:bold;"
    if block.italic:
        style += "font-style:italic;color:#444;"
    return f'<p style="{style}">{block.text}</p>'


def _html_bullets(block: Bullets) -> str:
    parts = ['<ul style="margin:4px 0 4px 18px;line-height:1.4;">']
    parts.extend(f"<li>{item}</li>" for item in block.items)
    parts.append("</ul>")
    return "".join(parts)


def _html_worksheet(block: Worksheet, crimson: str, crimson_dark: str) -> str:
    parts = [
        '<table cellspacing="0" cellpadding="4" '
        'style="border-collapse:collapse;margin:6px 0;width:100%;">'
    ]
    for row in block.rows:
        lbl_style = "border-bottom:1px solid #ddd;"
        val_style = "border-bottom:1px solid #ddd;text-align:right;"
        if row.emphasis:
            lbl_style = (
                f"border-top:2px solid {crimson};font-weight:bold;"
                f"color:{crimson_dark};"
            )
            val_style = (
                f"border-top:2px solid {crimson};font-weight:bold;"
                f"text-align:right;color:{crimson_dark};"
            )
        parts.append(
            f'<tr><td style="{lbl_style}">{row.label}</td>'
            f'<td style="{val_style}">{row.value}</td></tr>'
        )
    parts.append("</table>")
    return "".join(parts)


def _html_note(block: Note, crimson: str) -> str:
    return (
        f'<div style="margin:8px 0;padding:8px 10px;'
        f'background:#F3E6E8;border-left:3px solid {crimson};'
        f'font-style:italic;color:#333;line-height:1.4;">'
        f'{block.text}</div>'
    )


def _html_block(block: Block, crimson: str, crimson_dark: str) -> str:
    if isinstance(block, Para):
        return _html_para(block)
    if isinstance(block, Bullets):
        return _html_bullets(block)
    if isinstance(block, Worksheet):
        return _html_worksheet(block, crimson, crimson_dark)
    if isinstance(block, Note):
        return _html_note(block, crimson)
    return ""


def _html_sections(doc: ExplanationDoc, crimson: str, crimson_dark: str) -> List[str]:
    parts: List[str] = []
    for sec in doc.sections:
        if sec.heading:
            parts.append(
                f'<h3 style="color:{crimson};margin:14px 0 4px 0;">{sec.heading}</h3>'
            )
        parts.extend(_html_block(block, crimson, crimson_dark) for block in sec.blocks)
    return parts


def _html_closing(doc: ExplanationDoc, gray: str) -> List[str]:
    parts: List[str] = []
    if doc.closing:
        parts.append(f'<p style="margin:14px 0 2px 0;color:{gray};">{doc.closing}</p>')
    if doc.signer:
        parts.append(
            f'<div style="margin:2px 0;color:{gray};">{doc.signer}</div>'
        )
    if doc.signer_title:
        parts.append(
            f'<div style="margin:0 0 8px 0;color:{gray};">{doc.signer_title}</div>'
        )
    for line in doc.cc_lines:
        parts.append(f'<div style="margin:2px 0;color:{gray};font-size:10px;">{line}</div>')
    return parts


# ── HTML rendering (on-screen panel) ────────────────────────────────────────
def explanation_to_html(doc: ExplanationDoc) -> str:
    """Render an ExplanationDoc to a styled HTML fragment for QTextEdit."""
    crimson = "#8B1A2A"
    crimson_dark = "#5C0A14"
    slate = "#4A6FA5"
    gray = "#333"
    parts: List[str] = []

    parts.extend(_html_letter_frame(doc, crimson_dark, slate, gray))
    parts.extend(_html_report_title(doc, crimson_dark, slate))
    parts.extend(_html_sections(doc, crimson, crimson_dark))
    parts.extend(_html_closing(doc, gray))
    return "".join(parts)


# ── Word rendering (export) ─────────────────────────────────────────────────
def _new_docx_document():
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    for section in document.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    return document


def _docx_colors():
    from docx.shared import RGBColor

    return {
        "crimson": RGBColor(0x8B, 0x1A, 0x2A),
        "crimson_dark": RGBColor(0x5C, 0x0A, 0x14),
        "slate": RGBColor(0x4A, 0x6F, 0xA5),
        "gray": RGBColor(0x33, 0x33, 0x33),
    }


def _docx_letter_header(document, doc: ExplanationDoc, colors) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    for i, line in enumerate(doc.letterhead):
        lp = document.add_paragraph()
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.paragraph_format.space_after = Pt(0)
        lrun = lp.add_run(line)
        lrun.bold = (i == 0)
        lrun.font.size = Pt(13 if i == 0 else 8.5)
        lrun.font.color.rgb = colors["crimson_dark"] if i == 0 else colors["gray"]
    document.add_paragraph()
    dp = document.add_paragraph()
    dp.add_run(doc.letter_date).font.color.rgb = colors["gray"]
    for line in doc.recipient:
        rp = document.add_paragraph()
        rp.paragraph_format.space_after = Pt(0)
        rp.add_run(line).font.color.rgb = colors["gray"]
    document.add_paragraph()
    for line in doc.re_lines:
        rep = document.add_paragraph()
        rerun = rep.add_run(line)
        rerun.bold = True
        rerun.font.color.rgb = colors["crimson_dark"]
    if doc.salutation:
        sp = document.add_paragraph()
        sp.add_run(doc.salutation).font.color.rgb = colors["gray"]


def _docx_report_header(document, doc: ExplanationDoc, colors) -> None:
    from docx.shared import Pt

    title_p = document.add_paragraph()
    run = title_p.add_run(doc.title)
    run.bold = True
    run.font.size = Pt(17)
    run.font.color.rgb = colors["crimson_dark"]
    sub_p = document.add_paragraph()
    sub_run = sub_p.add_run(doc.subtitle.replace("\u2022", "|"))
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = colors["slate"]


def _docx_para(document, block: Para) -> None:
    from docx.shared import RGBColor

    para = document.add_paragraph()
    run = para.add_run(block.text)
    run.bold = block.bold
    run.italic = block.italic
    if block.italic:
        run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def _docx_bullets(document, block: Bullets) -> None:
    for item in block.items:
        document.add_paragraph(item, style="List Bullet")


def _docx_worksheet(document, block: Worksheet, colors) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    table = document.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    table.autofit = True
    for row in block.rows:
        cells = table.add_row().cells
        lbl_run = cells[0].paragraphs[0].add_run(row.label)
        val_para = cells[1].paragraphs[0]
        val_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        val_run = val_para.add_run(row.value)
        if row.emphasis:
            lbl_run.bold = True
            val_run.bold = True
            lbl_run.font.color.rgb = colors["crimson_dark"]
            val_run.font.color.rgb = colors["crimson_dark"]


def _docx_note(document, block: Note) -> None:
    from docx.shared import Inches, RGBColor

    para = document.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.2)
    run = para.add_run(block.text)
    run.italic = True
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def _docx_block(document, block: Block, colors) -> None:
    if isinstance(block, Para):
        _docx_para(document, block)
    elif isinstance(block, Bullets):
        _docx_bullets(document, block)
    elif isinstance(block, Worksheet):
        _docx_worksheet(document, block, colors)
    elif isinstance(block, Note):
        _docx_note(document, block)


def _docx_sections(document, doc: ExplanationDoc, colors) -> None:
    from docx.shared import Pt

    for sec in doc.sections:
        if sec.heading:
            head_p = document.add_paragraph()
            head_p.space_before = Pt(10)
            head_run = head_p.add_run(sec.heading)
            head_run.bold = True
            head_run.font.size = Pt(12.5)
            head_run.font.color.rgb = colors["crimson"]
        for block in sec.blocks:
            _docx_block(document, block, colors)


def _docx_closing(document, doc: ExplanationDoc, colors) -> None:
    from docx.shared import Pt

    if doc.closing:
        document.add_paragraph()
        cp = document.add_paragraph()
        cp.paragraph_format.space_after = Pt(0)
        cp.add_run(doc.closing).font.color.rgb = colors["gray"]
    if doc.signer:
        document.add_paragraph()  # room for a signature
        sgn = document.add_paragraph()
        sgn.paragraph_format.space_after = Pt(0)
        sgn.add_run(doc.signer).font.color.rgb = colors["gray"]
    if doc.signer_title:
        tp = document.add_paragraph()
        tp.paragraph_format.space_after = Pt(0)
        tp.add_run(doc.signer_title).font.color.rgb = colors["gray"]
    if doc.cc_lines:
        document.add_paragraph()
        for line in doc.cc_lines:
            ccp = document.add_paragraph()
            ccp.paragraph_format.space_after = Pt(0)
            ccrun = ccp.add_run(line)
            ccrun.font.size = Pt(9)
            ccrun.font.color.rgb = colors["gray"]


def explanation_to_docx(doc: ExplanationDoc, out_path: str) -> str:
    """Render an ExplanationDoc to a professional .docx file at ``out_path``.

    Returns the path written. Raises ImportError if python-docx is unavailable.
    """
    document = _new_docx_document()
    colors = _docx_colors()

    if doc.letterhead:
        _docx_letter_header(document, doc, colors)
    else:
        _docx_report_header(document, doc, colors)
    _docx_sections(document, doc, colors)
    _docx_closing(document, doc, colors)

    document.save(out_path)
    return out_path
