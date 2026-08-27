"""Dump a worksheet's used range as a compact TSV-ish grid (values only).

Usage:
    venv\\Scripts\\python.exe tools/office/dump_sheet_grid.py '<json>'
    json: {"path": "<xlsx>", "sheet": "Sheet2", "max_rows": 60, "max_cols": 30}
"""
from __future__ import annotations

import json
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def main() -> None:
    arg = sys.argv[1]
    if arg.startswith("@"):
        with open(arg[1:], "r", encoding="utf-8") as fh:
            cfg = json.load(fh)
    else:
        cfg = json.loads(arg)
    path = cfg["path"]
    sheet = cfg["sheet"]
    max_rows = int(cfg.get("max_rows", 60))
    max_cols = int(cfg.get("max_cols", 30))

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    nrows = min(ws.max_row, max_rows)
    ncols = min(ws.max_column, max_cols)
    for r in range(1, nrows + 1):
        cells = []
        for c in range(1, ncols + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            cells.append(f"{get_column_letter(c)}{r}={v}")
        if cells:
            print(" | ".join(cells))
    wb.close()


if __name__ == "__main__":
    main()
