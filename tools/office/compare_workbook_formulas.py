"""Diff the *formulas* of two workbooks, sheet by sheet.

Usage (single JSON arg):
    venv\\Scripts\\python.exe tools/office/compare_workbook_formulas.py '<json>'

    {"a": "<workbook A>", "b": "<workbook B>",
     "sheets": ["CalcEngine", "NSP"],      # omit/null = every shared sheet
     "skip_sheets": ["Rates_Control"],     # bulk data sheets you don't want walked
     "max_row": 3000, "max_col": 760,      # optional caps per sheet
     "max_diffs": 25}

Companion to compare_workbook_ranges.py, which compares cached *values*; this
one compares the formula text so you can tell whether two copies of a
calculation workbook actually compute the same thing.

Prints JSON: per-sheet formula counts, differing-cell count, and the first
few differing cells (ref, formula in A, formula in B).
"""
from __future__ import annotations

import json
import sys

import openpyxl
from openpyxl.utils import get_column_letter


def _formulas(ws, max_row=None, max_col=None) -> dict:
    """Map 'A1' -> formula text for every formula cell in the sheet."""
    out = {}
    rows = ws.iter_rows(
        min_row=1,
        max_row=min(max_row, ws.max_row) if max_row else ws.max_row,
        max_col=min(max_col, ws.max_column) if max_col else ws.max_column,
    )
    for row in rows:
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                out[f"{get_column_letter(cell.column)}{cell.row}"] = v
    return out


def main() -> None:
    cmd = json.loads(sys.argv[1])
    max_diffs = int(cmd.get("max_diffs", 25))
    skip = {s.lower() for s in cmd.get("skip_sheets", [])}

    wb_a = openpyxl.load_workbook(cmd["a"], data_only=False, read_only=True, keep_links=False)
    wb_b = openpyxl.load_workbook(cmd["b"], data_only=False, read_only=True, keep_links=False)
    try:
        names = cmd.get("sheets") or [n for n in wb_a.sheetnames if n in set(wb_b.sheetnames)]
        report = []
        totals = {"formula_cells": 0, "differing": 0, "only_a": 0, "only_b": 0}

        for name in names:
            if name.lower() in skip:
                report.append({"sheet": name, "skipped": True})
                continue
            if name not in wb_a.sheetnames or name not in wb_b.sheetnames:
                report.append({"sheet": name, "error": "missing in one workbook"})
                continue

            fa = _formulas(wb_a[name], cmd.get("max_row"), cmd.get("max_col"))
            fb = _formulas(wb_b[name], cmd.get("max_row"), cmd.get("max_col"))

            only_a = sorted(set(fa) - set(fb))
            only_b = sorted(set(fb) - set(fa))
            differing = sorted(r for r in (set(fa) & set(fb)) if fa[r] != fb[r])

            totals["formula_cells"] += len(fa)
            totals["differing"] += len(differing)
            totals["only_a"] += len(only_a)
            totals["only_b"] += len(only_b)

            report.append({
                "sheet": name,
                "formulas_a": len(fa),
                "formulas_b": len(fb),
                "differing": len(differing),
                "only_in_a": len(only_a),
                "only_in_b": len(only_b),
                "sample_differing": [
                    {"ref": r, "a": fa[r][:200], "b": fb[r][:200]} for r in differing[:max_diffs]
                ],
                "sample_only_in_a": [{"ref": r, "a": fa[r][:200]} for r in only_a[:max_diffs]],
                "sample_only_in_b": [{"ref": r, "b": fb[r][:200]} for r in only_b[:max_diffs]],
            })
    finally:
        wb_a.close()
        wb_b.close()

    print(json.dumps({"a": cmd["a"], "b": cmd["b"], "totals": totals, "sheets": report},
                     indent=2, default=str))


if __name__ == "__main__":
    main()
