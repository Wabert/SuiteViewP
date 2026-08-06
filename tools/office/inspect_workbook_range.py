r"""Inspect formulas and cached values for an exact workbook range.

Usage:
    venv\Scripts\python.exe tools\inspect_workbook_range.py ^
      "{\"path\":\"docs\\Illustration_UL\\RERUN (v20.0) local IUL.xlsm\",\"sheet\":\"Illustration Values\",\"range\":\"BR9:CD45\"}"
"""

from __future__ import annotations

import json
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


def _serializable(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Pass a JSON config argument.")
    arg = sys.argv[1]
    if arg.startswith("@"):
        with open(arg[1:], "r", encoding="utf-8") as fh:
            config = json.load(fh)
    else:
        config = json.loads(arg)
    path = config["path"]
    sheet_name = config["sheet"]
    cell_range = config["range"]
    compact = bool(config.get("compact", False))

    formula_book = load_workbook(
        path, read_only=True, data_only=False, keep_vba=True)
    value_book = load_workbook(
        path, read_only=True, data_only=True, keep_vba=True)
    try:
        formula_sheet = formula_book[sheet_name]
        value_sheet = value_book[sheet_name]
        min_col, min_row, _max_col, _max_row = range_boundaries(cell_range)
        rows = []
        for row_offset, (formula_row, value_row) in enumerate(zip(
            formula_sheet[cell_range], value_sheet[cell_range]
        )):
            row = []
            for col_offset, (formula_cell, value_cell) in enumerate(zip(
                formula_row, value_row
            )):
                coordinate = (
                    f"{get_column_letter(min_col + col_offset)}"
                    f"{min_row + row_offset}"
                )
                row.append({
                    "cell": coordinate,
                    "formula": _serializable(formula_cell.value),
                    "value": _serializable(value_cell.value),
                    "number_format": formula_cell.number_format,
                })
            if compact:
                row = [
                    cell for cell in row
                    if cell["formula"] is not None or cell["value"] is not None
                ]
            rows.append(row)
    finally:
        formula_book.close()
        value_book.close()

    print(json.dumps({
        "path": path,
        "sheet": sheet_name,
        "range": cell_range,
        "rows": rows,
    }, indent=2))


if __name__ == "__main__":
    main()
