"""Read a keyed row from the historical CVF workbook to verify source signs.

Usage: venv\\Scripts\\python.exe tools\\rates\\inspect_cvf_workbook.py @config.json
Config: path, sheet, rate_key, issue_age, optional company and output.
No workbook or database data is modified.
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from suiteview.core.json_store import write_json


def main():
    arg = sys.argv[1]
    config = json.loads(Path(arg[1:]).read_text(encoding="utf-8")) if arg.startswith("@") else json.loads(arg)
    path = Path(config["path"])
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[config["sheet"]]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        if not all(name in header for name in ("Company", "CV Plancode", "Issue Age", 0)):
            raise ValueError("Unsupported CVF workbook header.")
        company_column = header.index("Company")
        key_column = header.index("CV Plancode")
        age_column = header.index("Issue Age")
        rate_start = header.index(0)
        matches = []
        for number, row in enumerate(rows, 2):
            key = str(row[key_column] or "").strip().upper().replace("-", "")
            if key != config["rate_key"].upper() or row[age_column] != config["issue_age"]:
                continue
            if config.get("company") and str(row[company_column]).strip().zfill(2) != config["company"]:
                continue
            matches.append({
                "row": number, "company": row[company_column],
                "rate_key": row[key_column], "issue_age": row[age_column],
                "rates": [
                    {"cell": f"{get_column_letter(column)}{number}", "duration": duration, "value": value}
                    for column, (duration, value) in enumerate(
                        zip(header[rate_start:], row[rate_start:]), rate_start + 1
                    )
                    if value is not None and value != ""
                ],
            })
        report = {"path": str(path), "sheet": sheet.title, "matches": matches}
        if not matches:
            raise ValueError("No matching CVF workbook row found.")
        if config.get("output"):
            write_json(config["output"], report)
        print(json.dumps(report, indent=2, default=str))
    finally:
        book.close()


if __name__ == "__main__":
    main()
