"""Build the final RERUN vs CyberLife baseline evidence workbook.

Usage (from the repository root)::

    venv\\Scripts\\python.exe tools\\rerun\\baseline_workbook.py
        --selection <selected_policies.json> --results <run_dir>
        --plancodes <baseline_plancodes.tsv>
        --charge-findings <charge_findings.json> --av-findings <av_root_causes.json>
        --shadow-findings <shadow_root_causes.json> [--shadow-extra <shadow_check_results.json>]
        [--rates-repo <Rates_Database folder>] [--rates-date 2026-09-30]
        [--rates-audit <rates_audit.json>]
        --output <workbook.xlsx> --summary-md <FINAL_SUMMARY.md> --summary-json <summary.json>

``<run_dir>`` is a ``baseline_history_compare.py`` output folder (``policy_results``,
``run_manifest.json``) that also holds ``av_decomposition.json`` written by
``baseline_av_decompose.py fit --work <run_dir>\\decomp``.

The builder

1. classifies every replayed policy-month and measure (AV, COI, expense, other,
   total charges, MD, interest credited, loan balance) as ``Exact`` (<= $0.01),
   ``Rounding`` (<= $1.00), ``Explained`` (> $1.00 and fully attributed to a
   validated root cause), ``Material`` (> $1.00, attributed to a cause whose proof
   is still open) or ``Unexplained`` (> $1.00, not attributed). The rollback seed
   row is ``Seed`` and is excluded from match counts.
2. attributes AV and interest differences with the per-month components of the
   AV decomposition (day count, receipt-date interest, net premium, withdrawal,
   MD, CyberLife-rule residual, interest on the carried difference, unexplained)
   and charge/loan differences with the validated charge and loan rules;
3. writes the evidence sheets (Summary, Fix List, By Plancode, Feature Coverage,
   Policies, Monthly Detail, Unexplained, root-cause sheets, Shadow, Month-0 MD
   Check, Rates Loaded, Limitations, Not Tested / Errors, Method & Sources) with
   openpyxl as a saved file, plus a one-page Markdown summary and a JSON summary.

``ROOT_CAUSES`` records the root causes validated by the 9/30/2026 investigations
(charge triage, AV deep dive, shadow deep dive). Update it when fixes land.
Everything here is read-only: CyberLife and UL_Rates are not touched by this
script; ``--rates-repo`` reads the loader's progress/issue logs and
``ProductDB_Manager.xlsx`` (read-only).
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from tools.rerun.baseline_common import (  # noqa: E402
    EXACT_TOLERANCE,
    ROUNDING_TOLERANCE,
    json_dump,
    json_load,
    parse_date,
    resolve_requested_plancodes,
)


MEASURES = ("av", "coi", "expense", "other", "total_charges", "md", "interest", "loan")
MEASURE_LABELS = {
    "av": "AV", "coi": "COI", "expense": "Expense", "other": "Other", "total_charges": "Total charges",
    "md": "MD", "interest": "Interest", "loan": "Loan",
}
CHARGE_MEASURES = ("coi", "expense", "other", "total_charges", "md")
CLASS_ORDER = ("Exact", "Rounding", "Explained", "Material", "Unexplained")
COMPARED_CLASSES = set(CLASS_ORDER)
UNEXPLAINED_IDS = {"U-AV", "U-INT", "U-CHG", "U-LOAN", "U-WD", "U-DATE", "U-SHD"}
MINOR = 0.50

# Plancodes whose plancode_table Interest_Method contradicts PDF DIFFCMPD (AV C1b).
C1B_PLANS = {
    "1S133229", "1S133729", "1S133A29", "1S133B2X", "1S133C29", "1S133D2X", "1S133E29", "1S133F29",
    "1S133G29", "1S133H29", "1S135A00", "1S135C00", "1S135M0X", "1S135N00", "1S135W29", "1S135X2X",
}
# Plancodes whose plancode_table LoanChargeRate contradicts PDF DBSPLINT (AV C4b).
P02_PLANS = {
    "1G130A00", "1S133H29", "1S134400", "NU1F3H00", "NU1F3L00", "NU1F3M00", "NU1F3N00", "NU1F3O00",
    "NU1F3P00", "NU1FU100", "NU1FU200", "NU1L1N00", "NU1L2B00", "NU1L2C00", "NU1LAC00", "NU1LAE00",
    "NU1LAJ00", "NU1LAL00", "NU1LAN00",
}
R01_POLICIES = {"01_U0703756"}
R02_POLICIES = {"01_U0482280", "01_U0482386"}

ROOT_CAUSES: list[dict[str, str]] = [
    {
        "id": "E01", "area": "Engine", "proof": "validated", "status": "Needs Robert's ruling (Albert F01)",
        "title": "Interest day-count shift in CyberLife-timing replays",
        "description": "ProjectionTiming.CYBERLIFE_MONTHLIVERSARY credits the span previous MV -> current MV with "
                       "the day count of the current calendar month (the next span). Each 30/31-day boundary "
                       "shifts one day of interest; the error alternates in sign and largely self-cancels.",
        "evidence": "CyberLife rule (CyberDoc B11 pp.32-37) reproduces TOT_CRE_ITS_AMT to 2 cents in 1,561/1,630 "
                    "months; months without cash flow: prior-span daily 514/520 vs engine 92/520. Runtime patch "
                    "01/13497726 max |AV diff| 2.72 -> 0.03. Albert F01 agrees (26/000324178 -24.47 both).",
        "source": "AV deep: av_root_causes.md C1, fix #1",
        "fix": "calc_engine.credit_interest_pre_withdrawal: pass interest_days(previous MV, month_date); keep "
               "(1+i)^(1/12)-1 for monthly plans; golden test 01/13497726 (81.95, 79.63, 82.62, 80.29, 83.31, 83.66).",
    },
    {
        "id": "E02", "area": "Engine", "proof": "validated", "status": "Ready to fix",
        "title": "CVAT necessary-premium (NPT) allowance is 0 after TAMRA year 7",
        "description": "calc_engine.compute_allowances passes npt_premium=0.0; premium_allowance sets the CVAT NPT "
                       "allowance to it once tamra_year > 7, so every premium is rejected (also in normal RERUN "
                       "illustrations, conform_to_tamra defaults True). CyberLife accepted the premiums.",
        "evidence": "11 CVAT policies (1U145700/1U146200/1U146300) lost every premium after TAMRA year 7; runtime "
                    "patch (NPT = INF + E01) cuts max |AV diff| from 1,133.34 to <= 1.17 on all 11.",
        "source": "AV deep: C3, fix #2",
        "fix": "Implement RERUN vNPT_Premium (NPT_NSP/NPT_Premium) or return INF with a loud flag until done; "
               "add a CVAT tamra_year 8 regression test.",
    },
    {
        "id": "E03", "area": "Engine", "proof": "validated", "status": "Ready to fix",
        "title": "Loan repayment allocation (interest-first vs CyberLife principal)",
        "description": "loan_handler.repay_loan pays arrears interest first; CyberLife applies PL repayments to "
                       "principal while POL_LN_ITS_AMT keeps accruing. All repayment codes map to one kind.",
        "evidence": "26/000289723: LN_PRI_AMT falls 149.01 per month in CyberLife; engine reduces accrued instead.",
        "source": "AV deep: C4b, fix #6",
        "fix": "Support principal-only repayments (CyberLife PL) and map FH_FIXED repayment codes explicitly.",
    },
    {
        "id": "E04", "area": "Engine", "proof": "validated", "status": "Ready to fix",
        "title": "Future-dated coverage segment charged before its issue date",
        "description": "A coverage segment with issue date after the projection month is included in DB/NAR/COI "
                       "and in MD-based waiver charges.",
        "evidence": "26/000324822 (1U1F4M00): phase 12 issue 2026-10-26 face 25,000 charged every month "
                    "(COI +28.66, waiver +8.92, MD +37.58); month-0 md_check variance 37.59.",
        "source": "Charge triage: charge_findings.md (engine issues)",
        "fix": "Filter segments with issue_date > month date out of DB/NAR/COI and dependent waiver bases.",
    },
    {
        "id": "E05", "area": "Engine", "proof": "validated", "status": "Ready to fix (latent; confirm 0.65% date with product)",
        "title": "Interest bonus timing, qualification and tiers (VP/MS VPM020)",
        "description": "Engine adds the duration bonus at the anniversary (one month early), ignores PRO_BNS_RS_CD "
                       "and conditional qualification, and loads only the first tier; tRates_IntBonus dates "
                       "1U135D00/K00/P00 0.65% from 2018-01-01 (CyberLife 2020-06-01).",
        "evidence": "VPM020 UL96bonus.PMS f_adjusted_interest_with_bonus; INT_RT history 01/U0482041, U0482280. "
                    "No sampled policy crosses a bonus anniversary in the window: no $ impact here.",
        "source": "AV deep: C7, fix #12",
        "fix": "Bonus from the MV after the threshold anniversary; honour PRO_BNS_RS_CD; latest tier replaces; "
               "move the 0.65% entry to 2020-06-01.",
    },
    {
        "id": "E06", "area": "Engine", "proof": "validated", "status": "Ready to fix (fail loud)",
        "title": "Missing chargeable benefit rate silently charged $0",
        "description": "rate_loader._load_benefit_coi_rates returns [] and _safe_rate([]) returns 0.0, so an active "
                       "benefit without a schema BENCOI is charged $0 with no warning.",
        "evidence": "Code path confirmed; Albert F05 (1U1F4L00 3F under dbo). With schema rates every chargeable "
                    "benefit in this sample loads; empty BENCOI only for non-charging benefits (A0, V1-V3, U0/U1).",
        "source": "AV deep: fix #13; Albert F05 (corroborating)",
        "fix": "Separate no-charge benefits from missing rates; raise RateLookupError for a chargeable benefit "
               "without a schema schedule.",
    },
    {
        "id": "E07", "area": "Engine (shadow)", "proof": "validated", "status": "Ready to fix",
        "title": "Shadow interest day count (SGUL monthly-effective vs days/365)",
        "description": "shadow_calc uses (1+i)^(days/365) with the base plan's day count; CyberLife SGUL shadow "
                       "credits (1+i)^(1/12). FPUL2 fits days/365, so the basis is product-specific.",
        "evidence": "UN013474 +3.44 -> +0.01; U0703370 +6.10 -> -0.01; 12 SGUL18+ policies end within $0.06. "
                    "Albert decomposition (verified closure): total effect +34.86 on 34 SGUL policies.",
        "source": "Shadow deep: 4.1; Albert SGUL suiteview_diff (corroborating)",
        "fix": "Shadow-specific interest method in plancode_table (SGUL = monthly, FPUL2 = actual days).",
    },
    {
        "id": "E08", "area": "Engine + Rates (shadow)", "proof": "validated", "status": "Ready to fix",
        "title": "LTGUL/LTGUL08 shadow target read from MTP scale S (0) and APS205 load relief missing",
        "description": "ShadowTarget='Table' reads MTP scale S = 0.00, so the 45% excess load hits every premium "
                       "dollar; CyberLife (APS205 F_CCVtargetpremium) uses the CCV target = UL_Rates CTP scale S "
                       "and relieves the load (2T per year / (n+1)T cumulative).",
        "evidence": "VP/MS APS205 replay exact to the cent on U0588503, U0590177 (the 9.58M case), U0590680, "
                    "U0591700, U0912979, U0915039; U0591866 +1.19 after withdrawals (E09).",
        "source": "Shadow deep: 4.3; aps205_results.json",
        "fix": "Read the shadow target from CTP scale S for LTGUL/LTGUL08 and implement the APS205 relief rule.",
    },
    {
        "id": "E09", "area": "Engine (shadow)", "proof": "validated", "status": "Ready to fix",
        "title": "Shadow ignores partial withdrawals; option-B shadow NAR adds AV",
        "description": "shadow_calc sets shadow_wd_charges = 0; APS205 subtracts each gross withdrawal with interest "
                       "and reduces units. APS205 NAR = face / i^(1/12) - AV for every DB option; engine adds AV "
                       "for option B.",
        "evidence": "U0591866 +150.56 -> +1.19 (withdrawals); U0577137 -46.21 -> +22.78 (DB option; residual = PW "
                    "charge basis).",
        "source": "Shadow deep: 4.6",
        "fix": "Apply withdrawals in the shadow; use face-based NAR for every DB option (APS205 products, FPUL2).",
    },
    {
        "id": "E10", "area": "Engine (shadow)", "proof": "validated", "status": "Ready to fix (fail loud)",
        "title": "Missing shadow rates read as 0",
        "description": "rate_loader._safe_rate returns 0.0 for an empty list and _load_shadow_rates uses `or []` for "
                       "EPU, TPP, EPP, SHADOW_INT and DBD; a CCV with no shadow rates runs on all-zero rates.",
        "evidence": "SGUL15 DBD = 0 (R04); 1U146200 shadow = premium sum (UNS00924 15,058.00 vs XP 10,273.58) (R05).",
        "source": "Shadow deep: 4.5 (4.2, 4.4)",
        "fix": "Raise RateLookupError when a 'Table' shadow rate is missing or a CCV is active without shadow rates.",
    },
    {
        "id": "E11", "area": "Engine capability / harness", "proof": "validated",
        "status": "Needs Robert's ruling (replay-only)",
        "title": "Dated cash flows bucketed to the next monthliversary without receipt-date interest",
        "description": "Dated inputs (premiums, withdrawals, loans, repayments) are applied at the next "
                       "monthliversary with no interest from the receipt date; CyberLife credits from receipt "
                       "(receipt and MV day both counted). Affects rollback/from-issue replays only; inforce "
                       "projections pay modal premiums on monthliversaries.",
        "evidence": "Premium months fitting only one convention: receipt+MV inclusive 343 vs exclusive 38; e.g. "
                    "01/U0590177 -363.55; FPUL2 shadow U0436543 -164.46 -> +0.08 with receipt-date interest.",
        "source": "AV deep: fix #9 (F02 Albert agrees); Shadow deep: 5",
        "fix": "Add exact-date crediting for dated inputs (n = MV - receipt + 1 days) or keep marking as timing.",
    },
    {
        "id": "E12", "area": "Engine (shadow)", "proof": "validated",
        "status": "Confirmed by Robert. Priority LOW for illustrations (premiums assumed on the monthliversary); "
                  "REQUIRED for rollback, from-issue and dated mid-month premium replays",
        "title": "SGUL shadow late-payment-forgiveness feature (confirmed by Robert)",
        "description": "SGUL product line: any premium paid during the month earns a full month of shadow interest "
                       "from the previous monthliversary (Robert: MV 8/16, $100 paid 8/25 earns interest 8/16 -> 9/16, "
                       "31 days); it is added after that month's deduction (does not reduce that month's NAR) and the "
                       "net premium is rounded to cents. shadow_calc moves the premium to the NEXT monthliversary, "
                       "applies it before the COI and gives it no interest.",
        "evidence": "Robert 9/30 (authoritative). Largest SGUL shadow defect: total effect -879.92 over 34 runnable "
                    "in-scope SGUL policies (Albert decomposition; engine values equal ours 48/48, closure $0.00, totals "
                    "re-summed here); e.g. U0697261 -186.03, U0703034 -145.60, U0653421 -126.27. Shadow deep 5: "
                    "UN004363 -57.47 -> +2.84. Load rounding +2.09.",
        "source": "Robert 9/30; Shadow deep 0b, 5; Albert SGUL suiteview_diff (corroborating, verified)",
        "fix": "Implement late-payment forgiveness in shadow_calc for SGUL (premium post-deduction at the MV on/before "
               "receipt, full month of interest) and cent rounding of the shadow net premium.",
    },
    {
        "id": "E13", "area": "Engine", "proof": "attributed", "status": "Needs Robert's ruling (external evidence)",
        "title": "4M target-based PWoT multiplied by table factor (1 + 0.25 x table)",
        "description": "monthly_deduction.target_waiver_charge multiplies by the table factor; CyberLife does not.",
        "evidence": "Albert F06 2/2 (26/000272626, 26/000299857); code path confirmed; no table-rated FFL policy "
                    "with an active 4M in this sample, so not reproduced here.",
        "source": "AV deep: Albert cross-reference F06",
        "fix": "Confirm with product; drop the table factor from the 4M target-based waiver if confirmed.",
    },
    {
        "id": "E14", "area": "Engine", "proof": "validated", "status": "Ready to fix (new 9/30)",
        "title": "Monthliversary dates drift for MV day 31 (projection steps from the prior date)",
        "description": "Projecting from a day-31 valuation, engine month dates step from the previous date "
                       "(3-31, 4-30, 5-30, ... 8-30) instead of from the issue date (CyberLife 5-31, 7-31, 8-31). "
                       "Dated inputs and state dates then fall on different days from CyberLife.",
        "evidence": "01/UE219641 (issue 2023-07-31): RERUN states 2026-05-30/07-30/08-30 vs CyberLife MV "
                    "2026-05-31/07-31/08-31; those 3 months cannot be compared. Only day-31 policy in the sample.",
        "source": "This baseline (engine states in decomp/engine/01_UE219641.json)",
        "fix": "Derive each projected monthliversary from the issue date (clamped to month end), not from the "
               "previous month's date; add a day-31 regression test.",
    },
    {
        "id": "P01", "area": "Plan config", "proof": "validated", "status": "Ready to fix",
        "title": "Interest_Method contradicts CyberLife PDF DIFFCMPD (16 1S13 plans)",
        "description": "15 1S13xxxx plans coded '1/12th' credit daily in CyberLife; 1S135A00 coded 'ExactDays' "
                       "credits monthly.",
        "evidence": "01/S1378831 (1S135A00) matches (1+i)^(1/12)-1 only; 01/S0509476 matches daily only.",
        "source": "AV deep: C1b, fix #4",
        "fix": "Set Interest_Method from PDF DIFFCMPD; add a plancode_table vs CYBERLIFE_PDF consistency test.",
    },
    {
        "id": "P02", "area": "Plan config", "proof": "validated", "status": "Ready to fix",
        "title": "LoanChargeRate contradicts PDF DBSPLINT (19 plancodes)",
        "description": "NU1F3H/L/M/N/O/P00 and NU1FU100/200 coded 8% while CyberLife charges 6.5%; NU1L*, "
                       "1G130A00, 1S133H29, 1S134400 to be checked.",
        "evidence": "26/000289723: CyberLife accrual 224.49 vs engine 276.92 for 31 days; live LN_CRG_ITS_RT 6.500.",
        "source": "AV deep: C4b, fix #5",
        "fix": "Correct LoanChargeRate from PDF DBSPLINT / live LN_CRG_ITS_RT.",
    },
    {
        "id": "P03", "area": "Plan config", "proof": "validated", "status": "Ready to fix",
        "title": "Withdrawal minimum face = hard-coded 25,000; WithdrawalFee not populated",
        "description": "plancode_table has no MinFaceAfterWD/WithdrawalFee; the default 25,000 caps "
                       "withdrawals under DB option A (1U130N2X PDF minimum 10,000).",
        "evidence": "01/U0107043: engine applied 0 of a 2,400 net (2,425 gross) withdrawal; 2,463.00 -> 2.43 with "
                    "min face 10,000.",
        "source": "AV deep: C5, fix #3",
        "fix": "Populate MinFaceAfterWD from PDF DBSMIAMT and WithdrawalFee per plan.",
    },
    {
        "id": "R01", "area": "Rates", "proof": "validated", "status": "Data (reload CIRF)",
        "title": "CIRF: 'ANICO2019 R' key not loaded; CIRF after 2026-05-01 missing",
        "description": "UL_Rates rates.RATE_ASSIGN_FUND has only ANICO2019; CyberLife credited the R key 2.55% / "
                       "2.30%. The source CIRF report is dated 05/01/26.",
        "evidence": "01/U0703756 over-credited: CyberLife-rule residual +2.46, final AV +2.20.",
        "source": "AV deep: C2, fix #10",
        "fix": "Reload CIRF from a current report; load '<key> R' variants once the routing attribute is known.",
    },
    {
        "id": "R02", "area": "Rates / engine", "proof": "validated", "status": "Data (identify DB2 source)",
        "title": "Policy-level guaranteed crediting floor (1U135K00 state 42 = 3.25%)",
        "description": "declared_rates floors at plancode GINT 3.00%; CyberLife floors at the policy's operative "
                       "guaranteed rate (3.25% for state 42).",
        "evidence": "01/U0482280, 01/U0482386: GUA_RT_ERN_ITS_AMT implies 3.25%; INT_RT history = max(CIRF, 3.25) "
                    "+ bonus at all six changes since 2018; impact -2.33 / -3.17 over 6 months.",
        "source": "AV deep: C2, fix #11",
        "fix": "Floor at the policy's guaranteed rate once its DB2 field is identified.",
    },
    {
        "id": "R03", "area": "Rates", "proof": "validated", "status": "Rates pending IAF pull (Robert) - 1U539600, ~3 in force",
        "title": "Rates absent from schema rates (loud RateLookupError)",
        "description": "Round 3 blocked 17 policies: TBL1MTP (08228000, 1U133100, 1U135000, MLUL, 1U144100, "
                       "1U144500), 1U135400 shadow scale S, riders 1U5CH800 and 1U539600. All loaded 9/30 except "
                       "rider 1U539600 (PDF only, no IAF/dbo source); Robert will pull its IAF.",
        "evidence": "RateLookupError list per policy (Not Tested / Errors sheet); bulk-loader-progress 9/30 21:02/21:24.",
        "source": "AV deep: fix #14; rates loader 9/30",
        "fix": "Robert pulls the 1U539600 IAF; load and re-run 04/UXE00064 and 04/UXE00073.",
    },
    {
        "id": "R04", "area": "Rates (shadow)", "proof": "validated", "status": "Data (confirm with rate owner)",
        "title": "SGUL15 shadow DB_DISCOUNT scale S missing (NAR undiscounted)",
        "description": "1U145700/1U146000/1U146100/1U146300 have ShadowDBDRate='Table' but no DB_DISCOUNT S row.",
        "evidence": "With DBD = SHADOW_INT: U0652930 -8.55 -> -0.19, UN004368 -34.05 -> +0.07, UXE00047 -25.61 -> +0.22. "
                    "Albert decomposition (verified closure): total effect -177.65 on the SGUL15 family.",
        "source": "Shadow deep: 4.2; Albert SGUL suiteview_diff (corroborating)",
        "fix": "Load DB_DISCOUNT S = SHADOW_INT for these plans (confirm with the rate owner).",
    },
    {
        "id": "R05", "area": "Rates + plan config (shadow)", "proof": "validated", "status": "Ready to fix",
        "title": "1U146200 (SGUL15S NY): active CCV but no ShadowPlancode and no scale S rates",
        "description": "has_shadow_account is True but shadow rates are skipped, so the shadow runs on zero COI, "
                       "interest, loads and DBD (RERUN shadow = premium sum).",
        "evidence": "4 policies; UNS00924 RERUN 15,058.00 vs XP 10,273.58.",
        "source": "Shadow deep: 4.4",
        "fix": "Load scale S (legacy CCV46100-style), set ShadowPlancode and raise when CCV has no shadow rates.",
    },
    {
        "id": "R06", "area": "Rates", "proof": "attributed", "status": "Needs Robert's ruling (issues #37-#44)",
        "title": "Rates loaded 9/30 from dbo fallback or awaiting source decisions",
        "description": "34 no-IAF UL/IUL/rider plans loaded with IAF-only grids from legacy dbo RATE_* (marked "
                       "IAF_SUBSTITUTE_SOURCE); SCR table 01 for 1X130100/1G130A00, MTP_TBL1 for 1U135K00 and TBL1 "
                       "add-ons from dbo RATE_TRGPREM; 1U135400 shadow S from the legacy CCV plancode.",
        "evidence": "review/bulk-load-issues.md #37-#44; dbo_check_loaded_plans_mismatches.json (271,642 cell "
                    "differences across print-vs-dbo rows for Robert's review).",
        "source": "Rates loader 9/30",
        "fix": "Robert to accept the dbo-sourced rows or pull IAF prints; replace rows when prints arrive.",
    },
    {
        "id": "H01", "area": "Harness", "proof": "validated", "status": "Fixed in this run",
        "title": "Loan balance compared on policy_debt (incl. next month's accrual)",
        "description": "Now compared on regular/preferred/variable principal + accrued at the monthliversary.",
        "evidence": "35 loan policies agree to $0.05 on the MV basis (AV deep C4a).",
        "source": "AV deep: C4a, fix #7",
        "fix": "Done (baseline_history_compare._state_values).",
    },
    {
        "id": "H02", "area": "Harness", "proof": "validated", "status": "Fixed in this run",
        "title": "SN/SM per-values-phase rows replayed as separate net withdrawals",
        "description": "Rows of one withdrawal event are now grouped; lead-row NET_AMT is the net request.",
        "evidence": "26/UNS01657 $25 fee charged twice; 01/U0107043 2,425 read as 13 net requests.",
        "source": "AV deep: C5, fix #8",
        "fix": "Done (baseline_history_compare._actual_transactions).",
    },
    {
        "id": "H04", "area": "Harness", "proof": "validated", "status": "Fixed in this run",
        "title": "PW waiver credits replayed as premium",
        "description": "PW is excluded from the replay and the policy flagged (the engine has no waiver-credit mode).",
        "evidence": "Albert F07; no PW row in any replay window here (flag count on the Policies sheet).",
        "source": "AV deep: fix #15",
        "fix": "Done; engine waiver-credit mode is a separate request.",
    },
    {
        "id": "H05", "area": "Harness", "proof": "validated", "status": "Fixed in this run",
        "title": "Rollback seed row classified as a difference",
        "description": "The seed row is CyberLife's own starting point; now 'Seed' and excluded from all counts.",
        "evidence": "Charge triage: 5 seed-only COI differences (1.02-2.16) under the current-target basis.",
        "source": "Charge triage (seed cluster)",
        "fix": "Done.",
    },
    {
        "id": "H06", "area": "Harness / CyberLife convention", "proof": "validated", "status": "Fixed in this run",
        "title": "COI vs OTHER split: CyberLife books the table/flat COI load in OTH_PRM_AMT",
        "description": "RERUN carries the substandard load in total_coi_charge; totals reconcile. Compared on the "
                       "total-charges measure (CINS+OTH vs COI+benefit+rider+asset); COI/OTHER rows marked H06.",
        "evidence": "Charge triage: 15 policies, COI diff + OTHER diff = 0.00, MD diff 0.00; Albert F08 agrees.",
        "source": "Charge triage (split cluster)",
        "fix": "Done (comparison). Optional: present the substandard load as OTHER for split matching.",
    },
    {
        "id": "H07", "area": "Harness", "proof": "validated", "status": "Fixed in this run",
        "title": "CyberLife MD recorded 0 while components are populated",
        "description": "MD is compared with the CINS+EXP+OTH component sum on those rows (flagged D01).",
        "evidence": "26/000218382, 26/UN005972: component sum = RERUN MD to the cent.",
        "source": "Charge triage (CyberLife anomalies)",
        "fix": "Done.",
    },
    {
        "id": "H08", "area": "Harness (limitation)", "proof": "validated", "status": "Limitation",
        "title": "Historical rider/benefit charge absent from the current-coverage basis",
        "description": "Under the harness basis (current coverage/targets) a rider that existed historically is "
                       "missing, so CyberLife OTH has no RERUN counterpart.",
        "evidence": "01/U0114057 OTH 18.57, 26/000291629 OTH 34.66: RERUN trace has no benefit/rider at the date.",
        "source": "Charge triage (historical rider cluster)",
        "fix": "Recover historical benefit/rider rows at the rollback date.",
    },
    {
        "id": "H09", "area": "Harness (limitation)", "proof": "attributed", "status": "Open (VPM007)",
        "title": "Premium load / target basis under the current-target harness basis",
        "description": "Net premium differs from CyberLife NET_AMT where the load split (up-to/over target, premiums "
                       "paid in the policy year) depends on historical targets/accumulators the harness does not "
                       "recover.",
        "evidence": "01/UE015003 NET_AMT 116.88/119.04 vs RERUN 114.08 on 124.00; 26/UNS01657 35.88 vs 37.44 on "
                    "39.00. Not proven to the cent (VPM007 open).",
        "source": "AV deep: open items",
        "fix": "Recover historical targets/premium accumulators at the rollback date; check VPM007.",
    },
    {
        "id": "H10", "area": "Harness (shadow)", "proof": "validated", "status": "Fixed (baseline_shadow_check.py)",
        "title": "Shadow compared on shadow_eav and without cash-with-application premiums",
        "description": "XP is post-deduction, pre-interest: compare shadow_av, and apply pre-issue premiums at issue.",
        "evidence": "Same engine run: 0/68 -> 15/68 within $1 on the corrected quantity.",
        "source": "Shadow deep: 1-2",
        "fix": "Done in tools/rerun/baseline_shadow_check.py.",
    },
    {
        "id": "D01", "area": "CyberLife data", "proof": "validated", "status": "Data (CyberLife)",
        "title": "CyberLife CD monthly deduction 0 with populated CINS/EXP/OTH",
        "description": "See H07.", "evidence": "26/000218382 (2026-04-15, 08-15), 26/UN005972.",
        "source": "Charge triage", "fix": "None in RERUN; report to CyberLife support if needed.",
    },
    {
        "id": "D02", "area": "CyberLife data", "proof": "validated", "status": "Data (raise with product)",
        "title": "CyberLife charges 1U1F4* WPMP (4M) past BNF_PAY_UP_DT",
        "description": "Engine stops at pay-up (Robert's illustration ruling); CyberLife keeps charging 1U1F4*.",
        "evidence": "Albert F12 5/5; re-read on DB2: 26/000308961 OTH 23.74 = 84.777 x 0.28 after 2024-04-09 "
                    "pay-up. No 1U1F4* policy past pay-up in this sample.",
        "source": "AV deep: Albert cross-reference F12", "fix": "Raise with CyberLife/product.",
    },
    {
        "id": "D03", "area": "CyberLife data", "proof": "validated", "status": "Data (anchor on a later value)",
        "title": "Unreplayable CyberLife history (AV adjustments S7/P7, CA/CC, historical rider charges)",
        "description": "From-issue shadow replays cannot represent manual AV adjustments or historical charge offsets.",
        "evidence": "U0437024 +4,994 (S7 989 in 2006 + CA 2022-24); SGUL15+CTR fixed offsets; anchored roll-"
                    "forwards from the June/July RGA seriatim close within $0.05.",
        "source": "Shadow deep: 3, 6", "fix": "Anchor on the seriatim/known value; no RERUN change.",
    },
    {
        "id": "D04", "area": "CyberLife data", "proof": "validated", "status": "Data (timing artifact)",
        "title": "Same-day processing artifact (9/30 monthliversary processed during the session)",
        "description": "UE219641's 9/30 MV was processed while the data were read; the premium was not yet in FH_FIXED.",
        "evidence": "Shadow -204.82 at 9/30 vs +0.07 at 8/31.", "source": "Shadow deep: per-plancode 1U147600",
        "fix": "Re-pull after processing completes.",
    },
    {
        "id": "D05", "area": "CyberLife history (shadow)", "proof": "attributed",
        "status": "Open (current processing reproduced)",
        "title": "Historical shadow offset with unidentified cause",
        "description": "From-issue shadow differs but the roll-forward anchored on the June/July 2026 seriatim "
                       "reproduces XP within $0.05, so today's processing matches and the offset sits in earlier "
                       "history (SGUL15+CTR rider charge history, LTGUL08 R-class, two 1035 lumps, SGUL20/21 timing).",
        "evidence": "Seriatim-anchored roll-forward within $0.05 on every such policy (shadow_root_causes.json).",
        "source": "Shadow deep: 0c, 6", "fix": "Research history only if a from-issue rebuild is ever required.",
    },
    {
        "id": "D06", "area": "CyberLife history (shadow)", "proof": "attributed",
        "status": "Open question for Robert/IT (not an engine fix)",
        "title": "Possibly CTR charges not taken from the shadow before about Dec 2019; no single cutoff fits every policy",
        "description": "The remaining SGUL misses (UE046797, UXE00064, UXE00073, UE134041, the 1U146100 group) are fixed "
                       "historical offsets that grow only at the shadow interest rate across the 3/6/7/8-2026 RGA "
                       "seriatim snapshots; current monthly processing (incl. CTR and PWSTP charges) matches. Sign: our "
                       "replay is LOWER than XP - the opposite of the PT/PF defect (D07), so a different issue.",
        "evidence": "Albert (verified: his seriatim values equal ours 78/78): UE046797 126.49 -> 128.73, UXE00064 "
                    "160.31 -> 163.21, UXE00073 187.02 -> 190.32, UE134041 564.51 -> 575.94, UE079106 11.68 -> 11.90 "
                    "(3/31 -> 8/31); 1U146100 close to -0.26..+0.67 with no CTR charge at 2019-12-20; no single cutoff "
                    "fits UE046797; UE134041 has no rider (year-1 1035 lumps are the lead there).",
        "source": "Robert 9/30 (keep open); Albert SGUL suiteview_diff README (corroborating); Shadow deep 6",
        "fix": "Confirm CyberLife's historical SGUL shadow CTR handling with Robert/IT; do not code as a rule.",
    },
    {
        "id": "D07", "area": "CyberLife data (shadow)", "proof": "validated",
        "status": "CyberLife known defect (SR in progress)",
        "title": "PT and PF premiums never went into the CyberLife shadow account",
        "description": "Robert 9/30: known CyberLife defect; on policies with PT/PF premiums XP is LOWER than "
                       "contractually correct. RERUN's from-issue replay (baseline_shadow_check PREMIUM_CODES include "
                       "PT/PF) is the contractually correct view; excluding PT/PF reproduces CyberLife as processed. "
                       "Any residual PT/PF explains is a CyberLife defect, not an engine issue.",
        "evidence": "Sample: 0 of 80 shadow policies have a PT/PF premium in their history and 0 of 302 six-month "
                    "replay windows contain one, so no residual here depends on it (both views identical). Out-of-"
                    "sample check (Albert's PT/PF examples, replayed here both ways): UE128242 as-processed +0.77 vs "
                    "correct +27,228.13; U0673236 -12.93 vs +26,231.19; UE029367 -7.11 vs +1,918.66; UE031057 -17.39 "
                    "vs +1,359.59 (as-processed residuals = E12/R04). Albert's preliminary population: 161 in-force "
                    "SGUL with PT/PF, ~$1.56M of shadow (corroborating, not verified).",
        "source": "Robert 9/30; final/shadow_check/ptpf_shadow.json and ptpf_examples.json; Albert SGUL README",
        "fix": "None in RERUN. Reconcile to CyberLife as-is by excluding PT/PF; report the contractually correct value "
               "(PT/PF included) for the SR.",
    },
]
ROOT_CAUSE_BY_ID = {item["id"]: item for item in ROOT_CAUSES}
FIX_PRIORITY = [
    "E02", "P03", "E04", "P02", "E03", "E01", "E11", "E14", "P01",
    "E12", "E08", "R05", "E10", "R04", "E07", "E09",
    "R03", "R01", "R02", "R06", "E05", "E06", "E13",
    "H01", "H02", "H04", "H05", "H06", "H07", "H08", "H09", "H10",
    "D01", "D02", "D03", "D04", "D05", "D06", "D07",
]
D06_POLICIES = {"01_UE046797", "04_UXE00064", "04_UXE00073", "01_UE079106", "01_UE079233", "01_UE079240",
                "01_UE079247", "01_UE079253", "01_UE134041"}
UNEXPLAINED_LABELS = {
    "U-AV": "AV change not reproduced by the decomposition",
    "U-INT": "Interest not reproduced by CyberLife's documented rule or the carried difference",
    "U-CHG": "Charge difference with no validated rule",
    "U-LOAN": "Loan difference with no loan cash flow or known rate issue",
    "U-WD": "Withdrawal applied differently with no cap",
    "U-DATE": "RERUN has no state on the CyberLife monthliversary date",
    "U-SHD": "Shadow residual not identified",
}


def _proof(rc_id: str) -> str:
    if rc_id in UNEXPLAINED_IDS:
        return "unexplained"
    if rc_id == "RND":
        return "rounding"
    return ROOT_CAUSE_BY_ID.get(rc_id, {}).get("proof", "unexplained")


# ── classification ───────────────────────────────────────────────────────────

def _num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _attr_text(attr: Counter, *, minimum: float = MINOR, limit: int = 4) -> str:
    items = [(k, v) for k, v in attr.items() if k != "RND" and abs(v) >= minimum]
    items.sort(key=lambda kv: -abs(kv[1]))
    return "; ".join(f"{k} {v:+,.2f}" for k, v in items[:limit])


def _decide(diff: float | None, attr: Counter, *, seed: bool = False) -> str:
    """Class for one measure: Seed / Not comparable / Exact / Rounding / Explained / Material / Unexplained."""
    if seed:
        return "Seed"
    if diff is None:
        return "Not comparable"
    if abs(diff) <= EXACT_TOLERANCE:
        return "Exact"
    if abs(diff) <= ROUNDING_TOLERANCE:
        return "Rounding"
    unexplained = sum(v for k, v in attr.items() if _proof(k) == "unexplained")
    attributed = sum(v for k, v in attr.items() if _proof(k) == "attributed")
    explained = sum(v for k, v in attr.items() if _proof(k) in ("validated", "rounding"))
    if abs(unexplained) > ROUNDING_TOLERANCE or abs(diff - explained - attributed) > ROUNDING_TOLERANCE:
        return "Unexplained"
    if abs(attributed) > ROUNDING_TOLERANCE:
        return "Material"
    return "Explained"


def _spread(amount: float, shares: Counter) -> Counter:
    """Distribute ``amount`` over the IDs of ``shares`` in proportion to their (signed) size."""
    total = sum(shares.values())
    out: Counter = Counter()
    if abs(total) < 1e-9:
        return out
    for key, value in shares.items():
        out[key] += amount * value / total
    return out


def _dominant(shares: Counter) -> Counter:
    """Keep the non-rounding IDs that make up the carried difference (drop tiny/offsetting cents)."""
    kept = Counter({k: v for k, v in shares.items() if k != "RND" and abs(v) >= 0.05})
    return kept if kept else Counter(shares)


def _day_count_id(plancode: str) -> str:
    return "P01" if plancode in C1B_PLANS else "E01"


def _factor(rate: float, days: float) -> float:
    try:
        return (1.0 + rate) ** (days / 365.0) - 1.0
    except (TypeError, ValueError, OverflowError):
        return 0.0


def _classify_charges(result: dict[str, Any], month: dict[str, Any], carried: Counter) -> dict[str, Counter]:
    """Attribute material charge differences (COI/expense/other/total/MD) for one month.

    Rules, in order: future-dated segment (E04); COI/OTHER split with offsetting
    diffs (H06); historical rider/benefit charge absent under the harness basis
    (H08); NAR-driven difference whose size matches (RERUN NAR - CyberLife NAR_AMT)
    x RERUN COI rate, attributed to the AV difference entering the month; else U-CHG.
    """
    when = parse_date(month["date"])
    diffs = {meas: _num((month.get(meas) or {}).get("diff")) for meas in CHARGE_MEASURES}
    future = [
        seg for seg in result.get("future_segments") or []
        if parse_date(seg.get("issue_date")) and parse_date(seg.get("issue_date")) > when
    ]
    coi_d, oth_d = diffs["coi"], diffs["other"]
    split = (
        coi_d is not None and oth_d is not None and abs(coi_d) > ROUNDING_TOLERANCE
        and abs(oth_d) > ROUNDING_TOLERANCE and coi_d * oth_d < 0
    )
    moved = math.copysign(min(abs(coi_d), abs(oth_d)), coi_d) if split else 0.0
    coi_r = None if coi_d is None else coi_d - moved
    oth_r = None if oth_d is None else oth_d + moved
    other = month.get("other") or {}
    hist_rider = (
        str(result.get("rollback_basis", "")).startswith("harness basis")
        and (_num(other.get("cyberlife")) or 0.0) > ROUNDING_TOLERANCE
        and abs(_num(other.get("rerun")) or 0.0) <= EXACT_TOLERANCE
        and not split
    )
    nar = month.get("nar") or {}
    coi = month.get("coi") or {}
    rr_nar, cl_nar, rr_coi = _num(nar.get("rerun")), _num(nar.get("cyberlife")), _num(coi.get("rerun"))
    expected_coi = None
    if rr_nar and cl_nar is not None and rr_coi is not None:
        expected_coi = (rr_nar - cl_nar) * rr_coi / rr_nar
    coi_via_nar = (
        expected_coi is not None and coi_r is not None and abs(expected_coi) > MINOR
        and abs(coi_r - expected_coi) <= max(ROUNDING_TOLERANCE, 0.15 * abs(expected_coi))
    )
    out: dict[str, Counter] = {}
    for meas in CHARGE_MEASURES:
        diff = diffs[meas]
        attr: Counter = Counter()
        if diff is None:
            out[meas] = attr
            continue
        rest = diff
        if split and meas == "coi":
            attr["H06"] += moved
            rest = coi_r
        elif split and meas == "other":
            attr["H06"] -= moved
            rest = oth_r
        if abs(rest) <= ROUNDING_TOLERANCE:
            attr["RND"] += rest
        elif future:
            attr["E04"] += rest
        elif meas in ("other", "total_charges", "md") and hist_rider and rest < 0:
            attr["H08"] += rest
        elif coi_via_nar and (
            meas == "coi"
            or (meas in ("total_charges", "md")
                and abs(rest - (coi_r or 0.0) - (oth_r or 0.0)) <= ROUNDING_TOLERANCE)
            or (meas == "other" and (coi_r or 0.0) * rest > 0)
        ) and abs(sum(carried.values())) > MINOR:
            attr.update(_spread(rest, _dominant(carried)))
        else:
            attr["U-CHG"] += rest
        out[meas] = attr
    return out


def _classify_loan(result: dict[str, Any], diff: float | None) -> Counter:
    attr: Counter = Counter()
    if diff is None or abs(diff) <= ROUNDING_TOLERANCE:
        if diff is not None:
            attr["RND"] = diff
        return attr
    kinds = {tx.get("kind") for tx in result.get("transactions_replayed") or []}
    plancode = str(result.get("plancode") or "").upper()
    if plancode in P02_PLANS:
        attr["P02"] = diff
    elif "loan_repayment" in kinds:
        attr["E03"] = diff
    elif "loan" in kinds:
        attr["E11"] = diff
    else:
        attr["U-LOAN"] = diff
    return attr


def _loan_contributors(result: dict[str, Any]) -> list[str]:
    kinds = {tx.get("kind") for tx in result.get("transactions_replayed") or []}
    plancode = str(result.get("plancode") or "").upper()
    out = []
    if plancode in P02_PLANS:
        out.append("P02")
    if "loan_repayment" in kinds:
        out.append("E03")
    if kinds & {"loan", "loan_repayment"}:
        out.append("E11")
    return out


def classify_policy(result: dict[str, Any], decomposition: dict[str, Any] | None) -> dict[str, Any]:
    """Classify every month and measure of one compared policy.

    Returns ``{"months": [...], "av_cum": Counter, "ids": Counter}`` where each month
    holds per-measure ``cyberlife``, ``rerun``, ``diff``, ``class``, ``attr`` (Counter of
    root-cause ID -> dollars, summing to the diff) and ``text``.
    """
    key = result["key"]
    plancode = str(result.get("plancode") or "").upper()
    dc_id = _day_count_id(plancode)
    dec_rows = {row["cur"]: row for row in (decomposition or {}).get("rows", [])}
    cum: Counter = Counter()
    months = []
    loan_extra = _loan_contributors(result)
    for month in result.get("monthly") or []:
        if "av" not in month:
            months.append({
                "date": month.get("date"), "seed": False, "status": month.get("status", ""),
                "measures": {meas: {"cyberlife": None, "rerun": None, "diff": None, "class": "Not comparable",
                                    "attr": Counter({"E14": 0.0}), "text": "E14: no RERUN state on this MV date"}
                             for meas in MEASURES},
                "flags": ["E14: " + str(month.get("status", ""))],
            })
            continue
        seed = bool(month.get("seed_row"))
        measures: dict[str, dict[str, Any]] = {}
        for meas in MEASURES:
            item = month.get(meas) or {}
            measures[meas] = {
                "cyberlife": _num(item.get("cyberlife")), "rerun": _num(item.get("rerun")),
                "diff": _num(item.get("diff")), "attr": Counter(), "text": "",
            }
        flags = []
        if month.get("md_component_sum_substituted"):
            flags.append("D01/H07: CyberLife MD recorded 0; compared with CINS+EXP+OTH")
        if month.get("cyberlife_md_from_components"):
            flags.append("CyberLife MD = CINS+EXP+OTH (no recorded CD)")
        if month.get("cyberlife_components_null_as_zero"):
            flags.append("Null " + "/".join(month["cyberlife_components_null_as_zero"]) + " read as 0")
        if seed:
            cum["RND"] += measures["av"]["diff"] or 0.0
            for meas in MEASURES:
                measures[meas]["class"] = "Seed"
            months.append({"date": month["date"], "seed": True, "status": "", "measures": measures, "flags": flags})
            continue
        d = dec_rows.get(month["date"])
        gap_month = False
        if d is not None and parse_date(d["cur"]) and parse_date(d["prev"]):
            gap_month = (parse_date(d["cur"]) - parse_date(d["prev"])).days > 31
            if gap_month:
                d = None
                flags.append("E14: previous CyberLife MV has no RERUN state; AV change spans several months")
        av_diff = measures["av"]["diff"]
        prev_total = sum(cum.values())
        interest_attr: Counter = Counter()
        flow_attr: Counter = Counter()
        if d is not None:
            comps = d["components_rerun_minus_cyber"]
            engine = d.get("engine") or {}
            interest_attr[dc_id] += comps.get("day_count_or_compounding", 0.0)
            interest_attr["E11"] += (
                comps.get("premium_interest_from_receipt", 0.0) + comps.get("withdrawal_interest_adjustment", 0.0)
                + comps.get("loan_tx_interest_adjustment", 0.0)
            )
            crr = comps.get("cyber_rule_minus_recorded", 0.0)
            if key in R01_POLICIES:
                interest_attr["R01"] += crr
            elif key in R02_POLICIES:
                interest_attr["R02"] += crr
            elif abs(crr) <= 0.25:
                interest_attr["RND"] += crr
            else:
                interest_attr["U-INT"] += crr
            carried = comps.get("interest_on_carried_difference_and_other", 0.0)
            expected = prev_total * _factor(_num(engine.get("rate")) or 0.0, _num(engine.get("days")) or 30.0)
            if abs(carried - expected) <= max(MINOR, 0.25 * abs(expected)):
                if abs(sum(cum.values())) > 1e-6:
                    interest_attr.update(_spread(carried, _dominant(cum)))
                else:
                    interest_attr["RND"] += carried
            else:
                if abs(sum(cum.values())) > 1e-6:
                    interest_attr.update(_spread(expected, _dominant(cum)))
                interest_attr["U-INT"] += carried - expected
            npd = comps.get("net_premium_rerun_minus_cyber", 0.0)
            if abs(npd) > 0.05:
                allowance = engine.get("allowance") or {}
                npt = _num(allowance.get("NPT Allowance0"))
                if engine.get("premium_capped_by_tamra") and npt is not None and abs(npt) < 0.005:
                    flow_attr["E02"] += npd
                else:
                    flow_attr["H09"] += npd
            else:
                flow_attr["RND"] += npd
            wdd = comps.get("withdrawal_rerun_minus_cyber", 0.0)
            if abs(wdd) > 0.05:
                capped = (_num(engine.get("input_withdrawal")) or 0.0) > (_num(engine.get("max_net_withdrawal")) or 0.0) + 0.01
                flow_attr["P03" if capped else "U-WD"] += wdd
            else:
                flow_attr["RND"] += wdd
            unexplained = comps.get("unexplained", 0.0)
            flow_attr["U-AV" if abs(unexplained) > 0.05 else "RND"] += unexplained
            md_component = comps.get("md_rerun_minus_cyber", 0.0)
        else:
            md_component = None
        partial = Counter(cum)
        for source in (interest_attr, flow_attr):
            for k, v in source.items():
                partial[k] += v
        charge_attr = _classify_charges(result, month, partial)
        for meas in CHARGE_MEASURES:
            measures[meas]["attr"] = charge_attr[meas]
        md_attr = Counter()
        if md_component is not None:
            md_measure = charge_attr.get("md") or Counter()
            md_diff = measures["md"]["diff"]
            if md_diff is not None and abs(md_diff) > ROUNDING_TOLERANCE and abs(md_component) > 0.05:
                md_attr = _spread(md_component, Counter({k: v for k, v in md_measure.items()}))
            elif abs(md_component) > ROUNDING_TOLERANCE:
                md_attr = _spread(md_component, Counter(charge_attr.get("total_charges") or {"U-CHG": 1.0}))
                if not md_attr:
                    md_attr = Counter({"U-CHG": md_component})
            else:
                md_attr = Counter({"RND": md_component})
        month_attr = Counter()
        for source in (interest_attr, flow_attr, md_attr):
            for k, v in source.items():
                month_attr[k] += v
        if d is None and av_diff is not None:
            month_attr["U-DATE" if gap_month else "U-AV"] += av_diff - sum(cum.values())
        for k, v in month_attr.items():
            cum[k] += v
        if av_diff is not None:
            residual = av_diff - sum(cum.values())
            if abs(residual) > 0.05:
                cum["U-AV"] += residual
            elif residual:
                cum["RND"] += residual
        measures["av"]["attr"] = Counter(cum)
        measures["interest"]["attr"] = interest_attr
        measures["loan"]["attr"] = _classify_loan(result, measures["loan"]["diff"])
        for meas in MEASURES:
            m = measures[meas]
            m["class"] = _decide(m["diff"], m["attr"])
            m["text"] = _attr_text(m["attr"])
            if meas == "loan" and m["class"] not in ("Exact", "Rounding") and len(loan_extra) > 1:
                m["text"] = f"{m['text']} [contributors: {', '.join(loan_extra)}]"
            if m["class"] in ("Exact", "Rounding") and meas not in ("av", "interest"):
                m["text"] = ""
        if measures["coi"]["class"] == "Explained" and "H06" in measures["coi"]["attr"]:
            flags.append("H06: CyberLife books table/flat COI in OTH_PRM_AMT")
        months.append({"date": month["date"], "seed": False, "status": "", "measures": measures, "flags": flags})
    ids: Counter = Counter()
    for month in months:
        if month["seed"]:
            continue
        for meas, m in month["measures"].items():
            for k, v in m["attr"].items():
                if k != "RND" and (abs(v) > ROUNDING_TOLERANCE or k in ("U-DATE", "E14")):
                    ids[k] += 1
    return {"key": key, "months": months, "av_cum": cum, "ids": ids}


# ── loading and aggregation ─────────────────────────────────────────────────

def load_run(results_dir: Path) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]], dict[str, Any]]:
    """Return (policy results, decomposition by key, run manifest) for a run folder."""
    results = [json_load(path) for path in sorted((results_dir / "policy_results").glob("*.json"))]
    decomp_path = results_dir / "av_decomposition.json"
    decomposition = {}
    if decomp_path.exists():
        decomposition = {row["key"]: row for row in json_load(decomp_path).get("policies", [])}
    manifest_path = results_dir / "run_manifest.json"
    manifest = json_load(manifest_path) if manifest_path.exists() else {"runs": []}
    return results, decomposition, manifest


def classify_run(results: list[dict[str, Any]], decomposition: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {
        result["key"]: classify_policy(result, decomposition.get(result["key"]))
        for result in results if result.get("status") == "compared"
    }


def measure_stats(classified: dict[str, dict[str, Any]], keys: set[str] | None = None) -> dict[str, Counter]:
    """Per measure: Counter of classes over non-seed months (optionally restricted to ``keys``)."""
    stats: dict[str, Counter] = {meas: Counter() for meas in MEASURES}
    for key, item in classified.items():
        if keys is not None and key not in keys:
            continue
        for month in item["months"]:
            if month["seed"]:
                continue
            for meas in MEASURES:
                stats[meas][month["measures"][meas]["class"]] += 1
    return stats


def pct(counter: Counter, *classes: str) -> float | None:
    compared = sum(counter[c] for c in CLASS_ORDER)
    if not compared:
        return None
    return sum(counter[c] for c in classes) / compared


# ── workbook styling ─────────────────────────────────────────────────────────

CLASS_FILLS = {
    "Exact": "C6EFCE", "Rounding": "E2EFDA", "Explained": "DDEBF7", "Material": "FFEB9C",
    "Unexplained": "FFC7CE", "Seed": "EDEDED", "Not comparable": "EDEDED",
}
HEADER_FILL = "1F3864"
TITLE_COLOR = "1F3864"
MONEY_FMT = '#,##0.00;[Red]-#,##0.00;0.00'
PCT_FMT = '0.0%'
INT_FMT = '#,##0'


class Col:
    """Column spec for :func:`write_table`: header, kind (text/money/pct/int/class/wrap) and width."""

    def __init__(self, header: str, kind: str = "text", width: float | None = None):
        self.header = header
        self.kind = kind
        self.width = width


def _style_objects():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    return {
        "header_font": Font(name="Calibri", size=9, bold=True, color="FFFFFF"),
        "header_fill": PatternFill("solid", fgColor=HEADER_FILL),
        "body_font": Font(name="Calibri", size=9),
        "bold": Font(name="Calibri", size=9, bold=True),
        "title_font": Font(name="Calibri", size=14, bold=True, color=TITLE_COLOR),
        "h2_font": Font(name="Calibri", size=11, bold=True, color=TITLE_COLOR),
        "note_font": Font(name="Calibri", size=9, italic=True, color="595959"),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "top": Alignment(vertical="top"),
        "center": Alignment(horizontal="center", vertical="top"),
        "header_align": Alignment(wrap_text=True, vertical="center", horizontal="center"),
        "bottom_border": Border(bottom=Side(style="thin", color="BFBFBF")),
        "fills": {name: PatternFill("solid", fgColor=color) for name, color in CLASS_FILLS.items()},
    }


def _excel_value(value: Any) -> Any:
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, default=str)[:32000]
    if isinstance(value, str) and len(value) > 32000:
        return value[:32000]
    return value


def write_table(ws, columns: list[Col], rows: list[list[Any]], *, start_row: int = 1, start_col: int = 1,
                freeze: bool = True, autofilter: bool = True, conditional: bool = True) -> int:
    """Write a dense table (header + rows) and return the next free row.

    Money/percent/integer formats, wrapped text columns, frozen header and filters
    are applied; ``class`` columns get conditional fills (Exact/Rounding/Explained/
    Material/Unexplained).
    """
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter

    st = _style_objects()
    for offset, col in enumerate(columns):
        cell = ws.cell(row=start_row, column=start_col + offset, value=col.header)
        cell.font = st["header_font"]
        cell.fill = st["header_fill"]
        cell.alignment = st["header_align"]
    for r_offset, row in enumerate(rows, start=1):
        for offset, col in enumerate(columns):
            value = _excel_value(row[offset] if offset < len(row) else None)
            cell = ws.cell(row=start_row + r_offset, column=start_col + offset, value=value)
            cell.font = st["body_font"]
            if col.kind == "money":
                cell.number_format = MONEY_FMT
            elif col.kind == "pct":
                cell.number_format = PCT_FMT
            elif col.kind == "int":
                cell.number_format = INT_FMT
            if col.kind == "wrap":
                cell.alignment = st["wrap"]
            elif col.kind == "class":
                cell.alignment = st["center"]
            else:
                cell.alignment = st["top"]
    last_row = start_row + max(len(rows), 1)
    for offset, col in enumerate(columns):
        letter = get_column_letter(start_col + offset)
        width = col.width
        if width is None:
            sample = [len(str(col.header))] + [len(str(_excel_value(r[offset]) or "")) for r in rows[:400] if offset < len(r)]
            width = min(max(sample) + 2, 50)
            if col.kind in ("money",):
                width = max(width, 11)
        current = ws.column_dimensions[letter].width
        if not current or start_row == 1 or width > current:
            ws.column_dimensions[letter].width = width
        if col.kind == "class" and conditional and rows:
            rng = f"{letter}{start_row + 1}:{letter}{last_row}"
            for name, fill in st["fills"].items():
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{name}"'], fill=fill))
    ws.row_dimensions[start_row].height = 30
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=start_col)
    if autofilter and rows:
        ws.auto_filter.ref = (
            f"{get_column_letter(start_col)}{start_row}:{get_column_letter(start_col + len(columns) - 1)}{last_row}"
        )
    return last_row + 1


def write_title(ws, row: int, text: str, *, level: int = 1, col: int = 1) -> int:
    st = _style_objects()
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = st["title_font"] if level == 1 else st["h2_font"]
    return row + 1


def write_note(ws, row: int, text: str, *, col: int = 1, merge_to: int = 0) -> int:
    st = _style_objects()
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = st["note_font"]
    cell.alignment = st["wrap"]
    if merge_to:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=merge_to)
        ws.row_dimensions[row].height = max(15, 12 * (1 + len(text) // 160))
    return row + 1


def policy_label(key: str) -> str:
    return key.replace("_", "/", 1)


# ── evidence assembly ────────────────────────────────────────────────────────

WORST_ORDER = {"Unexplained": 5, "Material": 4, "Explained": 3, "Rounding": 2, "Exact": 1, "Not comparable": 0, "Seed": 0}
SHADOW_CATEGORY_IDS = {
    "comparison-definition": ["H10"],
    "input-history": ["H10"],
    "engine rule (interest basis)": ["E07"],
    "rate data (DB discount)": ["R04", "E10"],
    "rate data (shadow target)": ["E08"],
    "engine rule (LTGUL load relief)": ["E08"],
    "rate data (missing)": ["R03"],
    "rate data / engine rule (no shadow config)": ["R05", "E10"],
    "engine rule (DB option)": ["E09"],
    "engine rule (withdrawals)": ["E09"],
    "CyberLife anomaly / unreplayable history": ["D03"],
    "unresolved residual": ["U-SHD"],
}


def _shadow_ids(record: dict[str, Any]) -> list[str]:
    ids: list[str] = []
    family = str(record.get("family") or "")
    anchored = [abs(record[k]) for k in ("anchored_from_2026-06-30_diff_vs_xp", "anchored_from_2026-07-31_diff_vs_xp")
                if record.get(k) is not None]
    for category in record.get("root_cause_category") or []:
        if category == "input-timing":
            mapped = ["E12"] if family.startswith("SGUL") else ["E11"]
        elif category == "unresolved residual":
            mapped = ["D05"] if anchored and max(anchored) <= 0.05 else ["U-SHD"]
        else:
            mapped = SHADOW_CATEGORY_IDS.get(category, ["U-SHD"])
        for rc_id in mapped:
            if rc_id not in ids:
                ids.append(rc_id)
    if record.get("key") == "01_UE219641":
        ids = [i for i in ids if i not in ("U-SHD", "D05")] + ["D04"]
    if record.get("key") in D06_POLICIES and "D06" not in ids:
        ids.append("D06")
    return ids


def _shadow_effects(record: dict[str, Any]) -> dict[str, float]:
    """Dollar effect per root-cause ID from the investigation's per-category effects."""
    effects: dict[str, float] = defaultdict(float)
    family = str(record.get("family") or "")
    for category, amount in (record.get("category_effects_dollars") or {}).items():
        base = re.sub(r"\s*\(.*$", "", category) if category.startswith("unresolved") else category
        if base == "input-timing":
            mapped = ["E12"] if family.startswith("SGUL") else ["E11"]
        else:
            mapped = SHADOW_CATEGORY_IDS.get(base, [])
        for rc_id in mapped:
            effects[rc_id] = max(effects[rc_id], abs(float(amount or 0.0)))
    return dict(effects)


def read_albert_sgul(path: Path | None) -> dict[str, dict[str, float | None]]:
    """Per-policy SGUL rule effects from Albert's shadow decomposition CSV (corroborating input)."""
    import csv
    import io

    if path is None or not path.exists():
        return {}
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    out = {}
    for row in csv.DictReader(io.StringIO(text or "")):
        out[row["key"]] = {
            name: _num(row.get(column)) for name, column in (
                ("ref_minus_xp", "albert_minus_xp"), ("engine_minus_xp", "engine_minus_xp"),
                ("premium_timing", "effect_premium_timing"), ("dbd", "effect_dbd"),
                ("interest", "effect_interest"), ("load_rounding", "effect_load_rounding"),
                ("closure", "closure_residual"),
            )
        }
        out[row["key"]]["in_scope"] = str(row.get("in_scope", "")).strip().lower() in ("true", "1", "yes")
    return out


def shadow_records(shadow_findings: dict[str, Any] | None, shadow_extra: dict[str, Any] | None,
                   albert: dict[str, dict[str, Any]] | None = None,
                   aps205_extra: dict[str, dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    """Normalise the shadow investigation records plus any newly checked policies.

    ``aps205_extra`` maps a newly checked policy key to a VP/MS APS205 replay result
    (``aps205`` diff vs XP, ``rule_set``, ``ids``, ``family``).
    """
    albert = albert or {}
    aps205_extra = aps205_extra or {}
    rows: list[dict[str, Any]] = []
    for rec in (shadow_findings or {}).get("records", []):
        anchored = [rec.get(k) for k in ("anchored_from_2026-06-30_diff_vs_xp", "anchored_from_2026-07-31_diff_vs_xp")
                    if rec.get(k) is not None]
        effects = _shadow_effects(rec)
        alb = albert.get(rec["key"], {})
        for rc_id, name in (("E12", "premium_timing"), ("R04", "dbd"), ("E07", "interest")):
            if alb.get(name) is not None and abs(alb[name]) > effects.get(rc_id, 0.0):
                effects[rc_id] = abs(alb[name])
        ids = _shadow_ids(rec)
        for rc_id, name in (("E12", "premium_timing"), ("R04", "dbd"), ("E07", "interest")):
            if alb.get(name) is not None and abs(alb[name]) >= 0.01 and rc_id not in ids:
                ids.insert(1, rc_id)
        rows.append({
            "key": rec["key"], "plancode": rec.get("plancode"), "family": rec.get("family"),
            "xp": rec.get("cyberlife_value"), "rerun": rec.get("rerun_value"), "diff": rec.get("diff"),
            "prior_diff": rec.get("prior_tool_diff"),
            "after_rules": rec.get("diff_after_best_rules"), "rule_set": rec.get("best_rule_set", ""),
            "aps205": rec.get("vpms_replay_diff"), "vpms_model": rec.get("vpms_model", ""),
            "seriatim_jun": rec.get("seriatim_shadow_2026-06-30"), "seriatim_jul": rec.get("seriatim_shadow_2026-07-31"),
            "anchored_jun": rec.get("anchored_from_2026-06-30_diff_vs_xp"),
            "anchored_jul": rec.get("anchored_from_2026-07-31_diff_vs_xp"),
            "anchored_max": max((abs(v) for v in anchored), default=None),
            "primary": rec.get("primary_category", ""), "ids": ids, "effects": effects,
            "albert": alb,
            "root_cause": rec.get("root_cause", ""), "source": "Shadow deep (9/30) shadow_root_causes.json",
        })
    seen = {row["key"] for row in rows}
    for rec in (shadow_extra or {}).get("records", []):
        sens = rec.get("sensitivity_diff_vs_xp") or {}
        best_name, best = None, None
        for name in ("receipt_date_premium_interest", "monthly_interest+receipt_interest"):
            if sens.get(name) is not None and (best is None or abs(sens[name]) < abs(best)):
                best_name, best = name, sens[name]
        codes = (rec.get("history") or {}).get("unreplayed_value_codes") or {}
        ids = ["H10"]
        if best is not None and abs(best) <= ROUNDING_TOLERANCE:
            ids.append("E11")
        elif codes:
            ids.append("U-SHD")
        else:
            ids.append("U-SHD")
        row = {
            "key": rec["key"], "plancode": rec.get("plancode"), "family": "FPUL2" if rec.get("plancode") == "1U135400" else "",
            "xp": rec.get("cyberlife_xp"), "rerun": rec.get("rerun_shadow_av"), "diff": rec.get("diff"),
            "prior_diff": None, "after_rules": best, "rule_set": best_name or "",
            "aps205": None, "vpms_model": "", "seriatim_jun": None, "seriatim_jul": None,
            "anchored_jun": None, "anchored_jul": None, "anchored_max": None,
            "primary": "newly compared after the 9/30 scale-S load",
            "ids": ids, "effects": {}, "albert": {},
            "root_cause": (
                f"Engine as-is diff {rec.get('diff')}; best receipt-date variant {best_name} = {best}"
                + (f"; unreplayed CyberLife value codes {sorted(codes)}" if codes else "")
            ),
            "source": "baseline_shadow_check.py --sensitivity (final run)",
        }
        if rec["key"] in aps205_extra:
            extra = aps205_extra[rec["key"]]
            row.update({
                "family": extra.get("family", row["family"]), "after_rules": extra.get("aps205"),
                "aps205": extra.get("aps205"), "rule_set": extra.get("rule_set", ""), "ids": list(extra.get("ids") or ids),
                "root_cause": f"Engine as-is diff {rec.get('diff')} (shadow target MTP S = 0); VP/MS APS205 replay "
                              f"diff {extra.get('aps205')}",
                "source": "baseline_shadow_check.py (final run) + shadow_deep/vpms_aps205_replay.py",
            })
        if rec["key"] in seen:
            rows = [r for r in rows if r["key"] != rec["key"]]
        rows.append(row)
    return rows


def shadow_stats(rows: list[dict[str, Any]]) -> dict[str, Any]:
    compared = [r for r in rows if r.get("diff") is not None]
    no_config = [r for r in compared if "R05" in r["ids"]]
    core = [r for r in compared if "R05" not in r["ids"]]
    anchored = [r for r in rows if r.get("anchored_max") is not None]
    aps = [r for r in rows if r.get("aps205") is not None]
    aps_model = [r for r in aps if r.get("family") in ("LTGUL", "EXECUL")]
    aps_fpul = [r for r in aps if r.get("family") == "FPUL2"]
    return {
        "records": len(rows),
        "compared": len(compared),
        "compared_excl_1U146200": len(core),
        "asis_within_1": sum(abs(r["diff"]) <= ROUNDING_TOLERANCE for r in core),
        "after_rules_within_1": sum(r.get("after_rules") is not None and abs(r["after_rules"]) <= ROUNDING_TOLERANCE for r in core),
        "after_rules_within_5": sum(r.get("after_rules") is not None and abs(r["after_rules"]) <= 5 for r in core),
        "no_shadow_config": len(no_config),
        "anchored": len(anchored),
        "anchored_within_005": sum(r["anchored_max"] <= 0.05 for r in anchored),
        "anchored_exceptions": [r["key"] for r in anchored if r["anchored_max"] > 0.05],
        "aps205": len(aps),
        "aps205_within_001": sum(abs(r["aps205"]) <= 0.01 for r in aps),
        "aps205_within_1": sum(abs(r["aps205"]) <= ROUNDING_TOLERANCE for r in aps),
        "aps205_model": len(aps_model),
        "aps205_model_cent": sum(abs(r["aps205"]) <= 0.01 for r in aps_model),
        "aps205_fpul2": len(aps_fpul),
        "aps205_fpul2_within_035": sum(abs(r["aps205"]) <= 0.35 for r in aps_fpul),
        "not_compared": [r["key"] for r in rows if r.get("diff") is None],
        "unexplained": [r["key"] for r in rows if "U-SHD" in r["ids"]],
        "historical_open": [r["key"] for r in rows if "D05" in r["ids"]],
        "albert_in_scope": sum(1 for r in rows if (r.get("albert") or {}).get("in_scope")),
        "albert_ref_within_1": sum(
            1 for r in rows if (r.get("albert") or {}).get("in_scope")
            and r["albert"].get("ref_minus_xp") is not None and abs(r["albert"]["ref_minus_xp"]) <= ROUNDING_TOLERANCE
        ),
        "albert_effect_totals": {
            name: round(sum((r.get("albert") or {}).get(name) or 0.0 for r in rows if (r.get("albert") or {}).get("in_scope")), 2)
            for name in ("premium_timing", "dbd", "interest", "load_rounding")
        },
    }


def id_impacts(ev: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """Per root-cause ID: affected policies/plancodes/months and max $ by measure in this run + shadow."""
    out: dict[str, dict[str, Any]] = defaultdict(lambda: {
        "policies": set(), "plancodes": set(), "months": 0, "max_av": 0.0, "max_charge": 0.0,
        "max_interest": 0.0, "max_loan": 0.0, "shadow_policies": set(), "max_shadow": 0.0,
    })
    for key, item in ev["classified"].items():
        plancode = ev["by_key"][key].get("plancode")
        for month in item["months"]:
            if month["seed"]:
                continue
            month_ids = set()
            for meas, m in month["measures"].items():
                for rc_id, amount in m["attr"].items():
                    if rc_id == "RND" or (abs(amount) <= ROUNDING_TOLERANCE and rc_id not in ("E14", "U-DATE")):
                        continue
                    info = out[rc_id]
                    info["policies"].add(key)
                    info["plancodes"].add(plancode)
                    month_ids.add(rc_id)
                    bucket = {"av": "max_av", "interest": "max_interest", "loan": "max_loan"}.get(meas, "max_charge")
                    info[bucket] = max(info[bucket], abs(amount))
            for rc_id in month_ids:
                out[rc_id]["months"] += 1
    for row in ev["shadow_rows"]:
        residual = row.get("after_rules") if row.get("after_rules") is not None else row.get("diff")
        for rc_id in row["ids"]:
            info = out[rc_id]
            info["shadow_policies"].add(row["key"])
            info["plancodes"].add(row["plancode"])
            amount = row.get("effects", {}).get(rc_id)
            if amount is None and rc_id in ("U-SHD", "D05", "D03", "D04") and residual is not None:
                amount = abs(residual)
            if amount is not None:
                info["max_shadow"] = max(info["max_shadow"], abs(amount))
    for key, result in ev["by_key"].items():
        if result.get("status") == "error" and "RateLookupError" in str(result.get("error", "")):
            out["R03"]["policies"].add(key)
            out["R03"]["plancodes"].add(result.get("plancode"))
    return out


def impact_text(info: dict[str, Any]) -> str:
    parts = []
    if info["policies"]:
        parts.append(f"6-mo: {len(info['policies'])} pol / {info['months']} months")
    if info["shadow_policies"]:
        parts.append(f"shadow: {len(info['shadow_policies'])} pol")
    return "; ".join(parts) if parts else "none in this sample (latent / external)"


def max_impact(info: dict[str, Any]) -> float:
    return max(info["max_av"], info["max_charge"], info["max_interest"], info["max_loan"], info["max_shadow"])


# ── sheets: fix list, by plancode, feature coverage, policies ───────────────

def sheet_fix_list(wb, ev: dict[str, Any]) -> list[dict[str, Any]]:
    ws = wb.create_sheet("Fix List")
    impacts = ev["impacts"]
    rows = []
    records = []
    order = {rc_id: i for i, rc_id in enumerate(FIX_PRIORITY)}
    for item in sorted(ROOT_CAUSES, key=lambda it: order.get(it["id"], len(order))):
        info = impacts.get(item["id"]) or impacts.default_factory()
        record = {
            "id": item["id"], "area": item["area"], "title": item["title"], "status": item["status"],
            "proof": item["proof"], "policies": len(info["policies"]), "shadow_policies": len(info["shadow_policies"]),
            "plancodes": len({p for p in info["plancodes"] if p}), "months": info["months"],
            "max_av": info["max_av"], "max_charge": info["max_charge"], "max_interest": info["max_interest"],
            "max_loan": info["max_loan"], "max_shadow": info["max_shadow"], "max": max_impact(info),
            "affected": impact_text(info),
            "plancode_list": ", ".join(sorted(p for p in info["plancodes"] if p)),
        }
        records.append(record)
        rows.append([
            item["id"], item["area"], item["title"], item["description"], item["evidence"],
            record["affected"], record["plancodes"], record["policies"], record["shadow_policies"],
            record["max_av"] or None, record["max_charge"] or None, record["max_interest"] or None,
            record["max_loan"] or None, record["max_shadow"] or None, record["plancode_list"],
            item["source"], item["proof"], item["status"], item["fix"],
        ])
    row = write_title(ws, 1, "Fix List - confirmed issues (engine, plan config, rates, harness, CyberLife data)")
    row = write_note(ws, row, (
        "One row per confirmed issue. Counts/$ are from this run's classification (6-month replay, policy-months "
        "with > $1 attributed to the ID) and from the shadow evidence. 'Max' columns are the largest single-month "
        "attributed amount per measure (AV is cumulative within the window). Proof 'attributed' = cause assigned "
        "but not proven to the cent."), merge_to=12)
    write_table(ws, [
        Col("ID", width=6), Col("Area", width=14), Col("Issue", "wrap", 34), Col("Description", "wrap", 48),
        Col("Evidence summary", "wrap", 52), Col("Affected (this run)", "wrap", 18), Col("Plancodes", "int", 9),
        Col("6-mo policies", "int", 9), Col("Shadow policies", "int", 9), Col("Max AV $", "money", 11),
        Col("Max charge $", "money", 11), Col("Max interest $", "money", 11), Col("Max loan $", "money", 11),
        Col("Max shadow gap $", "money", 13), Col("Plancodes affected", "wrap", 26), Col("Source", "wrap", 22),
        Col("Proof", width=10), Col("Status", "wrap", 20), Col("Recommended fix", "wrap", 52),
    ], rows, start_row=row + 1)
    return records


def _requested_groups(ev: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for result in ev["results"]:
        groups[str(result.get("requested_plancode") or result.get("plancode") or "").upper()].append(result)
    return groups


def by_plancode_records(ev: dict[str, Any]) -> list[dict[str, Any]]:
    groups = _requested_groups(ev)
    selection = ev["selection"]
    shadow_by_plan: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in ev["shadow_rows"]:
        shadow_by_plan[str(row.get("plancode") or "").upper()].append(row)
    not_tested = Counter(str(row.get("requested_plancode") or "").upper() for row in selection.get("not_tested", []))
    records = []
    for req in ev["requested"]:
        code = req.requested
        info = (selection.get("by_requested") or {}).get(code, {})
        group = groups.get(code, [])
        compared = [r for r in group if r.get("status") == "compared"]
        errors = [r for r in group if r.get("status") != "compared"]
        keys = {r["key"] for r in compared}
        stats = measure_stats(ev["classified"], keys)
        months = sum(1 for k in keys for m in ev["classified"][k]["months"] if not m["seed"])
        ids = Counter()
        for k in keys:
            for rc_id in ev["classified"][k]["ids"]:
                ids[rc_id] += 1
        plans = {r.get("plancode") for r in group} | set(req.cyberlife_plancodes)
        shadow = [s for p in plans for s in shadow_by_plan.get(str(p).upper(), [])]
        sh_cmp = [s for s in shadow if s.get("diff") is not None]
        shadow_text = ""
        if shadow:
            shadow_text = (
                f"{len(sh_cmp)}/{len(shadow)} compared; as-is <= $1: {sum(abs(s['diff']) <= 1 for s in sh_cmp)}; "
                f"after rules <= $1: {sum(s.get('after_rules') is not None and abs(s['after_rules']) <= 1 for s in sh_cmp)}; "
                f"anchored <= $0.05: {sum(s.get('anchored_max') is not None and s['anchored_max'] <= 0.05 for s in shadow)}"
                f"/{sum(s.get('anchored_max') is not None for s in shadow)}"
            )
        candidates = info.get("candidate_count", 0)
        if not candidates and not group:
            status = "No in-force policies found"
        elif group and not compared:
            status = "Blocked: " + "; ".join(sorted({str(r.get('error') or r.get('blocker') or '')[:90] for r in errors}))
        elif errors:
            status = f"Replayed {len(compared)}/{len(group)}; {len(errors)} blocked"
        elif compared:
            status = "Replayed"
        else:
            status = "Not tested"
        if not_tested.get(code):
            status += f"; {not_tested[code]} candidate(s) failed selection"
        unexplained_months = sum(
            1 for k in keys for m in ev["classified"][k]["months"] if not m["seed"]
            and any(mm["class"] == "Unexplained" for mm in m["measures"].values())
        )
        shadow_ids = Counter(rc for s in shadow for rc in s["ids"])
        main_ids = [rc for rc, _ in (ids + shadow_ids).most_common() if rc not in ("RND",)][:6]
        records.append({
            "requested": code, "product": req.product_name, "cyberlife": ", ".join(sorted(p for p in plans if p)),
            "candidates": candidates, "selected": len(group), "replayed": len(compared), "blocked": len(errors),
            "months": months, "stats": stats, "unexplained_months": unexplained_months,
            "shadow": shadow_text, "status": status, "ids": main_ids, "note": req.resolution_note,
        })
    return records


def sheet_by_plancode(wb, ev: dict[str, Any]) -> None:
    ws = wb.create_sheet("By Plancode")
    rows = []
    for rec in ev["plancode_records"]:
        st = rec["stats"]
        rows.append([
            rec["requested"], rec["product"], rec["cyberlife"], rec["candidates"], rec["selected"], rec["replayed"],
            rec["blocked"], rec["months"],
            pct(st["av"], "Exact"), pct(st["av"], "Exact", "Rounding"), pct(st["coi"], "Exact"),
            pct(st["expense"], "Exact"), pct(st["other"], "Exact"), pct(st["total_charges"], "Exact"),
            pct(st["md"], "Exact"), pct(st["interest"], "Exact"), pct(st["interest"], "Exact", "Rounding"),
            pct(st["loan"], "Exact"), rec["unexplained_months"], rec["shadow"], rec["status"],
            ", ".join(f"{i} {ROOT_CAUSE_BY_ID.get(i, {}).get('title', UNEXPLAINED_LABELS.get(i, ''))[:40]}" for i in rec["ids"]),
        ])
    row = write_title(ws, 1, "By Plancode - every requested plancode")
    row = write_note(ws, row, (
        "% exact = share of compared non-seed policy-months within $0.01; '<= $1' adds Rounding. Candidates = "
        "in-force policies scanned for the plancode (capped at 12); up to 5 were selected for feature variety. Shadow: from-issue "
        "replay vs CyberLife XP (engine as-is, after identified rules) and roll-forward anchored on the June/July "
        "2026 RGA seriatim."), merge_to=14)
    write_table(ws, [
        Col("Requested", width=10), Col("Product", width=11), Col("CyberLife plancodes", "wrap", 16),
        Col("Candidates scanned (cap 12)", "int", 9), Col("Selected", "int", 8), Col("Replayed", "int", 8),
        Col("Blocked", "int", 8), Col("Months", "int", 8), Col("AV exact", "pct", 8), Col("AV <= $1", "pct", 8),
        Col("COI exact", "pct", 8), Col("Expense exact", "pct", 8), Col("Other exact", "pct", 8),
        Col("Total chg exact", "pct", 8), Col("MD exact", "pct", 8), Col("Interest exact", "pct", 8),
        Col("Interest <= $1", "pct", 8), Col("Loan exact", "pct", 8), Col("Months w/ unexplained", "int", 9),
        Col("Shadow", "wrap", 34), Col("Status", "wrap", 30), Col("Main root causes", "wrap", 60),
    ], rows, start_row=row + 1)


FEATURES = (
    ("DB A", lambda s, r: str(s.get("db_option") or "").upper() == "A"),
    ("DB B", lambda s, r: str(s.get("db_option") or "").upper() == "B"),
    ("DB other", lambda s, r: str(s.get("db_option") or "").upper() not in ("A", "B")),
    ("Table rated", lambda s, r: any("table" in t for t in s.get("substandard_tags") or [])
     or any((seg.get("table_rating") or 0) > 0 for seg in (r.get("features") or {}).get("segments", []))),
    ("Flat extra", lambda s, r: any("flat" in t for t in s.get("substandard_tags") or [])
     or any((seg.get("flat_extra") or 0) > 0 for seg in (r.get("features") or {}).get("segments", []))),
    ("Riders", lambda s, r: bool(s.get("rider_codes"))),
    ("Benefits", lambda s, r: bool(s.get("benefit_codes"))),
    ("Multi-segment", lambda s, r: "multi-segment" in (s.get("feature_tags") or [])),
    ("Loan balance", lambda s, r: (s.get("loan_balance") or 0) > 0),
    ("Loan tx in window", lambda s, r: any(t.get("kind") in ("loan", "loan_repayment") for t in r.get("transactions_replayed") or [])),
    ("Withdrawal in window", lambda s, r: bool(r.get("withdrawal_events"))),
    ("Premium in window", lambda s, r: any(t.get("kind") == "premium" for t in r.get("transactions_replayed") or [])),
    ("Shadow / CCV", lambda s, r: bool(s.get("shadow_applicable") or (r.get("features") or {}).get("ccv_active"))),
    ("IUL", lambda s, r: str(r.get("product_type") or s.get("product_type") or "").upper() == "IUL"),
    ("CVAT", lambda s, r: bool((r.get("features") or {}).get("is_cvat"))),
)


def sheet_feature_coverage(wb, ev: dict[str, Any]) -> None:
    ws = wb.create_sheet("Feature Coverage")
    selected = {f"{s['company']}_{s['policy_number']}": s for s in ev["selection"]["selected"]}
    groups = _requested_groups(ev)
    rows = []
    totals = Counter()
    rider_use: dict[str, set] = defaultdict(set)
    benefit_use: dict[str, set] = defaultdict(set)
    for req in ev["requested"]:
        group = [r for r in groups.get(req.requested, []) if r.get("status") == "compared"]
        counts = Counter()
        riders, benefits = Counter(), Counter()
        for result in group:
            sel = selected.get(result["key"], {})
            for name, test in FEATURES:
                try:
                    if test(sel, result):
                        counts[name] += 1
                except Exception:  # noqa: BLE001 - feature probes must not stop the workbook
                    pass
            for rider in sel.get("rider_codes") or []:
                riders[rider] += 1
                rider_use[rider].add(result["key"])
            for benefit in sel.get("benefit_codes") or []:
                benefits[benefit] += 1
                benefit_use[benefit].add(result["key"])
        totals.update(counts)
        totals["Replayed"] += len(group)
        rows.append([req.requested, req.product_name, len(group)] + [counts[name] for name, _ in FEATURES] + [
            ", ".join(f"{k} x{v}" for k, v in sorted(riders.items())),
            ", ".join(f"{k} x{v}" for k, v in sorted(benefits.items())),
        ])
    rows.append(["TOTAL", "", totals["Replayed"]] + [totals[name] for name, _ in FEATURES] + ["", ""])
    row = write_title(ws, 1, "Feature Coverage - replayed policies (6-month replay) by requested plancode")
    row = write_note(ws, row, (
        "Counts are replayed policies carrying each feature. Coverage is a 5-per-plancode feature-weighted sample, "
        "not a statistical sample: a feature with fewer than 3 policies is thin evidence (highlighted in the TOTAL "
        "row). Loan/withdrawal/premium 'in window' = cash flow inside the replayed months."), merge_to=14)
    columns = [Col("Requested", width=10), Col("Product", width=11), Col("Replayed", "int", 8)]
    columns += [Col(name, "int", 8) for name, _ in FEATURES]
    columns += [Col("Riders (plancode x policies)", "wrap", 36), Col("Benefits (type x policies)", "wrap", 30)]
    next_row = write_table(ws, columns, rows, start_row=row + 1)
    from openpyxl.styles import Font, PatternFill

    total_row = row + 1 + len(rows)
    for offset in range(len(columns)):
        cell = ws.cell(row=total_row, column=1 + offset)
        cell.font = Font(name="Calibri", size=9, bold=True)
        if 3 <= offset < 3 + len(FEATURES) and isinstance(cell.value, int) and cell.value < 3:
            cell.fill = PatternFill("solid", fgColor=CLASS_FILLS["Material"])
    next_row = write_title(ws, next_row + 1, "Rider plancodes and benefit types covered", level=2)
    detail = [["Rider", k, len(v), ", ".join(sorted(policy_label(x) for x in v))] for k, v in sorted(rider_use.items())]
    detail += [["Benefit", k, len(v), ", ".join(sorted(policy_label(x) for x in v))] for k, v in sorted(benefit_use.items())]
    write_table(ws, [Col("Kind", width=10), Col("Code", width=11), Col("Policies", "int", 8),
                     Col("Policies (company/policy)", "wrap", 90)], detail, start_row=next_row,
                freeze=False, autofilter=False)


def _worst(classes: list[str]) -> str:
    return max(classes, key=lambda c: WORST_ORDER.get(c, 0)) if classes else ""


def policy_records(ev: dict[str, Any]) -> list[dict[str, Any]]:
    selected = {f"{s['company']}_{s['policy_number']}": s for s in ev["selection"]["selected"]}
    out = []
    for result in ev["results"]:
        key = result["key"]
        sel = selected.get(key, {})
        item = ev["classified"].get(key)
        worst, maxdiff, final = {}, {}, {}
        ids = Counter()
        months = 0
        if item:
            body = [m for m in item["months"] if not m["seed"]]
            months = len(body)
            for meas in MEASURES:
                worst[meas] = _worst([m["measures"][meas]["class"] for m in body])
                diffs = [abs(m["measures"][meas]["diff"]) for m in body if m["measures"][meas]["diff"] is not None]
                maxdiff[meas] = max(diffs) if diffs else None
            final_rows = [m for m in body if m["measures"]["av"]["diff"] is not None]
            final["av"] = final_rows[-1]["measures"]["av"]["diff"] if final_rows else None
            ids = item["ids"]
        tx = Counter(t.get("kind") for t in result.get("transactions_replayed") or [])
        out.append({
            "key": key, "result": result, "sel": sel, "months": months, "worst": worst, "maxdiff": maxdiff,
            "final_av": final.get("av"), "ids": ids, "tx": tx,
        })
    return out


def sheet_policies(wb, ev: dict[str, Any]) -> None:
    ws = wb.create_sheet("Policies")
    rows = []
    for rec in ev["policy_records"]:
        result, sel, worst, maxdiff = rec["result"], rec["sel"], rec["worst"], rec["maxdiff"]
        m0 = result.get("month0") or {}
        rows.append([
            policy_label(rec["key"]), result.get("requested_plancode"), result.get("plancode"), result.get("product_name"),
            result.get("product_type") or sel.get("product_type"), result.get("status"),
            str(result.get("error") or result.get("blocker") or "")[:300],
            sel.get("issue_date"), sel.get("issue_state"), sel.get("issue_age"), sel.get("rate_sex"), sel.get("rate_class"),
            sel.get("db_option"), sel.get("face_amount"), sel.get("account_value"), sel.get("loan_balance"),
            ", ".join(sel.get("rider_codes") or []), ", ".join(sel.get("benefit_codes") or []),
            ", ".join(sel.get("substandard_tags") or []), ", ".join(sel.get("feature_tags") or []),
            result.get("rollback_date"), result.get("rollback_basis"), rec["months"],
            rec["tx"].get("premium", 0), rec["tx"].get("loan", 0) + rec["tx"].get("loan_repayment", 0),
            rec["tx"].get("withdrawal", 0), len(result.get("waiver_credit_in_window") or []),
            m0.get("variance"),
            worst.get("av"), maxdiff.get("av"), rec["final_av"], worst.get("coi"), worst.get("other"),
            worst.get("total_charges"), maxdiff.get("total_charges"), worst.get("md"), maxdiff.get("md"),
            worst.get("interest"), maxdiff.get("interest"), worst.get("loan"), maxdiff.get("loan"),
            ", ".join(k for k, _ in rec["ids"].most_common()),
        ])
    row = write_title(ws, 1, "Policies - every selected policy")
    row = write_note(ws, row, "Worst class over the replayed months per measure; max |diff| in dollars. PW = waiver "
                              "credits excluded from the replay (flag).", merge_to=12)
    write_table(ws, [
        Col("Policy", width=12), Col("Requested", width=10), Col("Plancode", width=10), Col("Product", width=10),
        Col("Type", width=5), Col("Status", width=9), Col("Error / blocker", "wrap", 40), Col("Issue date", width=10),
        Col("State", width=5), Col("Issue age", "int", 5), Col("Sex", width=4), Col("Class", width=5),
        Col("DB opt", width=5), Col("Face", "money", 12), Col("AV (current)", "money", 12), Col("Loan bal", "money", 10),
        Col("Riders", "wrap", 14), Col("Benefits", "wrap", 12), Col("Substandard", "wrap", 12),
        Col("Feature tags", "wrap", 22), Col("Rollback date", width=10), Col("Basis", width=16),
        Col("Months", "int", 6), Col("Premium tx", "int", 7), Col("Loan tx", "int", 6), Col("WD events", "int", 6),
        Col("PW flag", "int", 5), Col("Month-0 MD var", "money", 9),
        Col("AV worst", "class", 10), Col("AV max |diff|", "money", 10), Col("AV final diff", "money", 10),
        Col("COI worst", "class", 10), Col("Other worst", "class", 10), Col("Total chg worst", "class", 10),
        Col("Total chg max", "money", 9), Col("MD worst", "class", 10), Col("MD max", "money", 9),
        Col("Interest worst", "class", 10), Col("Interest max", "money", 9), Col("Loan worst", "class", 10),
        Col("Loan max", "money", 9), Col("Root-cause IDs", "wrap", 20),
    ], rows, start_row=row + 1)


# ── sheets: monthly detail, unexplained, root causes, shadow, month 0 ───────

def sheet_monthly_detail(wb, ev: dict[str, Any]) -> int:
    ws = wb.create_sheet("Monthly Detail")
    rows = []
    for key, item in ev["classified"].items():
        result = ev["by_key"][key]
        for month in item["months"]:
            row = [policy_label(key), result.get("plancode"), month["date"], "Seed" if month["seed"] else ""]
            for meas in MEASURES:
                m = month["measures"][meas]
                row += [m["cyberlife"], m["rerun"], m["diff"], m["class"], m["text"]]
            row.append("; ".join(month["flags"]))
            rows.append(row)
    columns = [Col("Policy", width=12), Col("Plancode", width=10), Col("MV date", width=10), Col("Seed", width=5)]
    for meas in MEASURES:
        label = MEASURE_LABELS[meas]
        columns += [Col(f"{label} CL", "money", 11), Col(f"{label} RERUN", "money", 11),
                    Col(f"{label} diff", "money", 9), Col(f"{label} class", "class", 10),
                    Col(f"{label} root cause", width=18)]
    columns.append(Col("Flags", "wrap", 40))
    row = write_title(ws, 1, "Monthly Detail - CyberLife vs RERUN per policy-month and measure")
    row = write_note(ws, row, (
        "CL = CyberLife LH_POL_MVRY_VAL (AV = CSV_AMT; COI = CINS_AMT; Expense = EXP_CRG_AMT; Other = OTH_PRM_AMT; "
        "Total charges = CINS+OTH vs COI+benefit+rider+asset; MD = recorded CD (or CINS+EXP+OTH); Interest = "
        "TOT_CRE_ITS_AMT; Loan = principal + accrued at the MV). Root cause = IDs with attributed $ (AV cumulative "
        "in the window). Seed rows are the rollback starting point."), merge_to=20)
    write_table(ws, columns, rows, start_row=row + 1)
    ws.freeze_panes = ws.cell(row=row + 2, column=5)
    return len(rows)


def unexplained_records(ev: dict[str, Any]) -> list[list[Any]]:
    out = []
    open_notes = {
        "01_U0433681": "AV deep open item: loan principal (6,518.30) exceeds CSV; loaned-CV crediting when debt > CV "
                       "not reproduced (CyberLife ~3.10% vs model 3.00%).",
        "01_U0590177": "LTGUL policy with negative CyberLife AV (-5.09M); residual is interest on the negative "
                       "carried difference.",
        "01_UE219641": "E14 date drift: the 6-30 -> 9-30 change spans three CyberLife MVs without RERUN states; "
                       "9/30 MV also processed during the session (D04).",
    }
    for key, item in ev["classified"].items():
        result = ev["by_key"][key]
        for month in item["months"]:
            if month["seed"]:
                continue
            for meas in MEASURES:
                m = month["measures"][meas]
                if m["class"] not in ("Unexplained", "Material") and not (
                    m["class"] == "Not comparable" and "E14" in m["attr"] and meas == "av"
                ):
                    continue
                unexpl = sum(v for k, v in m["attr"].items() if _proof(k) == "unexplained")
                attributed = sum(v for k, v in m["attr"].items() if _proof(k) == "attributed")
                out.append([
                    policy_label(key), result.get("plancode"), month["date"], MEASURE_LABELS[meas], m["cyberlife"],
                    m["rerun"], m["diff"], m["class"], m["text"] or ", ".join(m["attr"]), unexpl or None,
                    attributed or None, open_notes.get(key, ""),
                ])
    for row in ev["shadow_rows"]:
        residual = row.get("after_rules") if row.get("after_rules") is not None else row.get("diff")
        if residual is None or abs(residual) <= ROUNDING_TOLERANCE:
            continue
        if "U-SHD" in row["ids"] or "D05" in row["ids"]:
            out.append([policy_label(row["key"]), row["plancode"], "valuation", "Shadow (XP)", row["xp"], row["rerun"],
                        row["diff"], "Unexplained" if "U-SHD" in row["ids"] else "Material", ", ".join(row["ids"]),
                        row.get("after_rules"), None, row["root_cause"][:300]])
    return out


def sheet_unexplained(wb, ev: dict[str, Any]) -> None:
    ws = wb.create_sheet("Unexplained")
    row = write_title(ws, 1, "Unexplained and Material (open) items - nothing hidden")
    row = write_note(ws, row, (
        "Every policy-month/measure classed Unexplained (> $1 not attributed to a validated root cause) or Material "
        "(> $1 attributed to a cause whose proof is open, e.g. H09), every month RERUN cannot compare (E14), and "
        "every shadow policy whose residual is not identified. 'Unexplained $' is the portion not attributed."),
        merge_to=10)
    write_table(ws, [
        Col("Policy", width=12), Col("Plancode", width=10), Col("MV date", width=10), Col("Measure", width=10),
        Col("CyberLife", "money", 13), Col("RERUN", "money", 13), Col("Diff", "money", 10), Col("Class", "class", 11),
        Col("Attribution", "wrap", 36), Col("Unexplained $ / best rule", "money", 12), Col("Open-cause $", "money", 10),
        Col("Note", "wrap", 60),
    ], ev["unexplained_rows"], start_row=row + 1)


def sheet_charge_root_causes(wb, ev: dict[str, Any]) -> None:
    ws = wb.create_sheet("Charge Root Causes")
    findings = ev.get("charge_findings") or {}
    row = write_title(ws, 1, "Charge root causes (charge triage, validated) and this run's charge classification")
    row = write_note(ws, row, (
        "Upper table: clusters from charge_triage/charge_findings.json (round 2, material > $1 in COI/OTHER/MD). "
        "Lower table: this run's material charge months by root-cause ID and measure."), merge_to=10)
    cluster_rows = []
    id_map = {
        "CyberLife monthly-deduction field": "D01/H07", "Future-dated coverage": "E04",
        "Historical rider/benefit OTHER": "H08", "Pure deduction split": "H06", "Rollback seed": "H05",
    }
    for cluster in findings.get("cluster_summary", []):
        rc_id = next((v for k, v in id_map.items() if str(cluster.get("root_cause", "")).startswith(k)), "")
        cluster_rows.append([
            rc_id, cluster.get("category"), cluster.get("root_cause"), cluster.get("policy_count"),
            cluster.get("record_count"), cluster.get("max_abs_coi_diff"), cluster.get("max_abs_other_diff"),
            cluster.get("max_abs_md_diff"), ", ".join(cluster.get("policies") or []), cluster.get("suggested_fix"),
        ])
    row = write_table(ws, [
        Col("ID", width=8), Col("Category", width=18), Col("Root cause", "wrap", 50), Col("Policies", "int", 8),
        Col("Rows", "int", 6), Col("Max COI", "money", 9), Col("Max OTHER", "money", 9), Col("Max MD", "money", 9),
        Col("Policies (company/policy)", "wrap", 40), Col("Recommended fix", "wrap", 50),
    ], cluster_rows, start_row=row + 1, autofilter=False)
    counts: dict[tuple[str, str], dict[str, Any]] = defaultdict(lambda: {"months": 0, "policies": set(), "max": 0.0})
    for key, item in ev["classified"].items():
        for month in item["months"]:
            if month["seed"]:
                continue
            for meas in CHARGE_MEASURES:
                m = month["measures"][meas]
                if m["class"] in ("Exact", "Rounding", "Seed", "Not comparable"):
                    continue
                for rc_id, amount in m["attr"].items():
                    if rc_id == "RND" or abs(amount) <= ROUNDING_TOLERANCE:
                        continue
                    info = counts[(rc_id, meas)]
                    info["months"] += 1
                    info["policies"].add(key)
                    info["max"] = max(info["max"], abs(amount))
    run_rows = [[rc_id, MEASURE_LABELS[meas], ROOT_CAUSE_BY_ID.get(rc_id, {}).get("title", UNEXPLAINED_LABELS.get(rc_id, rc_id)),
                 len(info["policies"]), info["months"], info["max"], ", ".join(sorted(policy_label(k) for k in info["policies"]))]
                for (rc_id, meas), info in sorted(counts.items())]
    row = write_title(ws, row + 1, "This run: material charge months by root cause", level=2)
    write_table(ws, [
        Col("ID", width=8), Col("Measure", width=12), Col("Root cause", "wrap", 50), Col("Policies", "int", 8),
        Col("Months", "int", 8), Col("Max |$|", "money", 10), Col("Policies (company/policy)", "wrap", 60),
    ], run_rows, start_row=row, freeze=False, autofilter=False)


def sheet_av_root_causes(wb, ev: dict[str, Any]) -> None:
    ws = wb.create_sheet("AV Root Causes")
    av = ev.get("av_findings") or {}
    row = write_title(ws, 1, "AV and interest root causes (AV deep investigation, validated) and this run's attribution")
    stats = (ev.get("decomposition_stats") or {})
    row = write_note(ws, row, (
        f"Decomposition of this run: {stats.get('months', 0):,} policy-months; CyberLife roll-forward identity "
        f"(CSV + TOT_CRE_ITS_AMT + NET premiums - CINS-EXP-OTH - withdrawals) closes in {stats.get('identity_exact', 0):,}; "
        f"CyberLife's documented crediting rule reproduces TOT_CRE_ITS_AMT in {stats.get('cyber_rule_matches_recorded', 0):,}; "
        f"months without cash flow: prior-span daily rule {stats.get('cv_only_exact:cyberlife_daily_prior_span', 0)}/"
        f"{stats.get('months_without_cash_flows', 0)} vs engine current-month days "
        f"{stats.get('cv_only_exact:engine_daily_current_month_days', 0)}."), merge_to=10)
    verdicts = [[v.get("claim"), v.get("verdict"), v.get("category")] for v in av.get("verdict_table", [])]
    row = write_table(ws, [Col("First-pass claim", "wrap", 50), Col("Verdict", "wrap", 70), Col("Category", "wrap", 30)],
                      verdicts, start_row=row + 1, autofilter=False)
    row = write_title(ws, row + 1, "Confirmed fix list from the AV investigation (mapped to Fix List IDs)", level=2)
    av_map = {1: "E01", 2: "E02", 3: "P03", 4: "P01", 5: "P02", 6: "E03", 7: "H01", 8: "H02", 9: "E11", 10: "R01",
              11: "R02", 12: "E05", 13: "E06", 14: "R03", 15: "H04"}
    fixes = [[av_map.get(f.get("rank"), ""), f.get("rank"), f.get("type"), f.get("issue"), f.get("where"), f.get("fix"),
              f.get("evidence")] for f in av.get("confirmed_fixes", [])]
    row = write_table(ws, [Col("ID", width=6), Col("#", "int", 4), Col("Type", width=16), Col("Issue", "wrap", 40),
                           Col("Where", "wrap", 40), Col("Fix", "wrap", 50), Col("Evidence", "wrap", 60)],
                      fixes, start_row=row, freeze=False, autofilter=False)
    row = write_title(ws, row + 1, "This run: policies with a material AV month (> $1), final cumulative attribution", level=2)
    per_policy = []
    for key, item in ev["classified"].items():
        body = [m for m in item["months"] if not m["seed"] and m["measures"]["av"]["diff"] is not None]
        if not body:
            continue
        diffs = [m["measures"]["av"]["diff"] for m in body]
        if max(abs(d) for d in diffs) <= ROUNDING_TOLERANCE:
            continue
        worst = _worst([m["measures"]["av"]["class"] for m in body])
        per_policy.append([
            policy_label(key), ev["by_key"][key].get("plancode"), diffs[-1], max(abs(d) for d in diffs), worst,
            _attr_text(item["av_cum"], minimum=0.05, limit=6),
        ])
    per_policy.sort(key=lambda r: -r[3])
    write_table(ws, [Col("Policy", width=12), Col("Plancode", width=10), Col("Final AV diff", "money", 11),
                     Col("Max |AV diff|", "money", 11), Col("Worst class", "class", 11),
                     Col("Cumulative attribution at the last month (ID $)", "wrap", 80)],
                per_policy, start_row=row, freeze=False, autofilter=False)


def sheet_shadow(wb, ev: dict[str, Any]) -> None:
    ws = wb.create_sheet("Shadow")
    stats = ev["shadow_stats"]
    row = write_title(ws, 1, "Shadow account (XP / CCV / policy protection) - from issue, seriatim-anchored, APS205")
    row = write_note(ws, row, (
        f"XP = LH_COV_TARGET.TAR_PRM_AMT (TAR_TYP_CD 'XP'), post-deduction / pre-interest at the valuation date. "
        f"From-issue replay compared with RERUN shadow_av (engine as-is) for {stats['compared']} policies "
        f"({stats['no_shadow_config']} on 1U146200 with no shadow config); within $1: {stats['asis_within_1']}/"
        f"{stats['compared_excl_1U146200']} as-is, {stats['after_rules_within_1']}/{stats['compared_excl_1U146200']} "
        f"after the identified rules ({stats['after_rules_within_5']} within $5). Roll-forward anchored on the June/July "
        f"2026 RGA seriatim: {stats['anchored_within_005']}/{stats['anchored']} within $0.05 (exception: "
        f"{', '.join(policy_label(k) for k in stats['anchored_exceptions']) or 'none'} = D04 same-day artifact). "
        f"VP/MS APS205 replay: LTGUL/LTGUL08/EXECUL {stats['aps205_model_cent']}/{stats['aps205_model']} to the "
        f"cent; FPUL2 (APS205 mechanics, no FPUL2 model in the library) {stats['aps205_fpul2_within_035']}/"
        f"{stats['aps205_fpul2']} within $0.35. CyberLife keeps no shadow history in DB2, so only the current XP (and the "
        f"seriatim snapshots) can be compared. The 6-month replay does not compare shadow."), merge_to=14)
    eff = stats["albert_effect_totals"]
    row = write_note(ws, row, (
        "Robert 9/30 (authoritative): (1) SGUL late-payment forgiveness - any premium paid during the month earns a "
        "full month of shadow interest from the previous monthliversary (MV 8/16, $100 paid 8/25 earns 8/16 -> 9/16, "
        "31 days). This confirms E12; low priority for illustrations (premiums on the monthliversary) but required "
        "for rollback, from-issue and dated mid-month premium replays. (2) Known CyberLife defect: PT/PF premiums never "
        "went into the shadow, so XP is LOWER than contractually correct on those policies (D07, SR in progress). "
        "(3) CTR offsets stay an open question for Robert/IT (D06)."), merge_to=14)
    ptpf_rows = [r for r in ev["shadow_rows"] if r.get("ptpf") is not None]
    with_ptpf = [r for r in ptpf_rows if r["ptpf"].get("ptpf_count")]
    row = write_note(ws, row, (
        f"PT/PF check (D07): FH_FIXED history read for {len(ptpf_rows)} of {len(ev['shadow_rows'])} shadow policies; "
        f"{len(with_ptpf)} have a PT or PF premium (premium codes seen: "
        f"{', '.join(sorted({c for r in ptpf_rows for c in r['ptpf'].get('all_premium_codes') or []}))}). Our from-issue "
        "replays include PT/PF (baseline_shadow_check PREMIUM_CODES), the anchored roll-forwards use the same engine "
        "inputs, and the 6-month AV harness includes PT/PF as premiums (CyberLife's AV does take them); no PT/PF row "
        "falls in any 6-month replay window. With no PT/PF in the sample, the CyberLife as-processed view (PT/PF "
        "excluded) and the contractually correct view (included) are identical for every policy below; the "
        "out-of-sample table at the bottom shows both views where PT/PF exist. Direction: excluding PT/PF makes "
        "CyberLife LOWER than a correct replay; the D06 CTR offsets have the opposite sign (our replay is lower than "
        "XP), so they are a different issue."), merge_to=14)
    row = write_note(ws, row, (
        f"Corroboration (Albert, SGUL suiteview_diff; verified here: his engine values equal ours on 48/48 runnable "
        f"SGUL policies, his decomposition closes to $0.00, his seriatim values equal ours 78/78, effect totals "
        f"re-summed): over {stats['albert_in_scope']} in-scope SGUL policies shadow_calc departs from CyberLife by "
        f"premium timing {eff['premium_timing']:+,.2f} (E12 late-payment forgiveness, the top shadow defect), DB "
        f"discount {eff['dbd']:+,.2f} (R04, SGUL15), interest {eff['interest']:+,.2f} (E07) and load rounding "
        f"{eff['load_rounding']:+,.2f} (E12); COI, rider, EPU/MFEE inputs agree. His reference model is within $1 of XP "
        f"on {stats['albert_ref_within_1']}/{stats['albert_in_scope']} (same count as our identified rules). No SGUL "
        "VP/MS model exists; specs: ZZTaskRepo/RGA-July-2026-inventory-questions/01_Source/Product-Specs."), merge_to=14)
    rows = []

    def shadow_class(value: float | None, ids: list[str]) -> str | None:
        if value is None:
            return None
        if abs(value) <= 0.01:
            return "Exact"
        if abs(value) <= ROUNDING_TOLERANCE:
            return "Rounding"
        if "U-SHD" in ids:
            return "Unexplained"
        if any(_proof(i) == "attributed" for i in ids):
            return "Material"
        return "Explained"

    for rec in sorted(ev["shadow_rows"], key=lambda r: (str(r["plancode"]), r["key"])):
        within = shadow_class(rec.get("diff"), rec["ids"])
        after = shadow_class(rec.get("after_rules"), rec["ids"])
        alb = rec.get("albert") or {}
        pt = rec.get("ptpf") or {}
        has_pt = bool(pt.get("ptpf_count"))
        cl_view = pt.get("diff_without_ptpf") if has_pt else rec.get("diff")
        correct_view = pt.get("diff_with_ptpf") if has_pt else rec.get("diff")
        rows.append([
            policy_label(rec["key"]), rec["plancode"], rec["family"], rec["xp"], rec["rerun"], rec["diff"], within,
            rec.get("prior_diff"), rec.get("after_rules"), after, rec.get("rule_set"), rec.get("aps205"),
            rec.get("seriatim_jun"), rec.get("anchored_jun"), rec.get("seriatim_jul"), rec.get("anchored_jul"),
            alb.get("ref_minus_xp"), alb.get("premium_timing"), alb.get("dbd"), alb.get("interest"),
            alb.get("load_rounding"),
            ("not read" if not pt else pt.get("ptpf_count")), pt.get("ptpf_gross") if has_pt else None,
            pt.get("ptpf_worth_in_shadow") if has_pt else (0.0 if pt else None),
            "Yes (PREMIUM_CODES include PT/PF)" if pt else "", cl_view, correct_view,
            ", ".join(rec["ids"]), rec.get("primary"), rec.get("root_cause"), rec.get("source"),
        ])
    row = write_table(ws, [
        Col("Policy", width=12), Col("Plancode", width=10), Col("Family", width=9), Col("CyberLife XP", "money", 13),
        Col("RERUN shadow_av (as-is)", "money", 13), Col("Diff as-is", "money", 11), Col("As-is class", "class", 10),
        Col("Prior tool diff (shadow_eav)", "money", 11), Col("Diff after rules", "money", 11),
        Col("After-rules class", "class", 10), Col("Rule set", "wrap", 34), Col("APS205 replay diff", "money", 10),
        Col("Seriatim 6/30", "money", 12), Col("Anchored 6/30 -> XP", "money", 10), Col("Seriatim 7/31", "money", 12),
        Col("Anchored 7/31 -> XP", "money", 10), Col("Albert ref model - XP", "money", 10),
        Col("Albert: premium timing $", "money", 10), Col("Albert: DBD $", "money", 9),
        Col("Albert: interest $", "money", 9), Col("Albert: load rounding $", "money", 9),
        Col("PT/PF premiums (count)", width=9), Col("PT/PF gross $", "money", 10),
        Col("PT/PF worth in shadow $", "money", 10), Col("Our replay includes PT/PF", width=14),
        Col("Diff vs XP, CyberLife as-processed (PT/PF excl.)", "money", 13),
        Col("Diff vs XP, contractually correct (PT/PF incl.)", "money", 13),
        Col("Root-cause IDs", width=16), Col("Primary category", "wrap", 18),
        Col("Root cause / evidence", "wrap", 60), Col("Source", "wrap", 24),
    ], rows, start_row=row + 1)
    examples = ev.get("ptpf_examples") or []
    if examples:
        row = write_title(ws, row + 1, "PT/PF defect (D07) - out-of-sample examples replayed both ways", level=2)
        row = write_note(ws, row, (
            "Policies with PT/PF premiums (Albert's PT/PF examples), replayed here from issue with the engine twice. "
            "'As-processed' excludes PT/PF (reproduces CyberLife's defective XP up to the E12/R04 engine residuals); "
            "'contractually correct' includes them. The gap is what CyberLife's shadow is missing (CyberLife known "
            "defect, SR in progress) - not an engine issue."), merge_to=10)
        ex_rows = [[policy_label(e["key"]), e.get("plancode"), e.get("xp"), e.get("ptpf_count"), e.get("ptpf_gross"),
                    "; ".join(e.get("ptpf_dates") or []), e.get("engine_without_ptpf"), e.get("diff_without_ptpf"),
                    e.get("engine_with_ptpf"), e.get("diff_with_ptpf"), e.get("ptpf_worth_in_shadow")]
                   for e in examples]
        write_table(ws, [
            Col("Policy", width=12), Col("Plancode", width=10), Col("CyberLife XP", "money", 13),
            Col("PT/PF count", "int", 8), Col("PT/PF gross $", "money", 12), Col("PT/PF rows", "wrap", 34),
            Col("RERUN as-processed (excl.)", "money", 13), Col("Diff vs XP (excl.)", "money", 11),
            Col("RERUN correct (incl.)", "money", 13), Col("Diff vs XP (incl.)", "money", 11),
            Col("PT/PF worth in shadow $", "money", 12),
        ], ex_rows, start_row=row, freeze=False, autofilter=False)


def sheet_month0(wb, ev: dict[str, Any]) -> dict[str, int]:
    ws = wb.create_sheet("Month-0 MD Check")
    rows = []
    counts = Counter()
    for result in ev["results"]:
        m0 = result.get("month0") or {}
        var = _num(m0.get("variance"))
        cls = "Not comparable" if var is None else "Exact" if abs(var) <= 0.01 else "Rounding" if abs(var) <= 1 else (
            "Explained" if result.get("future_segments") else "Unexplained")
        counts[cls] += 1
        note = ""
        if result.get("future_segments"):
            note = "E04 future-dated coverage segment: " + ", ".join(
                f"phase {s.get('phase')} issue {s.get('issue_date')}" for s in result["future_segments"])
        elif var is None:
            note = str(result.get("error") or m0.get("note") or "")[:200]
        rows.append([
            policy_label(result["key"]), result.get("plancode"), result.get("status"), m0.get("system_md"),
            m0.get("calculated_md"), var, cls, m0.get("system_coi"), m0.get("calculated_coi"), m0.get("system_expense"),
            m0.get("calculated_expense"), m0.get("system_other"), m0.get("calculated_other"), note,
        ])
    row = write_title(ws, 1, "Month-0 MD check - RERUN md_check at the current valuation date, all policies")
    row = write_note(ws, row, (
        f"System MD = CyberLife current monthly deduction; Calculated = RERUN month-0 deduction from schema-rates. "
        f"Classes: {dict(counts)}. 'Not comparable' = rates missing (RateLookupError) so no calculation."), merge_to=10)
    write_table(ws, [
        Col("Policy", width=12), Col("Plancode", width=10), Col("Status", width=9), Col("System MD", "money", 11),
        Col("Calculated MD", "money", 11), Col("Variance", "money", 9), Col("Class", "class", 11),
        Col("System COI", "money", 10), Col("Calc COI", "money", 10), Col("System EXP", "money", 9),
        Col("Calc EXP", "money", 9), Col("System OTH", "money", 9), Col("Calc OTH", "money", 9), Col("Note", "wrap", 50),
    ], rows, start_row=row + 1)
    return dict(counts)


# ── sheets: rates, limitations, errors, method ──────────────────────────────

def read_rates_status(rates_repo: Path | None, rates_date: str, baseline_codes: set[str]) -> dict[str, Any]:
    """Read (read-only) the loader workbook rows loaded on ``rates_date`` and the open rate decisions."""
    out: dict[str, Any] = {"rows": [], "issues": [], "error": "", "workbook": "", "progress_tail": []}
    if rates_repo is None:
        return out
    workbook = rates_repo / "ProductDB_Manager.xlsx"
    out["workbook"] = str(workbook)
    try:
        from openpyxl import load_workbook

        wb = load_workbook(workbook, read_only=True, data_only=True)
        ws = wb["Plancodes"]
        header = None
        for values in ws.iter_rows(values_only=True):
            if header is None:
                if values and "Plancode" in [str(c or "") for c in values]:
                    header = [str(c or "") for c in values]
                continue
            rec = {h: ("" if v is None else str(v)) for h, v in zip(header, values)}
            code = rec.get("Plancode", "").strip().upper()
            if not rec.get("Load Date", "").startswith(rates_date):
                continue
            prod = rec.get("ProdType", "")
            if "UL" not in prod.upper() and code not in baseline_codes:
                continue
            detail = rec.get("Loader Detail", "").replace("\n", " ")
            dbo_parts = [part.strip() for part in re.split(r"\s\|\s", detail) if "dbo" in part.lower()]
            out["rows"].append({
                "plancode": code, "prod_type": prod, "status": rec.get("Status", ""), "load_date": rec.get("Load Date", ""),
                "loaded": rec.get("Rate Types Loaded", ""), "not_done": rec.get("Rate Types Not Done", ""),
                "dbo": "dbo-flagged" if dbo_parts else "IAF print",
                "dbo_note": " | ".join(dbo_parts)[:400],
                "detail": detail[-300:],
            })
        wb.close()
    except Exception as exc:  # noqa: BLE001 - the workbook may be locked by Excel; report, do not fail
        out["error"] = f"{type(exc).__name__}: {exc}"
    issues = rates_repo / "review" / "bulk-load-issues.md"
    if issues.exists():
        for line in issues.read_text(encoding="utf-8").splitlines():
            match = re.match(r"^\|\s*(3[7-9]|4[0-4])\s*\|\s*(.*?)\s*\|\s*(.*)\|\s*$", line)
            if match:
                out["issues"].append([match.group(1), match.group(2), match.group(3).strip()])
    progress = rates_repo / "review" / "bulk-loader-progress.md"
    if progress.exists():
        lines = progress.read_text(encoding="utf-8").splitlines()
        out["progress_tail"] = [ln for ln in lines if ln.startswith("### 2026-09-30") or ln.startswith("## 2026-09-30")]
    return out


def sheet_rates(wb, ev: dict[str, Any]) -> None:
    ws = wb.create_sheet("Rates Loaded 9-30")
    info = ev["rates_status"]
    row = write_title(ws, 1, "UL_Rates schema rates - plancodes loaded or reloaded on 9/30/2026 (UL family + baseline codes)")
    row = write_note(ws, row, (
        f"Source: {info.get('workbook')} (read-only), Plancodes sheet rows with Load Date {ev['args'].rates_date}. "
        "'dbo-flagged' = the loader detail records rows taken from legacy dbo tables (no IAF print) for Robert's "
        "review; everything else came from IAF prints. 'In baseline' counts selected policies using the plancode as "
        f"base or rider. Rate state for the final run: {ev['rate_state_text']}. "
        + (f"Workbook read error: {info['error']}" if info.get("error") else "")), merge_to=9)
    usage = Counter()
    for sel in ev["selection"]["selected"]:
        usage[str(sel.get("plancode") or "").upper()] += 1
        for rider in sel.get("rider_codes") or []:
            usage[str(rider).upper()] += 1
    rows = [[r["plancode"], r["prod_type"], r["status"], r["load_date"], r["dbo"], usage.get(r["plancode"], 0),
             r["loaded"], r["not_done"], r["dbo_note"], r["detail"]]
            for r in sorted(info["rows"], key=lambda r: (-usage.get(r["plancode"], 0), r["plancode"]))]
    row = write_table(ws, [
        Col("Plancode", width=10), Col("Type", width=8), Col("Status", width=8), Col("Load date", width=15),
        Col("Source", width=11), Col("In baseline", "int", 8), Col("Rate types loaded", "wrap", 55),
        Col("Not done", "wrap", 16), Col("dbo-sourced part (loader note)", "wrap", 50),
        Col("Loader detail (tail)", "wrap", 60),
    ], rows, start_row=row + 1)
    row = write_title(ws, row + 1, "Open rate decisions for Robert (review/bulk-load-issues.md #37-#44)", level=2)
    row = write_table(ws, [Col("#", width=5), Col("Item", "wrap", 40), Col("Question / decision", "wrap", 110)],
                      info["issues"], start_row=row, freeze=False, autofilter=False)
    row = write_title(ws, row + 1, "Loader log entries dated 9/30 (bulk-loader-progress.md headings)", level=2)
    write_table(ws, [Col("Entry", "wrap", 150)], [[ln.lstrip("# ")] for ln in info["progress_tail"]], start_row=row,
                freeze=False, autofilter=False)


def limitation_rows(ev: dict[str, Any]) -> list[list[Any]]:
    results = [r for r in ev["results"] if r.get("status") == "compared"]
    harness_basis = sum(1 for r in results if str(r.get("rollback_basis", "")).startswith("harness basis"))
    iul = sum(1 for r in results if r.get("iul_av_limitations"))
    null_zero = sum(1 for r in results for m in r.get("monthly") or [] if m.get("cyberlife_components_null_as_zero"))
    md_comp = sum(1 for r in results for m in r.get("monthly") or [] if m.get("cyberlife_md_from_components"))
    md0 = sum(1 for r in results for m in r.get("monthly") or [] if m.get("md_component_sum_substituted"))
    interest_nc = sum(1 for item in ev["classified"].values() for m in item["months"]
                      if not m["seed"] and m["measures"]["interest"]["class"] == "Not comparable")
    null_interest = [(r["key"], m["date"]) for r in results for m in r.get("monthly") or []
                     if not m.get("seed_row") and "interest" in m and m["interest"].get("cyberlife") is None]
    pw = sum(1 for r in results if r.get("waiver_credit_in_window"))
    return [
        ["Harness", "Rollback basis", f"{harness_basis}/{len(results)} replays start from the recorded LH_POL_MVRY_VAL "
         "row with CURRENT targets/accumulators/coverage (value_rollback blocks the canonical rollback). Target- or "
         "coverage-dependent effects (premium load split H09, historical riders H08) are limitations, not engine findings.",
         "H08, H09"],
        ["Harness", "Dated cash flows", "Premiums/withdrawals/loans are bucketed to the next monthliversary; receipt-date "
         "interest is not credited (E11). Attributed per month by the decomposition.", "E11"],
        ["Harness", "Seed row", "The first (rollback) month is CyberLife's own value and is excluded from all counts.", "H05"],
        ["Harness", "Null charge components", f"{null_zero} history rows had a null OTH_PRM_AMT/EXP_CRG_AMT read as 0.00 "
         "(CyberLife's roll-forward closes on that reading).", ""],
        ["Harness", "MD source", f"{md_comp} rows: MD = CINS+EXP+OTH because no CD row was recorded; {md0} rows: recorded "
         "CD = 0 with populated components (D01/H07).", "H07, D01"],
        ["Harness", "Interest not comparable", f"{interest_nc} non-seed months: TOT_CRE_ITS_AMT is null in CyberLife "
         f"({len(null_interest)} months on {len({k for k, _ in null_interest})} policies) or RERUN has no state on the "
         "date (E14).", ""],
        ["Harness", "PW waiver credits", f"{pw} replayed policies had a PW row in the window (excluded and flagged).", "H04"],
        ["Harness", "Shadow in the 6-month replay", "Shadow is excluded from the 6-month replay (CyberLife keeps no shadow "
         "history); shadow evidence is from-issue, seriatim-anchored and APS205 (Shadow sheet).", "H10"],
        ["Harness", "IUL", f"{iul} replays are IUL (engine product_type); the requested plancodes resolved to UL "
         "products, so IUL index-segment crediting (and the 9/30 index-rate rows) is not evidenced here.", ""],
        ["Data", "Sample", "Up to 5 in-force policies per requested plancode chosen for feature variety from <= 12 "
         "candidates; not statistically representative. Thin features are flagged on Feature Coverage.", ""],
        ["Data", "Window", "Six recorded monthliversaries (Mar-Sep 2026). No coverage changes, reinstatements or "
         "anniversary bonus crossings were exercised; E05/E06/E13/D02 have no $ in this window.", "E05, E06, E13, D02"],
        ["Data", "Shadow history", "No CyberLife shadow history exists in DB2; only today's XP and the June/July 2026 RGA "
         "seriatim snapshots can be compared.", "D03"],
        ["Data", "Same-day processing", "One policy's 9/30 monthliversary was processed while the data were read.", "D04"],
        ["Rates", "Schema rates only", "Every rate used by RERUN comes from UL_Rates schema rates (runtime audit on the "
         "Method & Sources sheet). dbo.CYBERLIFE_PDF (plan description metadata, not rates) is read only by the "
         "diagnostic decomposition for the compound-rule flag.", "R06"],
        ["Rates", "dbo-sourced rows", "Some 9/30 rows in schema rates were loaded from legacy dbo tables because no IAF "
         "print exists; they are flagged for Robert's review (Rates Loaded sheet).", "R06"],
        ["Code", "Working tree", f"Results depend on HEAD {ev['code_state'].get('head', '')[:7]} plus the uncommitted "
         "working tree (incl. the schema-rates loader ul_rates.py). A clean HEAD checkout would not reproduce them.", ""],
        ["Classification", "Attribution", "AV/interest differences are attributed with the per-month decomposition; "
         "charge differences driven by NAR are attributed to the AV difference entering the month (validated by "
         "RERUN NAR vs CyberLife NAR_AMT); loan differences to plan config / repayment / timing by policy facts.", ""],
    ]


def sheet_limitations(wb, ev: dict[str, Any]) -> None:
    ws = wb.create_sheet("Data & Harness Limits")
    row = write_title(ws, 1, "Data & Harness Limitations - what this evidence does and does not prove")
    write_table(ws, [Col("Area", width=14), Col("Topic", width=22), Col("Limitation", "wrap", 110),
                     Col("Related IDs", width=14)], limitation_rows(ev), start_row=row + 1)


def sheet_not_tested(wb, ev: dict[str, Any]) -> list[list[Any]]:
    ws = wb.create_sheet("Not Tested - Errors")
    rows = []
    for rec in ev["plancode_records"]:
        if rec["selected"] == 0:
            rows.append([rec["requested"], "", "", "Plancode not tested", rec["status"], ""])
    for item in ev["selection"].get("not_tested", []):
        rows.append([item.get("requested_plancode"), item.get("plancode"),
                     f"{item.get('company')}/{item.get('policy_number')}", "Selection failure",
                     "Candidate skipped at selection (UL_Rates query timeout during a rate load): "
                     + str(item.get("reason", ""))[:300], ""])
    for result in ev["results"]:
        if result.get("status") != "compared":
            error = str(result.get("error") or result.get("blocker") or "")
            rc_id = "R03" if "RateLookupError" in error else ""
            rows.append([result.get("requested_plancode"), result.get("plancode"), policy_label(result["key"]),
                         f"Replay {result.get('status')}", error[:400], rc_id])
    for key, item in ev["classified"].items():
        missing = [m["date"] for m in item["months"] if m.get("status")]
        if missing:
            rows.append([ev["by_key"][key].get("requested_plancode"), ev["by_key"][key].get("plancode"),
                         policy_label(key), "Months not comparable", "No RERUN state on CyberLife MV " + ", ".join(missing), "E14"])
    for rec in ev["shadow_rows"]:
        if rec.get("diff") is None:
            rows.append(["", rec["plancode"], policy_label(rec["key"]), "Shadow not compared", rec.get("root_cause", "")[:300],
                         ", ".join(rec["ids"])])
    row = write_title(ws, 1, "Not Tested / Errors - every selected policy or month that could not be compared")
    write_table(ws, [Col("Requested", width=10), Col("Plancode", width=10), Col("Policy", width=12), Col("What", width=20),
                     Col("Reason", "wrap", 100), Col("ID", width=8)], rows, start_row=row + 1)
    return rows


def sheet_method(wb, ev: dict[str, Any]) -> None:
    ws = wb.create_sheet("Method & Sources")
    args = ev["args"]
    audit = ev.get("rates_audit") or {}
    code = ev["code_state"]
    run = ev["final_run"]
    rows = [
        ["Purpose", "Baseline RERUN's UL/IUL engine and its UL_Rates schema `rates` against CyberLife's recorded "
                    "monthliversary history (AV, charges, interest, loans) and the shadow account (XP)."],
        ["Selection", f"{args.selection} (tools/rerun/baseline_select_policies.py; {args.plancodes})"],
        ["Replay", f"tools/rerun/baseline_history_compare.py --force --pdf-coding -> {args.results}"],
        ["Decomposition", "tools/rerun/baseline_av_decompose.py fit --work <run>/decomp (CyberLife rule CyberDoc B11 "
                          "pp.32-37; PDF DIFFCMPD compound flag)"],
        ["Workbook", "tools/rerun/baseline_workbook.py (this file's builder; classification + sheets)"],
        ["Shadow", f"{args.shadow_findings}; new checks: {args.shadow_extra or 'none'} (tools/rerun/baseline_shadow_check.py)"],
        ["Charge findings", str(args.charge_findings)],
        ["AV findings", str(args.av_findings)],
        ["CyberLife tables", "LH_POL_MVRY_VAL (CSV_AMT, CINS_AMT, EXP_CRG_AMT, OTH_PRM_AMT, TOT_CRE_ITS_AMT, "
                             "GUA_RT_ERN_ITS_AMT, NAR_AMT), FH_FIXED (PR/PA/PI/PF/PT/PB, LN*, PL/PP/..., SN/SM/SG/..., "
                             "CD, PW), LH_FND_VAL_LOAN, LH_COV_TARGET (XP) via PolicyInformation (CKPR, read-only)"],
        ["Rates", "UL_Rates schema `rates` only via the engine's rate_loader/ULRates (no dbo, no core.rates.Rates, no "
                  "Select_RATE_*)."],
        ["Rates audit", (f"{audit.get('policies_count', len(audit.get('policies', [])))} policies (one per requested "
                         f"plancode) replayed through compare_policy with pyodbc and core.rates.Rates instrumented: "
                         f"{audit.get('sql_statements', 'n/a')} SQL statements by DSN {audit.get('dsn_counts', {})}; "
                         f"{audit.get('schema_rates_statements', 'n/a')} UL_Rates statements name schema rates in the "
                         f"logged 300-character prefix and the rest are the rates.PLAN_DEF lookup in "
                         f"suiteview/core/rates_schema.py (FROM clause beyond the prefix); "
                         f"{len(audit.get('dbo_statements', [])) if audit else 'n/a'} statements against dbo/Select_RATE_*; "
                         f"{audit.get('core_rates_Rates_constructed', 'n/a')} core.rates.Rates objects. NEON_DSN = "
                         f"CyberLife DB2 (read-only). File: {args.rates_audit or 'not run'}")],
        ["Tolerances", f"Exact <= ${EXACT_TOLERANCE:.2f}; Rounding <= ${ROUNDING_TOLERANCE:.2f}; Material/Explained/"
                       f"Unexplained > ${ROUNDING_TOLERANCE:.2f}. Shadow anchored roll-forward reported at <= $0.05."],
        ["Classes", "Explained = attributed to validated IDs within $1; Material = > $1 attributed to an open-proof ID "
                    "(H09/R06/E13); Unexplained = > $1 not attributed (U-* IDs); Seed / Not comparable excluded."],
        ["Code state", f"HEAD {code.get('head')}; {code.get('modified_tracked_files')} modified tracked files, "
                       f"{code.get('untracked_entries')} untracked entries; tracked diff sha256 "
                       f"{code.get('tracked_diff_sha256')}; untracked (suiteview, tools/rerun) sha256 "
                       f"{code.get('untracked_sha256')}. Engine file hashes: "
                       + ", ".join(f"{Path(k).name}={v}" for k, v in (code.get("files") or {}).items())],
        ["Final run", f"started {run.get('started_at')}, finished {run.get('finished_at')}, {run.get('policy_count')} "
                      f"policies, status {run.get('status_counts')}, force={run.get('force')}"],
        ["Rate state", ev["rate_state_text"]],
        ["PT/PF check (D07)", f"final_work/ptpf_shadow.py (session helper; same build/run as baseline_shadow_check) -> "
                              f"{args.shadow_ptpf or 'not run'}; out-of-sample both-views replay -> "
                              f"{args.shadow_ptpf_examples or 'not run'}"],
        ["Robert 9/30 input", "Authoritative: E12 late-payment forgiveness confirmed (low priority for illustrations, "
                              "required for rollback/from-issue/dated replays); D07 PT/PF never in CyberLife shadow "
                              "(SR in progress); D06 CTR offsets remain an open question; R03 1U539600 IAF pull by Robert."],
        ["External corroboration", "Albert (unreviewed agent work) - cited only where verified: "
                                   "ZZTaskRepo/FFL-UL-6mo-History-Baseline-2026-09-30 (F01=E01, F02=E11, F03 in HEAD "
                                   "6e77a29, F05=E06, F06=E13, F07=H04, F08=H06, F12=D02) and "
                                   "ZZTaskRepo/SGUL-shadow-account-PT-PF-recalc (SGUL conventions, E07/E12); its "
                                   f"02_Working/Results/suiteview_diff ({args.shadow_albert or 'not used'}) verified "
                                   "here: engine values equal ours 48/48, closure $0.00, seriatim values equal ours "
                                   "78/78, rule totals re-summed (premium timing E12, DBD R04, interest E07, D06)."],
    ]
    row = write_title(ws, 1, "Method & Sources")
    write_table(ws, [Col("Item", width=22), Col("Detail", "wrap", 150)], rows, start_row=row + 1, autofilter=False)


# ── summary sheet, markdown, main ────────────────────────────────────────────

METHOD_BULLETS = [
    "Selected up to 5 in-force policies per requested plancode (62 codes) for variety: riders, benefits, table/flat "
    "substandard, DB options, loans, shadow, IUL.",
    "Started each policy at the earliest of its six recorded monthliversaries (CyberLife's own LH_POL_MVRY_VAL "
    "values; current targets) and replayed the recorded FH_FIXED cash flows through the RERUN engine "
    "(ProjectionTiming.CYBERLIFE_MONTHLIVERSARY), rates from UL_Rates schema `rates` only.",
    "Compared every monthliversary: AV, COI, expense, other, total charges, MD, interest credited "
    "(TOT_CRE_ITS_AMT) and loan balance (principal + accrued at the MV); seed month excluded.",
    "Decomposed each AV/interest difference month by month against CyberLife's documented crediting rule "
    "(roll-forward identity closes to the cent) and attributed every > $1 difference to a validated root-cause ID.",
    "Shadow: from-issue replay vs XP, CyberLife's VP/MS APS205 model replay, and roll-forward anchored on the "
    "June/July 2026 RGA seriatim shadow values.",
    "Month-0 MD check for every selected policy; rates loaded on 9/30 and dbo-sourced rows listed for review.",
]


def headline(ev: dict[str, Any]) -> dict[str, Any]:
    stats = measure_stats(ev["classified"])
    compared = [r for r in ev["results"] if r.get("status") == "compared"]
    months = sum(1 for item in ev["classified"].values() for m in item["months"] if not m["seed"])
    measures = {}
    for meas in MEASURES:
        c = stats[meas]
        measures[meas] = {
            "compared": sum(c[k] for k in CLASS_ORDER), "not_comparable": c["Not comparable"],
            **{k: c[k] for k in CLASS_ORDER},
            "pct_exact": pct(c, "Exact"), "pct_within_1": pct(c, "Exact", "Rounding"),
            "pct_explained_or_better": pct(c, "Exact", "Rounding", "Explained"),
        }
    six = [r for r in ev["unexplained_rows"] if r[3] != "Shadow (XP)"]
    net_of_timing = {}
    for meas in ("av", "interest"):
        exact = within = total = 0
        for item in ev["classified"].values():
            for month in item["months"]:
                m = month["measures"][meas]
                if month["seed"] or m["class"] not in CLASS_ORDER:
                    continue
                net = m["diff"] - sum(v for k, v in m["attr"].items() if k in ("E01", "P01", "E11"))
                total += 1
                exact += abs(net) <= EXACT_TOLERANCE + 0.005
                within += abs(net) <= ROUNDING_TOLERANCE
        net_of_timing[meas] = {"compared": total, "pct_exact": exact / total if total else None,
                               "pct_within_1": within / total if total else None}
    unexplained_policies = sorted({r[0] for r in six if r[7] == "Unexplained"})
    material_policies = sorted({r[0] for r in six if r[7] == "Material"})
    return {
        "policies_selected": len(ev["selection"]["selected"]),
        "selection_failures": len(ev["selection"].get("not_tested", [])),
        "requested_plancodes": len(ev["requested"]),
        "policies_replayed": len(compared),
        "policies_blocked": len(ev["results"]) - len(compared),
        "policy_months": months,
        "measures": measures,
        "net_of_timing": net_of_timing,
        "shadow": ev["shadow_stats"],
        "month0": ev.get("month0_counts", {}),
        "unexplained_policies": unexplained_policies,
        "material_policies": material_policies,
        "unexplained_rows": sum(1 for r in six if r[7] == "Unexplained"),
        "material_rows": sum(1 for r in six if r[7] == "Material"),
        "not_comparable_months": sum(1 for r in six if r[7] == "Not comparable"),
    }


def _kv(ws, row: int, label: str, value: Any, *, col: int = 1, span: int = 12) -> int:
    st = _style_objects()
    a = ws.cell(row=row, column=col, value=label)
    a.font = st["bold"]
    a.alignment = st["top"]
    b = ws.cell(row=row, column=col + 1, value=_excel_value(value))
    b.font = st["body_font"]
    b.alignment = st["wrap"]
    ws.merge_cells(start_row=row, start_column=col + 1, end_row=row, end_column=col + span)
    text = str(value or "")
    ws.row_dimensions[row].height = max(13, 12.5 * (1 + len(text) // 150))
    return row + 1


def sheet_summary(wb, ev: dict[str, Any]) -> None:
    ws = wb.active
    ws.title = "Summary"
    h = ev["headline"]
    code = ev["code_state"]
    row = write_title(ws, 1, "RERUN UL/IUL engine and UL_Rates schema `rates` vs CyberLife history - baseline evidence")
    row = _kv(ws, row + 1, "Purpose", (
        "Evidence that RERUN's UL/IUL engine, driven only by UL_Rates schema `rates`, reproduces CyberLife's recorded "
        "six-month monthliversary history (AV, charges, interest, loans) and shadow account for the requested "
        "plancodes, with every material difference traced to a root cause."))
    row = _kv(ws, row, "Date / for", f"{ev['generated_at']} - prepared for Robert Haessly (actuarial/product owner)")
    row = _kv(ws, row, "Code state", (
        f"git HEAD {code.get('head', '')[:7]} ({code.get('head', '')}) PLUS the uncommitted working tree "
        f"({code.get('modified_tracked_files')} modified tracked files, {code.get('untracked_entries')} untracked, incl. "
        "the schema-rates loader suiteview/illustration/core/ul_rates.py). A clean HEAD checkout would NOT reproduce "
        f"these numbers. Fingerprint: tracked diff {code.get('tracked_diff_sha256')}, untracked {code.get('untracked_sha256')} "
        "(engine file hashes on Method & Sources)."))
    row = _kv(ws, row, "Rate source / state", (
        "UL_Rates schema `rates` only (no dbo, no core.rates.Rates, no Select_RATE_*; runtime audit on Method & "
        f"Sources). {ev['rate_state_text']}"))
    row = _kv(ws, row, "Final run", (
        f"{ev['final_run'].get('started_at')} -> {ev['final_run'].get('finished_at')}; {ev['final_run'].get('policy_count')} "
        f"policies re-run with --force; status {ev['final_run'].get('status_counts')}"))
    row = write_title(ws, row + 1, "Method", level=2)
    for bullet in METHOD_BULLETS:
        row = _kv(ws, row, "-", bullet)
    row = _kv(ws, row, "Tolerances", (
        "Exact <= $0.01; Rounding <= $1.00; above $1.00 a difference is Explained (fully attributed to validated root "
        "causes), Material (attributed to a cause whose proof is open) or Unexplained (listed on the Unexplained sheet)."))
    row = write_title(ws, row + 1, "Headline numbers", level=2)
    row = _kv(ws, row, "Scope", (
        f"{h['requested_plancodes']} requested plancodes; {h['policies_selected']} policies selected "
        f"({h['selection_failures']} more candidate failed selection); {h['policies_replayed']} replayed, "
        f"{h['policies_blocked']} blocked; {h['policy_months']:,} replayed policy-months (seed months excluded)."))
    table = []
    for meas in MEASURES:
        m = h["measures"][meas]
        table.append([MEASURE_LABELS[meas], m["compared"], m["Exact"], m["Rounding"], m["Explained"], m["Material"],
                      m["Unexplained"], m["not_comparable"], m["pct_exact"], m["pct_within_1"], m["pct_explained_or_better"]])
    row = write_table(ws, [
        Col("Measure", width=16), Col("Compared months", "int", 11), Col("Exact", "int", 9), Col("Rounding", "int", 9),
        Col("Explained", "int", 9), Col("Material", "int", 9), Col("Unexplained", "int", 10),
        Col("Not comparable", "int", 10), Col("% exact", "pct", 9), Col("% within $1", "pct", 9),
        Col("% exact/rounding/explained", "pct", 12),
    ], table, start_row=row, freeze=False, autofilter=False)
    nt = h["net_of_timing"]
    row = _kv(ws, row, "Reading the table", (
        "Charges reproduce CyberLife to the cent (expense 100%; COI/OTHER differences are CyberLife booking the "
        "table/flat COI load in OTH, H06). AV and interest are rarely exact because of two timing rules: the engine's "
        "day-count shift (E01, alternating +/- one day of interest) and receipt-date interest on dated cash flows "
        f"(E11). Net of the dollars attributed to E01/P01/E11: AV {nt['av']['pct_exact']:.1%} exact and "
        f"{nt['av']['pct_within_1']:.1%} within $1; interest {nt['interest']['pct_exact']:.1%} exact and "
        f"{nt['interest']['pct_within_1']:.1%} within $1 (exact = within $0.015)."))
    sh = h["shadow"]
    row = _kv(ws, row, "Shadow (XP)", (
        f"From issue, engine as-is: {sh['asis_within_1']}/{sh['compared_excl_1U146200']} within $1; with the identified "
        f"rules: {sh['after_rules_within_1']}/{sh['compared_excl_1U146200']} within $1 ({sh['after_rules_within_5']} within "
        f"$5); 1U146200: {sh['no_shadow_config']} policies run on all-zero shadow rates (R05). Seriatim-anchored "
        f"roll-forward: {sh['anchored_within_005']}/{sh['anchored']} within $0.05. VP/MS APS205 replay: "
        f"LTGUL/LTGUL08/EXECUL {sh['aps205_model_cent']}/{sh['aps205_model']} to the cent, FPUL2 "
        f"{sh['aps205_fpul2_within_035']}/{sh['aps205_fpul2']} within $0.35. Historical offsets with current processing "
        f"reproduced (D05): {len(sh['historical_open'])}; residuals not identified: {len(sh['unexplained'])} policies."))
    m0 = h["month0"]
    row = _kv(ws, row, "Robert 9/30 rulings", (
        "E12 confirmed: SGUL shadow late-payment forgiveness (premium paid during the month earns a full month of "
        "shadow interest from the previous monthliversary) - low priority for illustrations, required for rollback/"
        "from-issue/dated mid-month replays. D07: PT/PF premiums never went into CyberLife's shadow (known defect, SR "
        "in progress) - XP is LOWER than contractually correct; no sampled shadow policy has PT/PF, so no residual here "
        "depends on it (4 out-of-sample policies show both views on the Shadow sheet). D06: pre-Dec-2019 CTR offsets "
        "stay an open question for Robert/IT (opposite sign to PT/PF). R03: 1U539600 rates pending Robert's IAF pull "
        "(~3 in force)."))
    row = _kv(ws, row, "Month-0 MD", f"All {sum(m0.values())} policies: {m0}")
    row = _kv(ws, row, "Unexplained / open", (
        f"{h['unexplained_rows']} measure-months Unexplained on {len(h['unexplained_policies'])} policies "
        f"({', '.join(h['unexplained_policies'])}); {h['material_rows']} measure-months Material (cause assigned, proof "
        f"open: H09 premium-load basis, D05 historical shadow offset) on {len(h['material_policies'])} policies. All are "
        "listed on the Unexplained sheet."))
    row = write_title(ws, row + 1, "Top findings (Fix List priority order; shadow: E12 premium timing first)", level=2)
    top = [r for r in ev["fix_records"] if r["area"].split()[0] in ("Engine", "Plan", "Rates")]
    row = write_table(ws, [
        Col("ID", width=6), Col("Area", width=16), Col("Issue", "wrap", 60), Col("Affected (this run)", "wrap", 26),
        Col("Max $ (any measure)", "money", 13), Col("Status", "wrap", 30),
    ], [[r["id"], r["area"], r["title"], r["affected"], r["max"] or None, r["status"]] for r in top],
        start_row=row, freeze=False, autofilter=False)
    row = write_title(ws, row + 1, "What is NOT proven", level=2)
    for text in [
        "Statistical representativeness: 5 policies per plancode chosen for feature variety, 6 months of history.",
        "Coverage changes, reinstatements, anniversary bonus crossings, waiver claims (PW) and MV day 29-31 schedules "
        "were barely or not exercised; E05/E06/E13/D02 have no $ in this window.",
        "Historical targets/accumulators: the replay keeps current targets (H09 premium-load residuals stay open).",
        "Shadow history before the June 2026 seriatim: from-issue gaps on SGUL15+CTR, U0437024 and 1U135400 rest on "
        "unreplayable CyberLife history, not proven engine behaviour.",
        "The PT/PF shadow defect (D07) is not exercised by the sample (no sampled shadow policy has PT/PF); its size "
        "rests on 4 out-of-sample replays and Albert's preliminary population figures.",
        "Illustration timing (ProjectionTiming.ILLUSTRATION) and guaranteed projections were not baselined here.",
        "IUL index crediting: every replayed policy is UL per the engine (no indexed policy among the requested "
        "plancodes), so index-segment crediting and the 9/30 index-rate rows are not evidenced.",
        "dbo-sourced rows loaded into schema rates on 9/30 (R06) are used as loaded; their correctness vs IAF prints "
        "awaits Robert's review.",
        "Results require the uncommitted working tree; they are not reproducible from a clean HEAD checkout.",
    ]:
        row = _kv(ws, row, "-", text)
    ws.column_dimensions["A"].width = 18
    ws.sheet_view.showGridLines = False


def write_markdown(path: Path, ev: dict[str, Any], workbook_path: Path) -> None:
    h = ev["headline"]
    sh = h["shadow"]
    code = ev["code_state"]
    lines = [
        "# RERUN vs CyberLife baseline - final summary (9/30/2026)",
        "",
        f"Workbook: `{workbook_path}`",
        "",
        f"**Scope.** {h['requested_plancodes']} requested plancodes, {h['policies_selected']} policies selected, "
        f"{h['policies_replayed']} replayed, {h['policies_blocked']} blocked, {h['policy_months']:,} policy-months "
        "(seed months excluded). Rates: UL_Rates schema `rates` only. Code: HEAD "
        f"{code.get('head', '')[:7]} + uncommitted working tree (a clean HEAD checkout would not reproduce). "
        f"Rate state: {ev['rate_state_text']}",
        "",
        "| Measure | Months | % exact | % within $1 | Explained | Material | Unexplained |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for meas in MEASURES:
        m = h["measures"][meas]
        lines.append(f"| {MEASURE_LABELS[meas]} | {m['compared']:,} | {(m['pct_exact'] or 0):.1%} | "
                     f"{(m['pct_within_1'] or 0):.1%} | {m['Explained']} | {m['Material']} | {m['Unexplained']} |")
    nt = h["net_of_timing"]
    lines += [
        "",
        f"Net of the dollars attributed to the two timing rules (E01 day-count shift, E11 receipt-date interest): AV "
        f"{nt['av']['pct_exact']:.1%} exact / {nt['av']['pct_within_1']:.1%} within $1; interest "
        f"{nt['interest']['pct_exact']:.1%} / {nt['interest']['pct_within_1']:.1%}.",
        "",
        f"**Shadow.** From issue as-is {sh['asis_within_1']}/{sh['compared_excl_1U146200']} within $1; with identified "
        f"rules {sh['after_rules_within_1']}/{sh['compared_excl_1U146200']}; seriatim-anchored roll-forward "
        f"{sh['anchored_within_005']}/{sh['anchored']} within $0.05; VP/MS APS205 LTGUL/EXECUL "
        f"{sh['aps205_model_cent']}/{sh['aps205_model']} to the cent, FPUL2 {sh['aps205_fpul2_within_035']}/"
        f"{sh['aps205_fpul2']} within $0.35.",
        "",
        "**Fix list (engine / plan config / rates).**",
        "",
        "| ID | Issue | Affected | Max $ | Status |",
        "|---|---|---|---:|---|",
    ]
    for r in ev["fix_records"]:
        if r["area"].split()[0] in ("Engine", "Plan", "Rates") or r["id"] in ("E11",):
            lines.append(f"| {r['id']} | {r['title']} | {r['affected']} | {r['max']:,.2f} | {r['status']} |")
    lines += [
        "",
        "**Robert 9/30.** E12 confirmed: SGUL shadow late-payment forgiveness (full month of shadow interest from the "
        "previous monthliversary) - low priority for illustrations, required for rollback/from-issue/dated mid-month "
        "replays. D07: PT/PF premiums never went into CyberLife's shadow (known defect, SR in progress; XP lower than "
        "correct) - none of the 80 sampled shadow policies has PT/PF, 4 out-of-sample policies show both views. D06: "
        "pre-Dec-2019 CTR offsets remain an open question for Robert/IT (opposite sign). R03: 1U539600 rates pending "
        "Robert's IAF pull (~3 in force).",
        "",
        "**Needs Robert's ruling / question.** " + "; ".join(
            f"{r['id']} {r['title']}" for r in ev["fix_records"]
            if "ruling" in r["status"].lower() or "question" in r["status"].lower()),
        "",
        f"**Unexplained.** {h['unexplained_rows']} measure-months on {', '.join(h['unexplained_policies']) or 'none'}; "
        f"{h['material_rows']} Material (cause assigned, proof open: H09 premium-load basis, D05 historical shadow "
        f"offset) on {', '.join(h['material_policies']) or 'none'}; shadow residuals "
        f"not identified: {', '.join(policy_label(k) for k in sh['unexplained']) or 'none'}. Details: Unexplained sheet.",
        "",
        "**Not proven.** Representativeness (5 policies x 6 months per plancode); coverage changes, bonus crossings, "
        "waiver claims; historical targets (H09); pre-seriatim shadow history; dbo-sourced 9/30 rate rows pending review.",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build(args) -> dict[str, Any]:
    from openpyxl import Workbook

    from tools.rerun.baseline_common import code_fingerprint, rate_state_note

    results, decomposition, manifest = load_run(Path(args.results))
    selection = json_load(Path(args.selection))
    ev: dict[str, Any] = {
        "args": args,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "selection": selection,
        "requested": resolve_requested_plancodes(Path(args.plancodes)),
        "results": results,
        "by_key": {r["key"]: r for r in results},
        "decomposition_stats": json_load(Path(args.results) / "av_decomposition.json").get("stats", {})
        if (Path(args.results) / "av_decomposition.json").exists() else {},
        "charge_findings": json_load(Path(args.charge_findings)) if args.charge_findings else {},
        "av_findings": json_load(Path(args.av_findings)) if args.av_findings else {},
        "rates_audit": json_load(Path(args.rates_audit)) if args.rates_audit and Path(args.rates_audit).exists() else {},
    }
    ev["classified"] = classify_run(results, decomposition)
    ev["shadow_rows"] = shadow_records(
        json_load(Path(args.shadow_findings)) if args.shadow_findings else None,
        json_load(Path(args.shadow_extra)) if args.shadow_extra else None,
        read_albert_sgul(Path(args.shadow_albert) if args.shadow_albert else None),
        json_load(Path(args.shadow_aps205)) if args.shadow_aps205 else None,
    )
    ptpf = {r["key"]: r for r in json_load(Path(args.shadow_ptpf))} if args.shadow_ptpf else {}
    for row in ev["shadow_rows"]:
        row["ptpf"] = ptpf.get(row["key"])
        if row["ptpf"] and row["ptpf"].get("ptpf_count") and "D07" not in row["ids"]:
            row["ids"].append("D07")
    ev["ptpf_examples"] = json_load(Path(args.shadow_ptpf_examples)) if args.shadow_ptpf_examples else []
    ev["shadow_stats"] = shadow_stats(ev["shadow_rows"])
    full_runs = [run for run in manifest.get("runs", []) if run.get("force") and not run.get("policy_keys")]
    ev["final_run"] = (full_runs or manifest.get("runs") or [{}])[-1]
    ev["code_state"] = ev["final_run"].get("code_state_end") or code_fingerprint()
    rate_state = ev["final_run"].get("rate_state_start") or rate_state_note()
    ev["rate_state_text"] = (
        (f"{args.rate_state_at}. " if args.rate_state_at else "")
        + f"Loader marker captured at run start {rate_state.get('captured_at')}: LOAD.lock present="
        f"{rate_state.get('load_lock_present')}; latest loader run {Path(rate_state.get('latest_run_report') or '').parent.name} "
        f"({rate_state.get('latest_run_mtime')})."
        + (f" {args.rate_note}" if args.rate_note else "")
    )
    partial = [run for run in manifest.get("runs", []) if run.get("force") and run.get("policy_keys")
               and str(run.get("started_at", "")) > str(ev["final_run"].get("started_at", ""))]
    if partial:
        ev["rate_state_text"] += " Follow-up forced re-runs after the full run: " + "; ".join(
            f"{run.get('started_at', '')[:19]} {run.get('policy_keys')}" for run in partial) + "."
    baseline_codes = {str(s.get("plancode") or "").upper() for s in selection["selected"]}
    baseline_codes |= {str(c).upper() for s in selection["selected"] for c in s.get("rider_codes") or []}
    ev["rates_status"] = read_rates_status(Path(args.rates_repo) if args.rates_repo else None, args.rates_date, baseline_codes)
    ev["impacts"] = id_impacts(ev)
    ev["plancode_records"] = by_plancode_records(ev)
    ev["policy_records"] = policy_records(ev)
    ev["unexplained_rows"] = unexplained_records(ev)

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    ev["fix_records"] = sheet_fix_list(wb, ev)
    sheet_by_plancode(wb, ev)
    sheet_feature_coverage(wb, ev)
    sheet_policies(wb, ev)
    ev["monthly_rows"] = sheet_monthly_detail(wb, ev)
    sheet_unexplained(wb, ev)
    sheet_charge_root_causes(wb, ev)
    sheet_av_root_causes(wb, ev)
    sheet_shadow(wb, ev)
    ev["month0_counts"] = sheet_month0(wb, ev)
    sheet_rates(wb, ev)
    sheet_limitations(wb, ev)
    ev["not_tested_rows"] = sheet_not_tested(wb, ev)
    sheet_method(wb, ev)
    ev["headline"] = headline(ev)
    sheet_summary(wb, ev)
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = {"Summary": "1F3864", "Fix List": "C00000", "Unexplained": "FF0000"}.get(ws.title, "8EA9DB")
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    write_markdown(Path(args.summary_md), ev, output)
    summary = {
        "generated_at": ev["generated_at"], "workbook": str(output), "headline": ev["headline"],
        "fix_list": ev["fix_records"], "final_run": ev["final_run"], "code_state": ev["code_state"],
        "rate_state": ev["rate_state_text"], "monthly_detail_rows": ev["monthly_rows"],
        "not_tested_rows": len(ev["not_tested_rows"]),
    }
    json_dump(Path(args.summary_json), summary)
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--selection", required=True)
    parser.add_argument("--results", required=True, help="baseline_history_compare.py output folder")
    parser.add_argument("--plancodes", required=True, help="requested plancodes TSV (Plancode, ProductName)")
    parser.add_argument("--charge-findings", default="")
    parser.add_argument("--av-findings", default="")
    parser.add_argument("--shadow-findings", default="")
    parser.add_argument("--shadow-extra", default="", help="shadow_check_results.json for newly checked policies")
    parser.add_argument("--shadow-albert", default="", help="corroborating SGUL shadow decomposition.csv (Albert)")
    parser.add_argument("--shadow-aps205", default="", help="JSON of APS205 replay results for newly checked policies")
    parser.add_argument("--shadow-ptpf", default="", help="ptpf_shadow.json: PT/PF premiums per shadow policy")
    parser.add_argument("--shadow-ptpf-examples", default="", help="ptpf JSON for out-of-sample PT/PF examples")
    parser.add_argument("--rates-repo", default="", help="Rates_Database folder (read-only)")
    parser.add_argument("--rates-date", default=date.today().isoformat())
    parser.add_argument("--rates-audit", default="", help="JSON from a runtime rates-source audit")
    parser.add_argument("--rate-note", default="", help="extra text for the rate-state line")
    parser.add_argument("--rate-state-at", default="", help="authoritative rate-state timestamp to report")
    parser.add_argument("--output", required=True)
    parser.add_argument("--summary-md", required=True)
    parser.add_argument("--summary-json", required=True)
    args = parser.parse_args()
    summary = build(args)
    print(json.dumps({"workbook": summary["workbook"], "headline": {
        k: v for k, v in summary["headline"].items() if k in ("policies_replayed", "policy_months", "unexplained_rows")
    }}, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
