"""Set individual Saved Cases cells in a RERUN workbook via Excel COM.

Targeted single-cell edits of an existing case column (e.g. correcting one
input after the fact) — rerun_build_case_inputs.py rewrites whole columns and
appends a NEW case, which is wrong for in-place fixes.

Usage (single JSON arg):
    venv\\Scripts\\python.exe tools/set_saved_case_value.py '<json>'

    {"workbook": "docs/Illustration_UL/RERUN (v20.0) local.xlsm",
     "case": 30,                       # case number or CaseID string
     "values": {"sINPUT_Variable_Loan_Rate": 0.057}}

Vector rows (names repeated over 121 rows) are not supported — only the FIRST
row bearing each name is written; the tool errors if the name is a vector.
"""
from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

ROOT = Path(__file__).resolve().parents[1]


def main() -> None:
    cmd = json.loads(sys.argv[1])
    workbook = (ROOT / (cmd.get("workbook")
                        or "docs/Illustration_UL/RERUN (v20.0) local.xlsm")).resolve() \
        if not Path(cmd.get("workbook", "")).is_absolute() else Path(cmd["workbook"])
    values: dict = cmd["values"]

    import openpyxl
    from rerun_com import _open_excel, _resolve_case_column, XL_CALC_MANUAL

    rwb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ws = rwb["Saved Cases"]
    col = _resolve_case_column(ws, cmd["case"])

    name_counts = Counter()
    name_row: dict[str, int] = {}
    old: dict[str, object] = {}
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if name is None:
            continue
        name = str(name).strip()
        name_counts[name] += 1
        if name in values and name not in name_row:
            name_row[name] = r
            old[name] = ws.cell(row=r, column=col).value
    rwb.close()

    missing = [n for n in values if n not in name_row]
    vectors = [n for n in values if name_counts.get(n, 0) > 1]
    if missing or vectors:
        print(json.dumps({"ok": False, "missing": missing, "vectors": vectors}))
        sys.exit(1)

    xl = _open_excel()
    try:
        wb = xl.Workbooks.Open(str(workbook), UpdateLinks=0, ReadOnly=False)
        xl.Calculation = XL_CALC_MANUAL
        sc = wb.Worksheets("Saved Cases")
        for name, val in values.items():
            sc.Cells(name_row[name], col).Value = val
        wb.Save()
        wb.Close(SaveChanges=False)
    finally:
        xl.Quit()

    print(json.dumps({"ok": True, "case": cmd["case"], "column": col,
                      "written": {n: {"old": old[n], "new": values[n]}
                                  for n in values}}, indent=2, default=str))


if __name__ == "__main__":
    main()
