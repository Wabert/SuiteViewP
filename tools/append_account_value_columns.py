r"""Append account-value + monthliversary-date columns to a GLP batch workbook.

Reads the records produced by ``fetch_last_mvry_av.py`` (JSON) and writes two new
columns immediately after the last populated header column of the batch sheet:

    <last+1>  "Account Value"        numeric, 2dp  (last-monthliversary CSV_AMT)
    <last+2>  "Monthliversary Date"  date          (that value's MVRY_DT)

Rows are matched by (Company col A, Policy col B), stripping whitespace exactly
like the fetch tool does (Policy cells are space-padded). Policies with no
account value are left blank. A timestamped backup of the workbook is written
before saving.

Usage:
    venv\Scripts\python.exe tools/append_account_value_columns.py "<workbook>" "<json>"
        [--sheet Batch] [--write]

Without --write it runs a dry run (reports what it would do, saves nothing).
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from copy import copy
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print(json.dumps({"error": "openpyxl not installed"}))
    sys.exit(1)


AV_HEADER = "Account Value"
DATE_HEADER = "Monthliversary Date"


def _key(company, policy):
    comp = str(company).strip() if company is not None else None
    pol = str(policy).strip() if policy is not None else None
    return (comp, pol)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("json_path")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--write", action="store_true")
    args = ap.parse_args()

    with open(args.json_path, "r", encoding="utf-8") as fh:
        payload = json.load(fh)

    av_by_key = {}
    date_by_key = {}
    for rec in payload["records"]:
        k = _key(rec.get("company"), rec.get("policy"))
        av = rec.get("account_value")
        dt = rec.get("monthliversary_date")
        if av is not None:
            av_by_key[k] = float(av)
        if dt:
            date_by_key[k] = datetime.strptime(dt[:10], "%Y-%m-%d")

    wb = openpyxl.load_workbook(args.workbook)
    ws = wb[args.sheet] if args.sheet else wb[wb.sheetnames[0]]

    # Locate the last populated header column and the Valuation Date column
    # (to copy its date number format).
    last_col = 0
    val_date_col = None
    header_template = None
    for cell in ws[1]:
        if cell.value is not None:
            last_col = cell.column
            header_template = cell
            if str(cell.value).strip() == "Valuation Date":
                val_date_col = cell.column

    date_number_format = "yyyy-mm-dd"
    if val_date_col is not None:
        sample = ws.cell(row=2, column=val_date_col)
        if sample.number_format and sample.number_format != "General":
            date_number_format = sample.number_format

    av_col = last_col + 1
    date_col = last_col + 2

    # Write headers, copying the style of an existing header cell.
    av_hdr = ws.cell(row=1, column=av_col, value=AV_HEADER)
    date_hdr = ws.cell(row=1, column=date_col, value=DATE_HEADER)
    if header_template is not None:
        for dst in (av_hdr, date_hdr):
            dst.font = copy(header_template.font)
            dst.fill = copy(header_template.fill)
            dst.alignment = copy(header_template.alignment)
            dst.border = copy(header_template.border)
            dst.number_format = "General"

    av_written = 0
    date_written = 0
    blank = 0
    for r in range(2, ws.max_row + 1):
        policy = ws.cell(row=r, column=2).value  # B
        if policy is None or not str(policy).strip():
            continue
        company = ws.cell(row=r, column=1).value  # A
        k = _key(company, policy)
        av = av_by_key.get(k)
        dt = date_by_key.get(k)
        if av is None and dt is None:
            blank += 1
            continue
        if av is not None:
            c = ws.cell(row=r, column=av_col, value=round(av, 2))
            c.number_format = "#,##0.00"
            av_written += 1
        if dt is not None:
            c = ws.cell(row=r, column=date_col, value=dt)
            c.number_format = date_number_format
            date_written += 1

    result = {
        "workbook": args.workbook,
        "sheet": ws.title,
        "av_header": AV_HEADER,
        "av_column": openpyxl.utils.get_column_letter(av_col),
        "date_header": DATE_HEADER,
        "date_column": openpyxl.utils.get_column_letter(date_col),
        "date_number_format": date_number_format,
        "rows_av_written": av_written,
        "rows_date_written": date_written,
        "rows_blank": blank,
        "written": bool(args.write),
    }

    if args.write:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup = f"{args.workbook}.bak-{stamp}"
        shutil.copy2(args.workbook, backup)
        result["backup"] = backup
        wb.save(args.workbook)
    else:
        result["note"] = "dry run - nothing saved; pass --write to persist"

    wb.close()
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
