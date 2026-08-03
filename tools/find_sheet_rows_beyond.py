"""Report the first non-empty rows beyond a row threshold on a big sheet.

Usage (single JSON arg):
    venv\\Scripts\\python.exe tools/find_sheet_rows_beyond.py '<json>'

    {"workbook": "<path>", "sheet": "Rates_Control",
     "min_row": 15492,          # first row to look at
     "limit": 15,               # stop after this many non-empty rows
     "max_cells": 12}           # populated cells to show per row

Answers "what is actually down there?" for sheets whose data sits in far
columns, without walking the whole grid. Prints JSON.
"""
from __future__ import annotations

import json
import sys

import openpyxl
from openpyxl.utils import get_column_letter


def main() -> None:
    cmd = json.loads(sys.argv[1])
    min_row = int(cmd.get("min_row", 1))
    limit = int(cmd.get("limit", 15))
    max_cells = int(cmd.get("max_cells", 12))

    wb = openpyxl.load_workbook(cmd["workbook"], data_only=True, read_only=True, keep_links=False)
    try:
        ws = wb[cmd["sheet"]]
        found = []
        scanned = 0
        for row in ws.iter_rows(min_row=min_row):
            scanned += 1
            cells = [c for c in row if c.value is not None and c.value != ""]
            if not cells:
                continue
            found.append({
                "row": cells[0].row,
                "populated": len(cells),
                "first_col": get_column_letter(cells[0].column),
                "last_col": get_column_letter(cells[-1].column),
                "sample": [
                    {f"{get_column_letter(c.column)}": str(c.value)[:40]}
                    for c in cells[:max_cells]
                ],
            })
            if len(found) >= limit:
                break
    finally:
        wb.close()

    print(json.dumps({
        "workbook": cmd["workbook"], "sheet": cmd["sheet"],
        "min_row": min_row, "rows_scanned": scanned,
        "non_empty_found": len(found), "rows": found,
    }, indent=2, default=str))


if __name__ == "__main__":
    main()
