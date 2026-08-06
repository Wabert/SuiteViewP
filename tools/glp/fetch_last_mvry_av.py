r"""Fetch each policy's account value as of its last monthliversary date.

Reads a GLP batch workbook (Company col A, Policy col B), resolves each policy's
``TCH_POL_ID`` from ``LH_BAS_POL``, then pulls the most-recent monthliversary row
from ``LH_POL_MVRY_VAL`` (ordered by ``MVRY_DT`` descending, ignoring the
9999 sentinel). The account value is that row's ``CSV_AMT`` and the effective
date is its ``MVRY_DT`` -- exactly the "Total AV" the PolView Advanced Product
tab shows via ``PolicyInformation.mv_av(0)`` / ``mv_date(0)``.

Only advanced/UL products have monthliversary rows; traditional or not-found
policies come back with ``account_value = None`` and ``monthliversary_date =
None``.

All lookups are batched with inline IN-lists (this DB2 ODBC driver rejects
parameter binding -- see policy_data.py), one query per chunk.

Usage:
    venv\Scripts\python.exe tools/glp/fetch_last_mvry_av.py "<workbook>" [--sheet Batch]
        [--region CKPR] [--limit N] [--out <path>]

Emits a JSON summary to stdout; writes the full record array to --out (or the
default path next to the workbook) as {"generated", "region", "source", "count",
"found", "with_av", "records":[{company, policy, tch_pol_id, account_value,
monthliversary_date, found}]}.
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, date
from decimal import Decimal

try:
    import openpyxl
except ImportError:
    print(json.dumps({"error": "openpyxl not installed"}))
    sys.exit(1)

sys.path.insert(0, ".")
from suiteview.core.db2_connection import DB2Connection  # noqa: E402


CHUNK = 400


def _norm_date(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = str(value).strip()
    return text or None


def read_workbook(path: str, sheet: str | None):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    records = []
    for row in ws.iter_rows(min_row=2):
        company = row[0].value if len(row) > 0 else None
        policy = row[1].value if len(row) > 1 else None
        if policy is None:
            continue
        policy_str = str(policy).strip()
        if not policy_str:
            continue
        records.append({
            "company": str(company).strip() if company is not None else None,
            "policy": policy_str,
        })
    wb.close()
    return records


def _chunks(seq, size):
    for i in range(0, len(seq), size):
        yield seq[i:i + size]


def _sql_literal(value: str) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def resolve_pol_ids(db: DB2Connection, by_company: dict, system_code: str = "I") -> dict:
    """Return {(company, policy): tch_pol_id} for policies found in LH_BAS_POL."""
    resolved = {}
    for company, policies in by_company.items():
        for chunk in _chunks(policies, CHUNK):
            in_list = ",".join(_sql_literal(p) for p in chunk)
            sql = (
                "SELECT CK_CMP_CD, CK_POLICY_NBR, TCH_POL_ID "
                "FROM DB2TAB.LH_BAS_POL "
                f"WHERE CK_SYS_CD = {_sql_literal(system_code)} "
                f"AND CK_CMP_CD = {_sql_literal(company)} "
                f"AND CK_POLICY_NBR IN ({in_list})"
            )
            _, rows = db.execute_query_with_headers(sql)
            for cmp_cd, pol_nbr, tch in rows:
                resolved[(str(cmp_cd).strip(), str(pol_nbr).strip())] = str(tch)
    return resolved


def fetch_last_mvry(db: DB2Connection, tch_ids: list) -> dict:
    """Return {tch_pol_id: (mvry_dt_iso, csv_amt_decimal)} for the latest
    (non-9999) monthliversary row of each policy."""
    latest: dict = {}
    for chunk in _chunks(tch_ids, CHUNK):
        in_list = ",".join(_sql_literal(t) for t in chunk)
        # Correlated MAX subquery -> one row per policy (its last monthliversary).
        sql = (
            "SELECT m.TCH_POL_ID, m.MVRY_DT, m.CSV_AMT "
            "FROM DB2TAB.LH_POL_MVRY_VAL m "
            f"WHERE m.TCH_POL_ID IN ({in_list}) "
            "AND m.MVRY_DT = ("
            "  SELECT MAX(m2.MVRY_DT) FROM DB2TAB.LH_POL_MVRY_VAL m2 "
            "  WHERE m2.TCH_POL_ID = m.TCH_POL_ID AND YEAR(m2.MVRY_DT) < 9999"
            ")"
        )
        _, rows = db.execute_query_with_headers(sql)
        for tch, mvry_dt, csv_amt in rows:
            tch_key = str(tch)
            dt_iso = _norm_date(mvry_dt)
            av = Decimal(str(csv_amt)) if csv_amt is not None else None
            prev = latest.get(tch_key)
            # Keep the newest date if the driver ever returns ties.
            if prev is None or (dt_iso or "") >= (prev[0] or ""):
                latest[tch_key] = (dt_iso, av)
    return latest


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--region", default="CKPR")
    ap.add_argument("--limit", type=int, default=None,
                    help="only process the first N rows (probe mode)")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    records = read_workbook(args.workbook, args.sheet)
    if args.limit:
        records = records[:args.limit]

    by_company: dict = {}
    for r in records:
        by_company.setdefault(r["company"], set()).add(r["policy"])
    by_company = {c: sorted(p) for c, p in by_company.items()}

    db = DB2Connection(region=args.region)
    resolved = resolve_pol_ids(db, by_company)

    tch_ids = sorted(set(resolved.values()))
    latest = fetch_last_mvry(db, tch_ids)

    out_records = []
    found_n = 0
    with_av = 0
    for r in records:
        key = (r["company"], r["policy"])
        tch = resolved.get(key)
        if tch is None:
            out_records.append({
                **r, "tch_pol_id": None,
                "account_value": None, "monthliversary_date": None, "found": False,
            })
            continue
        found_n += 1
        mv = latest.get(tch)
        if mv is None or mv[1] is None:
            out_records.append({
                **r, "tch_pol_id": tch,
                "account_value": None, "monthliversary_date": mv[0] if mv else None,
                "found": True,
            })
            continue
        with_av += 1
        out_records.append({
            **r, "tch_pol_id": tch,
            "account_value": f"{mv[1]:.2f}",
            "monthliversary_date": mv[0],
            "found": True,
        })

    payload = {
        "generated": datetime.now().isoformat(timespec="seconds"),
        "region": args.region,
        "source": args.workbook,
        "count": len(out_records),
        "found": found_n,
        "with_av": with_av,
        "records": out_records,
    }

    out_path = args.out
    if out_path is None and not args.limit:
        out_path = args.workbook + ".account_value.json"
    if out_path:
        with open(out_path, "w", encoding="utf-8") as fh:
            json.dump(payload, fh, indent=2, ensure_ascii=False)

    summary = {k: payload[k] for k in ("generated", "region", "count", "found", "with_av")}
    summary["out"] = out_path
    summary["sample"] = out_records[:8]
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
