r"""Append an 'Issue State' column to the GLP batch workbook.

Reads a GLP batch workbook (Company col A, Policy col B), batch-resolves each
policy's issue state code (``POL_ISS_ST_CD``) from ``LH_BAS_POL`` grouped by
company, translates the numeric code to a state abbreviation, then writes the
value into a new column immediately after the last populated header column of
the sheet, matching rows by (Company col A, Policy col B).

All lookups are batched with inline IN-lists (this DB2 ODBC driver rejects
parameter binding, matching the pattern in ``fetch_policy_debt.py``).

Usage:
    venv\Scripts\python.exe tools/glp/append_issue_state_column.py "<workbook>"
        [--sheet Batch] [--region CKPR] [--system-code I]

Emits a JSON summary to stdout.
"""
from __future__ import annotations

import argparse
import json
import sys

try:
    import openpyxl
except ImportError:
    print(json.dumps({"error": "openpyxl not installed"}))
    sys.exit(1)

# Make the suiteview package importable when run from repo root.
sys.path.insert(0, ".")
from suiteview.core.db2_connection import DB2Connection  # noqa: E402
from suiteview.polview.models.cl_polrec.policy_translations import (  # noqa: E402
    translate_state_code,
)


HEADER = "Issue State"
CHUNK = 400


def _sql_literal(value: str) -> str:
    """Single-quote a string value for inline SQL (this DB2 ODBC driver rejects
    parameter binding, so all callers inline literals — see fetch_policy_debt.py)."""
    return "'" + str(value).replace("'", "''") + "'"


def _chunks(seq, size):
    for i in range(0, len(seq), size):
        yield seq[i:i + size]


def _translate(code) -> str:
    if code is None:
        return ""
    text = str(code).strip()
    if text.isdigit():
        return translate_state_code(int(text))
    return text


def resolve_states(db: DB2Connection, by_company: dict, system_code: str) -> dict:
    """Return {(company, policy): issue_state_abbr} from LH_BAS_POL."""
    resolved = {}
    for company, policies in by_company.items():
        for chunk in _chunks(policies, CHUNK):
            in_list = ",".join(_sql_literal(p) for p in chunk)
            sql = (
                "SELECT CK_CMP_CD, CK_POLICY_NBR, POL_ISS_ST_CD "
                "FROM DB2TAB.LH_BAS_POL "
                f"WHERE CK_SYS_CD = {_sql_literal(system_code)} "
                f"AND CK_CMP_CD = {_sql_literal(company)} "
                f"AND CK_POLICY_NBR IN ({in_list})"
            )
            _, rows = db.execute_query_with_headers(sql)
            for cmp_cd, pol_nbr, st_cd in rows:
                key = (str(cmp_cd).strip(), str(pol_nbr).strip())
                resolved[key] = _translate(st_cd)
    return resolved


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--region", default="CKPR")
    ap.add_argument("--system-code", default="I")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook)
    ws = wb[args.sheet] if args.sheet else wb[wb.sheetnames[0]]

    # Collect unique (company, policy) pairs, grouped by company.
    by_company: dict = {}
    row_keys = {}
    for r in range(2, ws.max_row + 1):
        policy = ws.cell(row=r, column=2).value  # B
        if policy is None or not str(policy).strip():
            continue
        company = ws.cell(row=r, column=1).value  # A
        policy_str = str(policy).strip()
        company_str = str(company).strip() if company is not None else None
        by_company.setdefault(company_str, set()).add(policy_str)
        row_keys[r] = (company_str, policy_str)
    by_company = {c: sorted(p) for c, p in by_company.items()}

    db = DB2Connection(region=args.region)
    resolved = resolve_states(db, by_company, args.system_code)

    # Find the last populated header column in row 1.
    last_col = 0
    for cell in ws[1]:
        if cell.value is not None:
            last_col = cell.column
    target_col = last_col + 1
    ws.cell(row=1, column=target_col, value=HEADER)

    written = 0
    blank = 0
    for r, key in row_keys.items():
        state = resolved.get(key)
        if not state:
            blank += 1
            continue
        ws.cell(row=r, column=target_col, value=state)
        written += 1

    wb.save(args.workbook)

    print(json.dumps({
        "workbook": args.workbook,
        "sheet": ws.title,
        "header": HEADER,
        "column_index": target_col,
        "column_letter": openpyxl.utils.get_column_letter(target_col),
        "policies": sum(len(p) for p in by_company.values()),
        "rows_written": written,
        "rows_blank": blank,
    }, indent=2))


if __name__ == "__main__":
    main()
