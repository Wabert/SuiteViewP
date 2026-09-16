"""Historical output remains labeled after leaving the main-window toolbar."""

import io
import os
from datetime import date
from unittest.mock import Mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook, load_workbook
from PyQt6.QtWidgets import QApplication, QTreeWidgetItem

from suiteview.illustration.models.case_store import (
    decode_policy_snapshot, encode_policy_snapshot,
)
from suiteview.illustration.models.policy_data import IllustrationPolicyData
from suiteview.illustration.ui.values_tab import (
    LEDGER_COLUMNS, _ExportLabeledOverview, _export_values_summary,
)

_APP = None


def _policy():
    return IllustrationPolicyData(
        policy_number="ROLLBACK-OUTPUT", company_code="01", plancode="1U135D00",
        issue_date=date(2019, 3, 1), issue_age=43,
        valuation_date=date(2026, 8, 1), rollback_date=date(2026, 8, 1),
        rollback_source_date=date(2026, 9, 1),
        rollback_limitations=["Specified amounts are manual assumptions."],
        account_value=9_500, face_amount=300_000,
    )


def test_active_provenance_round_trips_in_policy_snapshot():
    policy = _policy()
    decoded = decode_policy_snapshot(encode_policy_snapshot(policy))
    assert decoded == policy
    assert decoded.rollback_date == date(2026, 8, 1)
    assert decoded.rollback_limitations == policy.rollback_limitations


def test_current_manual_provenance_is_persistent_and_not_historical():
    from suiteview.illustration.core.report_builder import build_ul_report
    from suiteview.illustration.models.calc_state import MonthlyState
    from suiteview.illustration.ui.values_tab import _issue_export_rows

    policy = _policy()
    policy.policy_number = "CURRENT-OUTPUT"
    policy.rollback_date = policy.rollback_source_date = None
    policy.rollback_limitations = []
    policy.starting_basis_assumptions = [
        "Account value was entered manually as 0.00; source policy is unchanged."]
    policy.starting_account_value_is_manual = True
    policy.account_value = 0
    decoded = decode_policy_snapshot(encode_policy_snapshot(policy))
    assert decoded == policy
    report = build_ul_report(decoded, [MonthlyState(date=policy.valuation_date)])
    text = "\n".join(report.basis_lines)
    assert "CURRENT INFORCE BASIS" in text
    assert "entered manually as 0.00" in text
    assert "ROLLBACK" not in text
    assert "HISTORICAL" not in text
    rows = _issue_export_rows(decoded)
    assert ("Opening account value", 0) in rows
    assert ("Current valuation date", "2026-08-01") in rows
    assert "rollback" not in str(rows).lower()


def test_historical_manual_av_is_not_reported_as_recorded():
    from suiteview.illustration.core.report_builder import rollback_output_basis

    policy = _policy()
    policy.starting_account_value_is_manual = True
    policy.starting_basis_assumptions = ["Account value was entered manually as 0.00."]
    text = "\n".join(rollback_output_basis(policy))
    assert "MANUAL POST-DEDUCTION AV" in text
    assert "RECORDED POST-DEDUCTION AV" not in text
    assert "entered manually as 0.00" in text


def test_ledger_export_identifies_rollback_without_changing_ledger_columns(monkeypatch):
    from suiteview.core import excel_export

    global _APP
    _APP = QApplication.instance() or QApplication([])
    excel, workbook, basis_sheet = Mock(), Mock(), Mock()
    dump = Mock(return_value=(excel, workbook, basis_sheet))
    write = Mock()
    monkeypatch.setattr(excel_export, "dump_to_new_workbook", dump)
    monkeypatch.setattr(excel_export, "write_table", write)
    overview = _ExportLabeledOverview()
    try:
        overview.display(_policy(), [])
        row = ["value"] * len(LEDGER_COLUMNS)
        overview.ledger.addTopLevelItem(QTreeWidgetItem(row))
        overview._dump_ledger(False)
        metadata = str(dump.call_args.args[1])
        assert "ROLLBACK" in metadata
        assert "08/01/2026" in metadata
        assert "Specified amounts are manual assumptions." in metadata
        assert write.call_args.args[1] == LEDGER_COLUMNS
        assert write.call_args.args[2] == [row]
    finally:
        overview.deleteLater()


def test_summary_workbook_has_rollback_basis_sheet_and_filename(monkeypatch, tmp_path):
    captured = {}
    save = Workbook.save

    def capture_save(workbook, path):
        stream = io.BytesIO()
        save(workbook, stream)
        captured["bytes"] = stream.getvalue()

    monkeypatch.setattr(Workbook, "save", capture_save)
    path = _export_values_summary(_policy(), [], [], tmp_path)
    workbook = load_workbook(io.BytesIO(captured["bytes"]))
    try:
        assert path.name.endswith("-ROLLBACK-2026-08-01.xlsx")
        assert workbook.sheetnames[0] == "Illustration Basis"
        text = str(list(workbook["Illustration Basis"].values))
        assert "ROLLBACK" in text
        assert "08/01/2026" in text
        assert "Specified amounts are manual assumptions." in text
    finally:
        workbook.close()


def test_iul_report_and_export_explain_total_only_basis_without_current_buckets():
    from suiteview.illustration.core.report_builder import build_ul_report
    from suiteview.illustration.core.value_rollback import apply_value_rollback
    from suiteview.illustration.models.calc_state import MonthlyState
    from suiteview.illustration.ui.values_tab import _issue_export_rows
    from tests.test_value_rollback_data import _policy, _complete_snapshot, WHEN

    policy = _policy()
    policy.product_type = "IUL"
    policy.plancode = "1U145500"
    policy.fund_values = {"SW": 2_000, "M1": 18_000}
    policy.premium_allocations = {"M1": 1.0}
    policy.rollback_snapshots = [_complete_snapshot()]
    historical = apply_value_rollback(policy, WHEN)
    report = build_ul_report(historical, [MonthlyState(date=WHEN)])
    assert report.is_iul
    assert report.iul_fund_values == []
    assert "total account value only" in "\n".join(report.basis_lines)
    assert report.iul_allocations
    rows = _issue_export_rows(historical)
    assert "total account value only" in str(rows)
    assert ("Opening account value", 10_000) in rows
    assert ("Rollback accumulated MTP", 2_500) in rows
    assert ("Rollback accumulated GLP", 15_000) in rows
