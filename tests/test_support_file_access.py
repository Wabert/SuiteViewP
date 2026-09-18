"""Support writes reauthorize at the boundary; browsing stays available."""

import os
from types import SimpleNamespace
from unittest.mock import Mock

import openpyxl
import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PyQt6.QtCore import QMimeData, Qt
from PyQt6.QtWidgets import QApplication, QListWidgetItem, QMessageBox

from suiteview.core import access_control as access
from suiteview.abrquote.ui import output_panel as abr
from suiteview.polview.ui.tabs import policy_support_tab as pol
from suiteview.ui.widgets import mini_explorer as mini
from suiteview.utils import excel_template


@pytest.fixture(scope="module")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture
def permissions(monkeypatch):
    state = {"writable": True, "unavailable": False}

    def load():
        if state["unavailable"]:
            raise access.AccessUnavailableError("Cannot verify permissions")
        return access.EffectiveAccess(
            "TEST", "TEST", True, False, state["writable"]
        )

    monkeypatch.setattr(access, "has_developer_access", lambda: False)
    monkeypatch.setattr(access, "_load_access", load)
    access.clear_access_cache()
    access.get_access()
    yield state
    access.clear_access_cache()


@pytest.fixture
def messages(monkeypatch):
    warning = Mock()
    monkeypatch.setattr(QMessageBox, "warning", warning)
    monkeypatch.setattr(QMessageBox, "information", Mock())
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.Yes)
    return warning


@pytest.mark.parametrize("extension", [".txt", ".xltx", ".xltm", ".xlt"])
@pytest.mark.parametrize("unavailable", [False, True])
def test_scoped_copy_denied_before_destination_open(
    tmp_path, permissions, extension, unavailable
):
    source = tmp_path / ("source" + extension)
    source.write_bytes(b"source must not be read or changed")
    destination = tmp_path / "existing.xlsx"
    destination.write_bytes(b"keep existing destination")
    permissions.update(writable=False, unavailable=unavailable)
    error = access.AccessUnavailableError if unavailable else access.AccessDeniedError

    with pytest.raises(error):
        excel_template.copy_as_workbook(str(source), str(destination), support_files=True)

    assert destination.read_bytes() == b"keep existing destination"
    assert source.read_bytes() == b"source must not be read or changed"


@pytest.mark.parametrize("support_files", [False, True])
def test_copy_scope_and_template_conversion(tmp_path, permissions, support_files):
    source = tmp_path / "template.xltx"
    workbook = openpyxl.Workbook()
    workbook.template = True
    workbook.active["A1"] = "Preserved"
    workbook.save(source)
    permissions["writable"] = support_files
    destination = tmp_path / "workbook.xlsx"

    excel_template.copy_as_workbook(str(source), str(destination), support_files=support_files)

    with open(destination, "rb") as stream:
        result = openpyxl.load_workbook(stream)
        assert result.template is False
        assert result.active["A1"].value == "Preserved"
        result.close()


@pytest.fixture(params=["polview", "abr"])
def panel(request, qapp, tmp_path, permissions, monkeypatch, messages):
    if request.param == "polview":
        monkeypatch.setattr(pol, "_load_user_tasks", lambda: [])
        monkeypatch.setattr(pol, "_get_onedrive_dir", lambda: str(tmp_path))
        widget = pol.PolicySupportTab()
        widget._policy_support_folder_path = str(tmp_path / "policy")
    else:
        monkeypatch.setattr(abr.OutputPanel, "_get_tools_root_path", lambda self: str(tmp_path))
        monkeypatch.setattr(abr.OutputPanel, "_is_onedrive_synced", lambda self: True)
        widget = abr.OutputPanel()
        widget._policy_folder_path = str(tmp_path / "policy")
        monkeypatch.setattr(widget, "_update_ui_state", lambda: None)
    widget._policy = SimpleNamespace(policy_number="TEST123", company_code="01")
    yield widget
    widget.close()
    widget.deleteLater()


@pytest.mark.parametrize("writable", [False, True])
def test_policy_folder_creation(panel, tmp_path, permissions, writable, messages):
    permissions["writable"] = writable
    assert access.can_write_support_files() is True  # The UI snapshot may be stale.
    panel._on_create_policy_folder()
    assert (tmp_path / "policy").exists() is writable
    assert bool(messages.call_count) is not writable


def test_authorization_unavailable_blocks_folder_creation(panel, tmp_path, permissions, messages):
    permissions["unavailable"] = True
    panel._on_create_policy_folder()
    assert not (tmp_path / "policy").exists()
    assert messages.call_count == 1


@pytest.mark.parametrize("writable", [False, True])
def test_policy_file_drop(panel, tmp_path, permissions, writable, messages):
    source = tmp_path / "form.txt"
    source.write_text("form", encoding="utf-8")
    destination = tmp_path / "policy"
    destination.mkdir()
    panel._subfolder_explorer.set_root(str(destination))
    permissions["writable"] = writable

    panel._on_file_dropped(str(source))

    assert (destination / "TEST123 - form.txt").exists() is writable
    assert bool(messages.call_count) is not writable


@pytest.mark.parametrize("writable", [False, True])
@pytest.mark.parametrize("root_exists", [False, True])
def test_category_drop_authorizes_each_folder(
    tmp_path, permissions, writable, root_exists, messages
):
    root = tmp_path / "policy"
    if root_exists:
        root.mkdir()
    panel = SimpleNamespace(
        _policy_support_folder_path=str(root),
        _subfolder_explorer=Mock(current_path=lambda: str(root)),
        _create_folder_btn=Mock(),
        _policy=None,
        _status_label=Mock(),
    )
    permissions["writable"] = writable

    pol.PolicySupportTab._on_category_dropped(panel, "Review")

    assert (root / "Review").exists() is writable
    assert root.exists() is (root_exists or writable)


@pytest.mark.parametrize("kind", ["polview", "scoped", "generic"])
@pytest.mark.parametrize("operation", ["rename", "delete"])
@pytest.mark.parametrize("directory", [False, True])
@pytest.mark.parametrize("writable", [False, True])
def test_browser_mutations(
    qapp, tmp_path, permissions, messages, monkeypatch,
    kind, operation, directory, writable,
):
    path = tmp_path / "original"
    if directory:
        path.mkdir()
        (path / "child.txt").write_text("keep", encoding="utf-8")
    else:
        path.write_text("keep", encoding="utf-8")
    if kind == "polview":
        explorer = pol._SubfolderExplorer()
        explorer.set_root(str(tmp_path))
    else:
        explorer = mini.MiniExplorer(root_path=str(tmp_path), support_files=kind == "scoped")
    item = QListWidgetItem("original")
    item.setData(Qt.ItemDataRole.UserRole, str(path))
    monkeypatch.setattr(mini.QInputDialog, "getText", lambda *a, **k: ("renamed", True))
    permissions["writable"] = writable

    getattr(explorer, f"_{operation}_entry")(item)

    allowed = writable or kind == "generic"
    assert path.exists() is not allowed
    assert (tmp_path / "renamed").exists() is (allowed and operation == "rename")
    assert bool(messages.call_count) is not allowed
    explorer.close()
    explorer.deleteLater()


@pytest.mark.parametrize("kind", ["polview", "scoped", "generic"])
@pytest.mark.parametrize("writable", [False, True])
def test_drop_signal_requires_fresh_authorization(qapp, permissions, messages, kind, writable):
    if kind == "polview":
        widget = pol._DropTargetSubfolderList()
    else:
        widget = mini.DropTargetSubfolderList()
        widget.support_files = kind == "scoped"
    listener = Mock()
    widget.file_dropped.connect(listener)
    mime = QMimeData()
    mime.setData(mini._MIME_TOOL_FILE, b"source.txt")
    event = Mock(mimeData=lambda: mime)
    permissions["writable"] = writable

    widget.dropEvent(event)

    allowed = writable or kind == "generic"
    assert bool(listener.call_count) is allowed
    assert bool(event.ignore.call_count) is not allowed
    widget.deleteLater()


@pytest.mark.parametrize("writable", [False, True])
def test_glp_detail_workbook_save(tmp_path, permissions, messages, writable):
    workbook = openpyxl.Workbook()
    panel = SimpleNamespace(
        _has_glp_quote_to_export=lambda: True,
        _glp_exception_folder_path=lambda: str(tmp_path),
        _glp_exception_file_name=lambda: "details.xlsx",
        _build_glp_quote_workbook=lambda: workbook,
    )
    permissions["writable"] = writable

    pol.PolicySupportTab._print_glp_quote_to_folder(panel)

    assert (tmp_path / "details.xlsx").exists() is writable
    assert bool(messages.call_count) is not writable


def test_read_only_ui_leaves_browsing_enabled(panel, permissions):
    permissions["writable"] = False
    access.get_access(refresh=True)
    panel._refresh_support_access()

    assert not panel._create_folder_btn.isEnabled()
    assert not panel._support_access_note.isHidden()
    assert panel._subfolder_explorer.isEnabled()
    assert not panel._subfolder_explorer._list.acceptDrops()
    if isinstance(panel, abr.OutputPanel):
        assert not panel._copy_up_btn.isEnabled()
        assert not panel._copy_left_btn.isEnabled()
        assert panel._subfolder_explorer._support_files
        assert panel._tools_explorer._support_files
