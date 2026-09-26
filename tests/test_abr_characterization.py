"""Characterization coverage for ABR refactors.

The assertions in this file pin the existing dedicated ABR Quote behavior before
the core-service split.  The fixtures are synthetic and in-memory; no live DB2,
UL_Rates, Excel COM, or support-file writes are required.
"""

from __future__ import annotations

from copy import deepcopy
from dataclasses import asdict
from datetime import date
from hashlib import sha256
import sys
from types import SimpleNamespace

import openpyxl
import pytest

from suiteview.abrquote.automation import (
    QuoteRequest,
    _assessment_port,
    calculate_quote,
)
from suiteview.abrquote.models.abr_data import (
    ABRPolicyData,
    ABRQuoteResult,
    MedicalAssessment,
)
from suiteview.abrquote.models.abr_database import using_quote_database
from suiteview.abrquote.ui.calc_viewer import CalcViewerDialog
from suiteview.abrquote.ui.output_panel import OutputPanel


class Rates:
    def get_vbt_qx(self, *args):
        return 10.0

    def get_effective_interest_rate(self, month):
        return (month, 0.05)

    def get_term_rate(self, *args):
        return 2.0

    def get_term_rate_schedule(self, *args):
        return [2.0] * 82

    def get_band(self, *args):
        return "A"

    def get_policy_fee(self, *args):
        return 60.0

    def get_modal_factor(self, *args):
        return 1.0

    def get_modal_fee_factor(self, *args):
        return 1.0

    def get_admin_fee(self, *args):
        return 250.0

    def get_per_diem(self, *args):
        return (420.0, 153300.0)

    def get_prem_cease_age(self, *args):
        return 95


@pytest.fixture
def policy() -> ABRPolicyData:
    return ABRPolicyData(
        policy_number="SYNTHETIC",
        company="01",
        region="CKPR",
        product_type="TERM",
        plan_code="TEST",
        issue_state="TX",
        issue_date=date(2020, 1, 15),
        maturity_date=date(2055, 1, 15),
        face_amount=100000,
        issue_age=40,
        attained_age=46,
        maturity_age=75,
        sex="F",
        rate_sex="F",
        rate_class="N",
        policy_year=7,
        policy_month=8,
        reinsurers="(none)",
        billing_mode=1,
    )


def _quote_request(assessment: dict, options: dict | None = None) -> dict:
    return {
        "policy_number": "SYNTHETIC",
        "company_code": "01",
        "region": "CKPR",
        "quote_date": "2026-09-03",
        "effective_date": "2026-09-03",
        "assessment": assessment,
        "options": {"min_face_amount": 50000, **(options or {})},
        "input_provenance": {"source": "synthetic characterization"},
    }


@pytest.mark.parametrize(
    ("name", "assessment", "expected"),
    [
        (
            "dual 5y+10y",
            {
                "rider_type": "Critical",
                "five_year_survival": 0.55,
                "ten_year_survival": 0.25,
                "return_after_10yr": True,
            },
            {
                "use_five_year": True,
                "use_ten_year": True,
                "derived_table_rating_5yr": 57.65929818153381,
                "derived_table_rating_10yr": 79.92386817932129,
                "computed_survival_5yr": 0.5499999987102724,
                "computed_survival_10yr": 0.2500008884529677,
            },
        ),
        (
            "single 5y return",
            {
                "rider_type": "Chronic",
                "five_year_survival": 0.60,
                "return_after_5yr": True,
            },
            {
                "use_five_year": True,
                "use_return_5yr": True,
                "derived_table_rating": 49.139559268951416,
                "computed_survival_5yr": 0.5999994710160301,
                "computed_survival_10yr": 0.5794357259318426,
            },
        ),
        (
            "single 10y return",
            {
                "rider_type": "Critical",
                "ten_year_survival": 0.40,
                "return_after_10yr": True,
            },
            {
                "use_ten_year": True,
                "use_return_10yr": True,
                "derived_table_rating": 45.11028528213501,
                "computed_survival_5yr": 0.6248768759806841,
                "computed_survival_10yr": 0.3999999767586792,
            },
        ),
        (
            "LE solve",
            {"rider_type": "Chronic", "life_expectancy_years": 8.0},
            {
                "use_le": True,
                "derived_table_rating": 60.4611006565392,
                "computed_le": 7.999999702507566,
            },
        ),
        (
            "direct table/flat add-ons",
            {
                "rider_type": "Critical",
                "table": {"value": 3, "start_year": 1, "stop_year": 4},
                "flat": {"value": 1.25, "start_year": 2, "stop_year": 5},
                "table_2": {"value": 2, "start_year": 5, "stop_year": 7},
                "flat_2": {"value": 0.75, "start_year": 7, "stop_year": 9},
                "increased_decrement": {"value": 200, "start_year": 9, "stop_year": 11},
            },
            {
                "use_table": True,
                "use_flat": True,
                "use_table_2": True,
                "use_flat_2": True,
                "use_increased_decrement": True,
                "derived_table_rating": 3.0,
                "derived_flat_extra": 1.25,
                "computed_survival_5yr": 0.9409489314184816,
            },
        ),
        (
            "terminal",
            {"rider_type": "Terminal"},
            {
                "rider_type": "Terminal",
                "computed_survival_5yr": 0.03125,
                "computed_survival_10yr": 0.0009765625,
            },
        ),
    ],
)
def test_assessment_branches_characterized(policy, name, assessment, expected):
    with using_quote_database(Rates()):
        result = _assessment_port(policy, QuoteRequest.from_dict(_quote_request(assessment)))

    for field, value in expected.items():
        actual = getattr(result, field)
        if isinstance(value, float):
            assert actual == pytest.approx(value, rel=0, abs=1e-9), name
        else:
            assert actual == value, name


@pytest.mark.parametrize(
    ("product", "assessment", "options", "expected"),
    [
        (
            "TERM",
            {"rider_type": "Terminal"},
            {},
            {
                "full_eligible_db": 100000,
                "full_actuarial_discount": 6789.55,
                "full_accel_benefit": 92960.45,
                "partial_eligible_db": 50000,
                "partial_actuarial_discount": 3394.78,
                "premium_before": "$260.00 Annual",
            },
        ),
        (
            "TERM",
            {"rider_type": "Chronic", "five_year_survival": 0.60},
            {},
            {
                "full_eligible_db": 100000,
                "full_actuarial_discount": 36216.54,
                "full_accel_benefit": 63533.46,
                "partial_eligible_db": 50000,
                "partial_actuarial_discount": 18108.27,
                "premium_before": "$260.00 Annual",
            },
        ),
        (
            "UL",
            {"rider_type": "Terminal"},
            {"level_annual_premium": 1000, "loan_payoff": 50, "surrender_value": 1234},
            {
                "full_eligible_db": 100000,
                "full_actuarial_discount": 7766.70,
                "full_loan_repayment": 50,
                "full_surrender_value": 1234,
                "full_accel_benefit": 91933.30,
                "premium_before": "$0.00",
            },
        ),
    ],
)
def test_quote_pipeline_snapshot_values(policy, product, assessment, options, expected):
    p = deepcopy(policy)
    p.product_type = product
    result = calculate_quote(
        _quote_request(assessment, options),
        p,
        database=Rates(),
        policy_provenance={"source": "synthetic"},
        eligible_riders={"Terminal", "Chronic", "Critical"},
    )

    numbers = result["numbers"]
    for field, value in expected.items():
        actual = numbers[field]
        if isinstance(value, float):
            assert actual == pytest.approx(value, rel=0, abs=0.01)
        else:
            assert actual == value
    assert result["assessment"]["rider_type"] == assessment["rider_type"]
    assert result["html"].startswith("<html>")
    assert result["text"].startswith("ABR Quote Summary")


def _workbook_policy() -> ABRPolicyData:
    return ABRPolicyData(
        policy_number="SYNTHETIC",
        company="01",
        region="CKPR",
        insured_name="Jane Example",
        product_type="TERM",
        plan_code="TEST",
        base_plancode="TEST",
        issue_state="TX",
        issue_date=date(2020, 1, 15),
        face_amount=100000,
        min_face_amount=50000,
        issue_age=40,
        attained_age=46,
        maturity_age=75,
        sex="F",
        rate_sex="F",
        rate_class="N",
        policy_year=7,
        policy_month=8,
        table_rating=1,
        flat_extra=0,
        reinsurers="(none)",
        billing_mode=1,
        modal_premium=260,
    )


def _workbook_result() -> ABRQuoteResult:
    return ABRQuoteResult(
        full_eligible_db=100000,
        full_actuarial_discount=1435.02,
        full_admin_fee=250,
        full_accel_benefit=98314.98,
        full_accelerated_benefit=98314.98,
        full_benefit_ratio=0.9831498,
        partial_eligible_db=50000,
        partial_actuarial_discount=717.51,
        partial_admin_fee=250,
        partial_accel_benefit=49032.49,
        partial_accelerated_benefit=49032.49,
        partial_benefit_ratio=0.9806498,
        premium_before="260.00 Annual",
        premium_after_full=0.0,
        premium_after_partial="130.00 Annual",
        plan_description="TEST",
        abr_interest_rate=0.05,
        quote_date=date(2026, 9, 3),
        apv_fb=98564.98,
        apv_fp=0.0,
        apv_fd=0.0,
        per_diem_daily=420,
        per_diem_annual=153300,
    )


def _workbook_assessment() -> MedicalAssessment:
    return MedicalAssessment(
        rider_type="Terminal",
        computed_survival_5yr=0.03125,
        computed_survival_10yr=0.0009765625,
        computed_le=1.5,
    )


def _mortality_rows() -> list[dict]:
    return [
        {
            "quote_month": 1,
            "duration_year": 7,
            "month_in_year": 8,
            "attained_age": 46,
            "qx_vbt": 0.5,
            "qx_multiplied": 0.5,
            "qx_improved": 0.5,
            "table_rating_applied": 0,
            "qx_table_rated": 0.5,
            "flat_extra_applied": 0,
            "qx_flat_extra": 0.5,
            "qx_capped": 0.5,
            "qx_monthly": 0.043478260869565216,
            "px_monthly": 0.9565217391304348,
            "cum_survival": 0.9565217391304348,
        }
    ]


def _apv_rows() -> list[dict]:
    return [
        {
            "month": 80,
            "t": 0,
            "qx_monthly": 0.043478260869565216,
            "px_monthly": 0.9565217391304348,
            "tp_x": 1.0,
            "v_benefit": 0.9959424073517916,
            "v_premium": 1.0,
            "death_benefit": 100000,
            "pvdb_t": 4.330184379790398,
            "pvdb_cum": 4.330184379790398,
            "prem_rate": 2.0,
            "pvfp_t": 2.0,
            "pvfp_cum": 2.0,
            "tp_x_end": 0.9565217391304348,
        }
    ]


def _apv_summary() -> dict:
    return {
        "pvfb_raw": 4.330184379790398,
        "cont_mort_adj": 1.0020309426016804,
        "pvfb_adjusted": 4338.978568149642,
        "pvfp": 2.0,
        "actuarial_discount": 95663.02,
        "monthly_rate": 0.0040741237836483535,
        "annual_rate": 0.05,
        "death_benefit": 100000,
    }


def test_output_panel_detail_workbook_characterization(tmp_path, monkeypatch):
    filepath = tmp_path / "output-detail.xlsx"
    panel = SimpleNamespace(
        _policy=_workbook_policy(),
        _result=_workbook_result(),
        _assessment=_workbook_assessment(),
        _derived_values={
            "std_survival_5yr": "0.9700  (97.00%)",
            "mod_survival_5yr": "0.0312  (3.12%)",
            "table_rating": "Annual Mortality = 0.5000",
        },
        _mort_detail=_mortality_rows(),
        _apv_detail=_apv_rows(),
        _apv_summary=_apv_summary(),
        _get_accel_inputs=lambda: (100000, 50000),
        _get_after_partial_deduction=lambda: "",
    )
    monkeypatch.setattr(
        "suiteview.abrquote.ui.output_panel.guard_support_files_writable",
        lambda *args, **kwargs: None,
    )

    OutputPanel._write_detail_workbook(panel, str(filepath))

    workbook = openpyxl.load_workbook(filepath, data_only=True)
    assert workbook.sheetnames == [
        "Policy Info",
        "Assessment",
        "Mortality Derivation",
        "Life Expectancy",
        "APV - Present Value",
    ]
    assert workbook["Policy Info"]["A1"].value == "ABR Quote — Policy Information"
    assert workbook["Policy Info"]["B4"].value == "SYNTHETIC"
    assert workbook["Assessment"]["A1"].value == "ABR Quote — Assessment"
    assert workbook["Assessment"]["B4"].value == "Terminal"
    assert workbook["Mortality Derivation"]["A1"].value == "Quote Month"
    assert workbook["Mortality Derivation"]["M2"].value == pytest.approx(0.043478260869565216)
    assert workbook["APV - Present Value"]["A1"].value == "Month"
    assert workbook["APV - Present Value"]["H2"].value == 100000
    workbook.close()


class _StyleProxy:
    def __getattr__(self, _name):
        return self

    def __setattr__(self, _name, _value):
        return None


class _ComCell:
    def __init__(self, worksheet, row, column):
        self._worksheet = worksheet
        self._row = row
        self._column = column
        self.Font = _StyleProxy()
        self.Interior = _StyleProxy()

    @property
    def Value(self):
        return self._worksheet.cell(self._row, self._column).value

    @Value.setter
    def Value(self, value):
        self._worksheet.cell(self._row, self._column).value = value

    @property
    def NumberFormat(self):
        return self._worksheet.cell(self._row, self._column).number_format

    @NumberFormat.setter
    def NumberFormat(self, value):
        self._worksheet.cell(self._row, self._column).number_format = value


class _ComRange:
    def __init__(self, worksheet, start=None, end=None):
        self._worksheet = worksheet
        self._start = start
        self._end = end or start
        self.Font = _StyleProxy()
        self.Interior = _StyleProxy()

    @property
    def Value(self):
        return None

    @Value.setter
    def Value(self, values):
        if not isinstance(values, (list, tuple)) or not self._start:
            return
        start_row, start_col = self._start
        for row_offset, row_values in enumerate(values):
            for col_offset, value in enumerate(row_values):
                self._worksheet.cell(
                    start_row + row_offset,
                    start_col + col_offset,
                    value=value,
                )

    @property
    def NumberFormat(self):
        return None

    @NumberFormat.setter
    def NumberFormat(self, _value):
        return None

    @property
    def HorizontalAlignment(self):
        return None

    @HorizontalAlignment.setter
    def HorizontalAlignment(self, _value):
        return None

    def Merge(self):
        return None

    def AutoFilter(self):
        return None

    def Select(self):
        return None


class _ColumnsProxy:
    def __call__(self, *_args):
        return self

    def __getattr__(self, _name):
        return self

    def __setattr__(self, _name, _value):
        return None

    def AutoFit(self):
        return None


class _ComWorksheet:
    def __init__(self, worksheet):
        self._worksheet = worksheet
        self.Columns = _ColumnsProxy()

    @property
    def Name(self):
        return self._worksheet.title

    @Name.setter
    def Name(self, value):
        self._worksheet.title = value

    def Cells(self, row, column):
        return _ComCell(self._worksheet, row, column)

    def Range(self, start, end=None):
        if isinstance(start, _ComCell):
            start_tuple = (start._row, start._column)
            end_tuple = (end._row, end._column) if isinstance(end, _ComCell) else start_tuple
            return _ComRange(self._worksheet, start_tuple, end_tuple)
        return _ComRange(self._worksheet)

    def Activate(self):
        return None


class _Worksheets:
    def __init__(self, workbook):
        self._workbook = workbook

    def __call__(self, index):
        return _ComWorksheet(self._workbook.worksheets[index - 1])

    def Add(self, After=None):
        if After is not None:
            index = self._workbook.worksheets.index(After._worksheet) + 1
        else:
            index = len(self._workbook.worksheets)
        worksheet = self._workbook.create_sheet(index=index)
        return _ComWorksheet(worksheet)


class _ComWorkbook:
    def __init__(self):
        self._workbook = openpyxl.Workbook()
        self.Worksheets = _Worksheets(self._workbook)


class _Workbooks:
    def __init__(self, excel):
        self._excel = excel

    def Add(self):
        self._excel.current_workbook = _ComWorkbook()
        return self._excel.current_workbook


class _ActiveWindow:
    FreezePanes = False


class _ComExcel:
    def __init__(self):
        self.Workbooks = _Workbooks(self)
        self.ActiveWindow = _ActiveWindow()
        self.current_workbook = None
        self.Visible = False
        self.ScreenUpdating = True


def test_calc_viewer_export_workbook_characterization(tmp_path, monkeypatch):
    filepath = tmp_path / "calc-viewer-detail.xlsx"
    fake_excel = _ComExcel()
    fake_dynamic = SimpleNamespace(Dispatch=lambda _name: fake_excel)
    monkeypatch.setitem(sys.modules, "win32com", SimpleNamespace())
    monkeypatch.setitem(sys.modules, "win32com.client", SimpleNamespace(dynamic=fake_dynamic))
    dialog = SimpleNamespace(
        _policy=_workbook_policy(),
        _result=_workbook_result(),
        _assessment=_workbook_assessment(),
        _derived_values={
            "std_survival_5yr": "0.9700  (97.00%)",
            "mod_survival_5yr": "0.0312  (3.12%)",
            "table_rating": "Annual Mortality = 0.5000",
        },
        _mort_rows=_mortality_rows(),
        _apv_rows=_apv_rows(),
        _apv_summary=_apv_summary(),
        _get_accel_inputs=lambda: (100000, 50000),
        _after_partial_override="",
        _accel_amount_input=100000,
        _min_face_amount_input=50000,
        _warnings=[],
    )

    CalcViewerDialog._on_export(dialog)
    fake_excel.current_workbook._workbook.save(filepath)

    workbook = openpyxl.load_workbook(filepath, data_only=True)
    assert workbook.sheetnames == [
        "Policy Info",
        "Assessment",
        "Mortality Derivation",
        "Life Expectancy",
        "APV Present Value",
    ]
    assert workbook["Policy Info"]["A1"].value == "ABR Quote — Policy Information"
    assert workbook["Policy Info"]["B4"].value == "SYNTHETIC"
    assert workbook["Assessment"]["A1"].value == "ABR Quote — Assessment"
    assert workbook["Assessment"]["B4"].value == "Terminal"
    assert workbook["Mortality Derivation"]["A1"].value == "Quote Month"
    assert workbook["Mortality Derivation"]["M2"].value == pytest.approx(0.043478260869565216)
    assert workbook["APV Present Value"]["A1"].value == "Month"
    assert workbook["APV Present Value"]["H2"].value == 100000
    workbook.close()


def test_explanation_document_text_characterization(policy):
    from suiteview.abrquote.core.abr_explanation import (
        build_explanation,
        explanation_to_html,
    )

    policy.insured_name = "Jane Example"
    result = _workbook_result()
    assessment = _workbook_assessment()

    doc = build_explanation(policy, result, assessment)
    html = explanation_to_html(doc)

    assert doc.title == "How Your Accelerated Benefit Was Determined"
    assert doc.letterhead[0] == "AMERICAN NATIONAL INSURANCE COMPANY"
    assert "Policy SYNTHETIC" in doc.subtitle
    assert doc.salutation == "Dear Jane Example:"
    assert "The mortality basis used in the calculation" in [
        section.heading for section in doc.sections
    ]
    assert "chance of death is 500 per 1,000 (50%)" in html
    assert "$98,314.98" in html


def test_email_summary_renderer_characterization():
    from suiteview.abrquote.core.quote_summary import render_quote_summary
    from suiteview.abrquote.ui.email_print_dialog import EmailPrintDialog

    policy = _workbook_policy()
    result = _workbook_result()
    assessment = _workbook_assessment()
    render = SimpleNamespace(
        _policy=policy,
        _assessment=assessment,
        _result=result,
        _fmt=EmailPrintDialog._fmt,
    )

    sections = EmailPrintDialog._build_summary_sections(render)
    html = EmailPrintDialog._build_clipboard_html(render, sections)
    text = EmailPrintDialog._build_clipboard_text(render, sections)
    core_sections, core_html, core_text = render_quote_summary(policy, result, assessment)

    assert (core_sections, core_html, core_text) == (sections, html, text)
    assert sha256(repr(sections).encode("utf-8")).hexdigest() == (
        "92985d8311176157115b375d15f78f7b602584ae6ac39fa6b2ee08ea0bd61b82"
    )
    assert sha256(html.encode("utf-8")).hexdigest() == (
        "d4a36a7a558c93e8000dc8d714fc91932e50b4902a4a30b9cfbf420cefa9afc1"
    )
    assert sha256(text.encode("utf-8")).hexdigest() == (
        "e224b9b3e64faa177d02cf589494a96f33c1f04904769bd21f6bc4ed903fca77"
    )


def test_quote_pipeline_result_shape_matches_snapshot_contract(policy):
    result = calculate_quote(
        _quote_request({"rider_type": "Terminal"}),
        policy,
        database=Rates(),
        policy_provenance={"source": "synthetic"},
        eligible_riders={"Terminal", "Chronic", "Critical"},
    )

    assert set(result) >= {
        "schema_version",
        "status",
        "html",
        "text",
        "numbers",
        "assessment",
        "policy",
        "provenance",
    }
    assert set(result["numbers"]) == set(asdict(ABRQuoteResult()))
