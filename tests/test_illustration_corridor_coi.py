"""Corridor mortality follows the latest active base segment, not coverage 1."""

from datetime import date

import pytest
from openpyxl import load_workbook

from suiteview.illustration.core import calc_engine
from suiteview.illustration.core.bonus_rates import BonusConfig
from suiteview.illustration.core.calc_engine import IllustrationEngine, ProjectionTiming
from suiteview.illustration.core.monthly_deduction import calculate_deduction
from suiteview.illustration.core.rate_loader import IllustrationRates
from suiteview.illustration.core.target_premium import TargetPremiumResult
from suiteview.illustration.debug.excel_export import export_projection_to_excel
from suiteview.illustration.models.plancode_config import PlancodeConfig
from suiteview.illustration.models.policy_data import CoverageSegment, IllustrationPolicyData


def _case(*, ratchet=False):
    policy = IllustrationPolicyData(
        plancode="TEST",
        issue_date=date(2020, 1, 15),
        valuation_date=date(2026, 9, 15),
        issue_age=40,
        attained_age=46,
        policy_year=7,
        policy_month=9,
        duration=81,
        maturity_age=121,
        db_option="A",
        face_amount=100_000.0,
        account_value=100_000.0,
        segments=[
            CoverageSegment(
                coverage_phase=1, issue_date=date(2020, 1, 15),
                face_amount=60_000.0, units=60.0,
            ),
            CoverageSegment(
                coverage_phase=7, issue_date=date(2024, 1, 15),
                face_amount=40_000.0, units=40.0,
            ),
        ],
    )
    config = PlancodeConfig(
        plancode="TEST", dbd=0.0, gint=0.0, corridor_code=2,
        rachet_banding=ratchet,
        table_rating_factor=0.25,
    )
    rates = IllustrationRates(
        coi=[0.0, 2.39],
        segment_coi={1: [0.0, 2.39], 7: [0.0, 2.55]},
        segment_coi_band1={1: [0.0, 2.39], 7: [0.0, 2.55]},
        segment_coi_band2={1: [0.0, 1.39], 7: [0.0, 1.55]},
        band_break=10_000.0,
    )
    return policy, config, rates


def _deduction(policy, config, rates):
    return calculate_deduction(
        policy.account_value, policy, config, rates,
        rate_year=7, attained_age=45, premiums_to_date=0.0,
        projection_date=policy.valuation_date,
    )


@pytest.mark.parametrize("ratchet", [False, True])
@pytest.mark.parametrize("inactive", [None, "depleted", "matured", "terminated", "future"])
def test_corridor_uses_latest_active_segment(ratchet, inactive):
    policy, config, rates = _case(ratchet=ratchet)
    if inactive:
        last = CoverageSegment(
            coverage_phase=9, face_amount=10_000.0, units=10.0,
            issue_date=date(2025, 1, 15),
        )
        if inactive == "depleted":
            last.face_amount = last.units = 0.0
        elif inactive == "matured":
            last.maturity_date = policy.valuation_date
        elif inactive == "terminated":
            last.status = "T"
        elif inactive == "future":
            last.issue_date = date(2027, 1, 15)
        policy.segments.append(last)
        rates.segment_coi[9] = [0.0, 9.99]
        rates.segment_coi_band1[9] = [0.0, 9.99]
        rates.segment_coi_band2[9] = [0.0, 8.99]

    result = _deduction(policy, config, rates)

    assert result.coi_rate == 2.39
    assert result.nar_corr > 0
    if ratchet:
        assert result.coi_band1_rates_by_coverage["corr"] == 2.55
        assert result.coi_band2_rates_by_coverage["corr"] == 1.55
        expected = (
            result.coi_band1_nar_by_coverage["corr"] * 2.55
            + result.coi_band2_nar_by_coverage["corr"] * 1.55
        ) / 1000.0
    else:
        expected = result.nar_corr / 1000.0 * 2.55
    assert result.coi_charge_corr == pytest.approx(expected)
    assert result.coi_rate_corr == result.coi_rates_by_coverage["cov2"] == 2.55
    assert result.total_coi_charge == pytest.approx(
        sum(result.coi_charges_by_coverage.values()) + expected)
    assert result.av_after_deduction == pytest.approx(
        policy.account_value - result.total_deduction)


@pytest.mark.parametrize("ratchet", [False, True])
def test_corridor_returns_to_prior_segment_after_latest_is_fully_reduced(ratchet):
    policy, config, rates = _case(ratchet=ratchet)
    calc_engine._reduce_base_face(
        policy, 40_000.0, rates, policy.valuation_date, 7,
        charge_scr=False, config=config,
    )
    result = _deduction(policy, config, rates)
    assert result.coi_rate_corr == result.coi_rates_by_coverage["cov1"] == 2.39
    if ratchet:
        assert result.coi_band2_rates_by_coverage["corr"] == 1.39


@pytest.mark.parametrize("ratchet", [False, True])
def test_corridor_reuses_adjusted_rate_and_its_own_coverage_duration(ratchet):
    policy, config, rates = _case(ratchet=ratchet)
    last = policy.segments[-1]
    last.table_rating = 2
    last.flat_extra = 12.11
    rates.segment_coi[7] = [0.0, 8.0, 7.0, 2.55, 6.0]
    rates.segment_coi_band1[7] = list(rates.segment_coi[7])
    result = _deduction(policy, config, rates)
    assert result.coi_rate_corr == pytest.approx(2.55 * 1.5 + 1.0)
    assert result.coi_rate_corr == result.coi_rates_by_coverage["cov2"]


@pytest.mark.parametrize("ratchet", [False, True])
def test_corridor_keeps_zero_rate_on_latest_active_segment(ratchet):
    policy, config, rates = _case(ratchet=ratchet)
    rates.segment_coi[7] = [0.0, 0.0]
    rates.segment_coi_band1[7] = [0.0, 0.0]
    rates.segment_coi_band2[7] = [0.0, 0.0]
    result = _deduction(policy, config, rates)
    assert result.coi_rate_corr == 0.0
    assert result.coi_charge_corr == 0.0


@pytest.mark.parametrize("ratchet", [False, True])
def test_corridor_has_no_rate_when_no_segment_remains_active(ratchet):
    policy, config, rates = _case(ratchet=ratchet)
    for segment in policy.segments:
        segment.face_amount = segment.units = 0.0
    result = _deduction(policy, config, rates)
    assert result.coi_rate_corr == 0.0
    assert result.coi_charge_corr == 0.0


@pytest.mark.parametrize("count", [1, 2, 4, 7])
def test_corridor_selection_has_no_segment_limit_and_does_not_choose_highest_rate(count):
    policy, config, rates = _case()
    policy.segments = [
        CoverageSegment(
            coverage_phase=index, face_amount=100_000.0 / count,
            units=100.0 / count,
        )
        for index in range(1, count + 1)
    ]
    rates.segment_coi = {index: [0.0, 4.0] for index in range(1, count + 1)}
    rates.segment_coi[count] = [0.0, 1.55]
    result = _deduction(policy, config, rates)
    assert result.coi_rate_corr == result.coi_rates_by_coverage[f"cov{count}"] == 1.55
    assert result.coi_charge_corr == pytest.approx(result.nar_corr / 1000.0 * 1.55)


def test_unsegmented_policy_retains_its_base_rate():
    policy, config, rates = _case()
    policy.segments = []
    result = _deduction(policy, config, rates)
    assert result.coi_rate_corr == result.coi_rate == 2.39
    assert result.coi_charge_corr == pytest.approx(result.nar_corr / 1000.0 * 2.39)


@pytest.mark.parametrize("timing", list(ProjectionTiming))
@pytest.mark.parametrize("guaranteed", [False, True])
def test_corridor_rate_survives_projection_and_workbook_export(
    monkeypatch, tmp_path, timing, guaranteed,
):
    policy, config, rates = _case()
    if guaranteed:
        rates.segment_coi = {1: [0.0, 3.39], 7: [0.0, 3.55]}
    expected = 3.55 if guaranteed else 2.55
    monkeypatch.setattr(calc_engine, "load_plancode", lambda _plan: config)
    monkeypatch.setattr(
        calc_engine, "compute_target_premiums",
        lambda *_args, **_kwargs: TargetPremiumResult(),
    )
    states = IllustrationEngine().project(
        policy, months=2, timing=timing, stop_on_lapse=False,
        options=calc_engine.IllustrationOptions(guaranteed_assumption=guaranteed),
        rates_override=rates, bonus_override=BonusConfig(),
    )
    assert len(states) == 3
    for state in states:
        assert state.coi_rate_corr == expected
        assert state.coi_rate_corr == state.coi_rates_by_coverage["cov2"]
        assert state.coi_charge_corr == pytest.approx(
            state.nar_corr / 1000.0 * expected)

    for include_policy in (False, True):
        output = tmp_path / f"corridor-{include_policy}.xlsx"
        export_projection_to_excel(
            states, output, policy_data=policy if include_policy else None)
        workbook = load_workbook(output, read_only=True, data_only=True)
        try:
            sheet = workbook["Projection"]
            rows = sheet.iter_rows(min_row=2, values_only=True)
            headers = next(rows)
            assert headers.count("coi_rate_corr") == 1
            column = headers.index("coi_rate_corr")
            assert [row[column] for row in rows] == [expected] * len(states)
        finally:
            workbook.close()
