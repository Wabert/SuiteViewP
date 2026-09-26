from __future__ import annotations

from decimal import Decimal

import pytest
from openpyxl import Workbook

from suiteview.ratemanager import mortality_loader


def _workbook(tmp_path):
    path = tmp_path / "mortality.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet.cell(2, 5, "Mortality Table")
    sheet.cell(2, 6, "Cyberlife Mortality Table Code")
    sheet.cell(2, 7, "First Age")
    sheet.cell(2, 8, "Last Age")
    for age in range(122):
        sheet.cell(2, 9 + age, age)

    codes = ["UN", "NW", "NX", "NY", "NZ", "N2", "N3"]
    codes.extend(f"T{number:02d}" for number in range(58))
    for row_number, code in enumerate(codes, start=3):
        sheet.cell(row_number, 5, f"Description {code}")
        sheet.cell(row_number, 6, code)
        sheet.cell(row_number, 7, 0)
        sheet.cell(row_number, 8, 120)
        for age in range(121):
            sheet.cell(row_number, 9 + age, min(Decimal(age) / 120, Decimal("1")))
    workbook.save(path)
    return path


def test_parse_applies_overrides_and_converts_rates(tmp_path, monkeypatch):
    monkeypatch.setattr(mortality_loader, "EXPECTED_RATE_COUNT", 7765)
    package = mortality_loader.parse_workbook(_workbook(tmp_path))

    tables = {table.code: table for table in package.tables}
    assert tables["UN"].first_age == 16
    assert all(tables[code].last_age == 99 for code in ("NW", "NX", "NY", "NZ"))
    assert next(
        rate.rate_per_1000
        for rate in package.rates
        if rate.code == "UN" and rate.attained_age == 120
    ) == Decimal("1000")
    assert len(package.rates) == 7765
    assert package.warnings == (
        "N2 and N3 have identical source schedules; loaded as supplied pending correction.",
    )


def test_parse_rejects_missing_in_range_rate(tmp_path, monkeypatch):
    path = _workbook(tmp_path)
    workbook = mortality_loader.load_workbook(path)
    workbook["Rates"].cell(3, 9 + 20).value = None
    workbook.save(path)
    monkeypatch.setattr(mortality_loader, "EXPECTED_RATE_COUNT", 7765)

    with pytest.raises(ValueError, match="UN: missing or nonnumeric qx at attained age 20"):
        mortality_loader.parse_workbook(path)


def test_database_rows_bind_rates_as_floats():
    rows = mortality_loader._rate_rows([
        mortality_loader.MortalityRate("N", 20, Decimal("1.23456789"))
    ])

    assert rows == [("N", 20, 1.23456789)]
    assert isinstance(rows[0][2], float)

def test_locked_workbook_falls_back_to_dispatchex_when_excel_not_running(monkeypatch, tmp_path):
    import sys
    import types

    import pywintypes

    def com_fail(*_args, **_kwargs):
        raise pywintypes.com_error(-2147221021, "Operation unavailable", None, None)

    class FakeRange:
        Value2 = ((1, 2), (3, 4))

    class FakeWorkbook:
        closed = False

        def Worksheets(self, _name):
            return types.SimpleNamespace(Range=lambda _ref: FakeRange())

        def Close(self, SaveChanges):
            FakeWorkbook.closed = True

    class FakeExcel:
        quit_called = False

        def __init__(self):
            self.Workbooks = types.SimpleNamespace(Open=lambda *_a, **_k: FakeWorkbook())

        def Quit(self):
            FakeExcel.quit_called = True

    client = types.ModuleType("win32com.client")
    client.GetObject = com_fail
    client.GetActiveObject = com_fail
    client.DispatchEx = lambda _prog_id: FakeExcel()
    package = types.ModuleType("win32com")
    package.client = client
    monkeypatch.setitem(sys.modules, "win32com", package)
    monkeypatch.setitem(sys.modules, "win32com.client", client)

    def locked(*_args, **_kwargs):
        raise PermissionError("locked by Excel")

    monkeypatch.setattr(mortality_loader, "load_workbook", locked)

    rows = mortality_loader._read_source_rows(tmp_path / "locked.xlsx")

    assert rows == ((1, 2), (3, 4))
    assert FakeWorkbook.closed and FakeExcel.quit_called
