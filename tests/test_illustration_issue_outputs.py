"""Issue-run outputs carry their modeling basis without changing inforce grids."""
import io
import os
from datetime import date, datetime

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import load_workbook
from PyQt6.QtWidgets import QApplication, QTreeWidgetItem

from suiteview.illustration.core.report_builder import (
    ISSUE_OUTPUT_LABEL,
    build_ul_report,
    issue_output_basis,
)
from suiteview.illustration.debug.summary_export import (
    CURRENT_SHEET,
    GUARANTEED_SHEET,
    build_summary_workbook,
)
from suiteview.illustration.models.calc_state import MonthlyState
from suiteview.illustration.models.policy_data import BenefitInfo, IllustrationPolicyData, RiderInfo
from suiteview.illustration.ui.report_tab import format_report_pages
from suiteview.illustration.ui.values_tab import (
    LEDGER_COLUMNS,
    _ExportLabeledOverview,
    _export_values_summary,
    _issue_export_rows,
)


def _policy(issue=True):
    return IllustrationPolicyData(
        policy_number="TEST123", plancode="1U143900", insured_name="TEST INSURED",
        company_code="01", issue_date=date(2019, 11, 9), issue_age=50,
        attained_age=50, valuation_date=date(2019, 10, 9),
        illustration_date=date(2026, 9, 13), run_from_issue=issue,
        face_amount=100000, account_value=0, premiums_paid_to_date=0,
        benefits=[BenefitInfo(benefit_type="3", benefit_subtype="9", is_active=True)],
    )


def _results():
    return [
        MonthlyState(
            date=date(2019, 10, 9), policy_year=0, policy_month=12,
            duration=0, attained_age=50, gsp=30000, glp=2500,
        ),
        MonthlyState(
            date=date(2019, 11, 9), policy_year=1, policy_month=1,
            duration=1, attained_age=50, annual_interest_rate=0.06,
            av_end_of_month=125, ending_sv=100, ending_db=100000,
        ),
    ]


def test_issue_report_labels_every_page_and_uses_issue_not_technical_date():
    report = build_ul_report(
        _policy(), _results(), guaranteed_results=_results(),
        run_date=date(2026, 9, 14),
    )
    pages = format_report_pages(report, include_expense_report=True)
    for page in pages:
        text = "\n".join(page)
        assert ISSUE_OUTPUT_LABEL in text
        assert "ORIGINAL ISSUE DATE: 11/09/2019" in text
        assert "RATES BASIS DATE: 09/13/2026" in text
        assert "CURRENT SCALE 1; NOT A HISTORICAL RECONSTRUCTION" in text
        assert "CONTRACTUAL GUARANTEED ASSUMPTIONS" in text
        assert "10/09/2019" not in text
        assert all(len(line) <= 112 for line in page)
    text = "\n".join("\n".join(page) for page in pages)
    assert "HYPOTHETICAL INFORCE ILLUSTRATION" not in text
    assert "CURRENT SPECIFIED AMOUNT:" not in text
    assert "ACTUAL PREMIUMS PAID:" not in text
    assert "MODELED SPECIFIED AMOUNT:" in text
    assert "MODELED ISSUE OPENING ACCUMULATION VALUE: $0.00" in text
    assert "RIDERS AND BENEFITS INCLUDED IN THE MODELED ISSUE CONDITIONS" in text
    assert "PREMIUM WAIVER" in text
    assert report.as_of_date == _policy().issue_date
    assert report.ledger[0].guar_accum == 125


def test_inforce_report_retains_existing_identity_and_shape():
    report = build_ul_report(_policy(False), _results())
    text = "\n".join("\n".join(page) for page in format_report_pages(report))
    assert report.basis_lines == []
    assert issue_output_basis(_policy(False)) == []
    assert "HYPOTHETICAL INFORCE ILLUSTRATION" in text
    assert "CURRENT SPECIFIED AMOUNT:" in text
    assert "ACTUAL PREMIUMS PAID:" in text
    assert "AS OF 10/09/2019" in text
    assert ISSUE_OUTPUT_LABEL not in text


def test_missing_rates_basis_is_explicit_not_replaced_with_run_or_issue_date():
    policy = _policy()
    policy.illustration_date = None
    report = build_ul_report(policy, _results(), run_date=date(2026, 9, 14))
    assert "RATES BASIS DATE: NOT PROVIDED" in "\n".join(report.basis_lines)


@pytest.mark.parametrize("issue", [False, True])
def test_summary_export_preserves_data_sheets_and_identifies_issue_basis(monkeypatch, issue):
    from openpyxl import Workbook

    policy = _policy(issue)
    captured = {}
    save = Workbook.save

    def capture_save(workbook, path):
        stream = io.BytesIO()
        save(workbook, stream)
        captured["bytes"] = stream.getvalue()

    monkeypatch.setattr(Workbook, "save", capture_save)
    path = _export_values_summary(policy, _results(), _results(), "tests")
    actual = load_workbook(io.BytesIO(captured["bytes"]))
    expected = build_summary_workbook(policy, _results(), _results())
    try:
        for name in (CURRENT_SHEET, GUARANTEED_SHEET):
            normalized = [
                tuple(value.date() if isinstance(value, datetime) else value for value in row)
                for row in actual[name].values
            ]
            assert normalized == list(expected[name].values)
            assert actual[name].freeze_panes == expected[name].freeze_panes
        if issue:
            assert actual.sheetnames[0] == "Illustration Basis"
            assert ISSUE_OUTPUT_LABEL in path.name
            text = str(list(actual["Illustration Basis"].values))
            assert ISSUE_OUTPUT_LABEL in text
            assert "11/09/2019" in text and "09/13/2026" in text
            assert "CONTRACTUAL GUARANTEED" in text
            assert "not historical inforce balances" in text
        else:
            assert actual.sheetnames == expected.sheetnames
            assert ISSUE_OUTPUT_LABEL not in path.name
    finally:
        actual.close()
        expected.close()


@pytest.fixture(scope="session")
def app():
    return QApplication.instance() or QApplication([])


@pytest.mark.parametrize("include_months", [False, True])
@pytest.mark.parametrize("guaranteed", [False, True])
def test_ledger_excel_has_basis_sheet_and_unchanged_columns(
    app, monkeypatch, include_months, guaranteed,
):
    from unittest.mock import Mock
    from suiteview.core import excel_export

    excel, workbook, basis_sheet, ledger_sheet = Mock(), Mock(), Mock(), Mock()
    workbook.Worksheets.Add.return_value = ledger_sheet
    dump = Mock(return_value=(excel, workbook, basis_sheet))
    write = Mock()
    monkeypatch.setattr(excel_export, "dump_to_new_workbook", dump)
    monkeypatch.setattr(excel_export, "write_table", write)
    overview = _ExportLabeledOverview()
    try:
        overview.display(_policy(), [])
        overview.export_guaranteed = guaranteed
        annual = ["annual"] * len(LEDGER_COLUMNS)
        monthly = ["monthly"] * len(LEDGER_COLUMNS)
        item = QTreeWidgetItem(annual)
        item.addChild(QTreeWidgetItem(monthly))
        overview.ledger.addTopLevelItem(item)
        overview._dump_ledger(include_months)
        metadata = str(dump.call_args.args[1])
        assert ISSUE_OUTPUT_LABEL in metadata
        assert "11/09/2019" in metadata and "09/13/2026" in metadata
        assert ("Displayed values', 'CONTRACTUAL GUARANTEED" in metadata) == guaranteed
        assert write.call_args.args[1] == LEDGER_COLUMNS
        assert write.call_args.args[2] == ([annual, monthly] if include_months else [annual])
        overview.clear()
        assert overview._export_policy is None
    finally:
        overview.close()


def test_inforce_ledger_excel_retains_single_sheet_export(app, monkeypatch):
    from unittest.mock import Mock
    from suiteview.core import excel_export

    dump = Mock()
    monkeypatch.setattr(excel_export, "dump_to_new_workbook", dump)
    overview = _ExportLabeledOverview()
    try:
        overview.display(_policy(False), [])
        row = ["inforce"] * len(LEDGER_COLUMNS)
        overview.ledger.addTopLevelItem(QTreeWidgetItem(row))
        overview._dump_ledger(False)
        dump.assert_called_once_with(
            LEDGER_COLUMNS, [row], sheet_name="Illustration Ledger")
    finally:
        overview.close()


def test_issue_export_basis_uses_edited_conditions_and_only_included_riders():
    policy = _policy()
    policy.face_amount = 175000
    policy.db_option = "B"
    policy.riders = [
        RiderInfo(
            coverage_phase=2, plancode="TERMPLAN", description="EDITED TERM",
            face_amount=50000, is_active=True,
        ),
        RiderInfo(description="EXCLUDED RIDER", is_active=False),
    ]
    policy.benefits.append(BenefitInfo(benefit_type="1", is_active=False))
    rows = _issue_export_rows(policy)
    assert ("Modeled specified amount", 175000) in rows
    assert ("Modeled DB option", "B - Increasing Death Benefit") in rows
    text = str(rows)
    assert "Phase 2: EDITED TERM (TERMPLAN); face $50,000.00" in text
    assert "PREMIUM WAIVER (39)" in text
    assert "EXCLUDED RIDER" not in text
    assert "ACCIDENTAL DEATH BENEFIT" not in text
    policy.riders = []
    policy.benefits = []
    assert ("Included riders / benefits", "NONE") in _issue_export_rows(policy)


@pytest.mark.parametrize("failure_stage", ["helper", "ledger"])
def test_issue_ledger_excel_failures_log_traceback_and_notify(
    app, monkeypatch, caplog, failure_stage,
):
    from unittest.mock import Mock
    from pywintypes import com_error
    from suiteview.core import excel_export
    from suiteview.illustration.ui import values_tab

    excel, workbook, basis_sheet = Mock(), Mock(), Mock()
    dump = Mock(return_value=(excel, workbook, basis_sheet))
    if failure_stage == "helper":
        dump.side_effect = excel_export.ExcelExportError("Excel unavailable")
    else:
        workbook.Worksheets.Add.side_effect = com_error(-1, "Write failed", None, None)
    warning = Mock()
    monkeypatch.setattr(excel_export, "dump_to_new_workbook", dump)
    monkeypatch.setattr(values_tab.QMessageBox, "warning", warning)
    overview = _ExportLabeledOverview()
    try:
        overview.display(_policy(), [])
        overview._dump_ledger(False)
        warning.assert_called_once()
        assert "Could not export:" in warning.call_args.args[2]
        assert any(
            record.exc_info and "Excel export failed" in record.message
            for record in caplog.records
        )
        if failure_stage == "ledger":
            assert excel.ScreenUpdating is True
    finally:
        overview.close()
